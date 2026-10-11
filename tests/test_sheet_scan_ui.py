"""真实工作簿验证总览联动、后台进度、取消、缓存和现有清洗流程。"""

import hashlib
import threading
import time
from unittest.mock import patch

from openpyxl import Workbook
import pandas as pd
import pytest
from PySide6.QtCore import QTimer, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from app.core.detector import analyze_column
from app.ui.main_window import MainWindow
from app.ui.sheet_overview_panel import STALE_SCAN_MESSAGE


def wait_until(predicate, timeout=10):
    deadline = time.monotonic() + timeout
    while not predicate():
        QApplication.processEvents()
        QTest.qWait(5)
        assert time.monotonic() < deadline, "后台任务超时"
    QApplication.processEvents()


def wait_for_idle(window):
    wait_until(lambda: not window._is_busy())


@pytest.fixture
def window(tmp_path, qapp):
    path = tmp_path / "整表扫描测试.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待清洗"
    sheet.append(["商品编码", "客户名称", "备注", "全空字段"])
    for index in range(100):
        sheet.append([
            f" P{index:03d}" if index < 2 else f"P{index:03d}",
            f"客户{index}",
            f"备注{index}\n" if index < 10 else f"备注{index}",
            None,
        ])
    other = workbook.create_sheet("其他表")
    other.append(["名称"])
    other.append(["甲"])
    other.append(["乙"])
    workbook.create_sheet("仅表头").append(["编码", "名称"])
    workbook.save(path)
    workbook.close()
    widget = MainWindow()
    widget.show()
    qapp.processEvents()
    widget._start_load(str(path))
    wait_for_idle(widget)
    yield widget
    wait_for_idle(widget)
    widget.close()
    widget.deleteLater()
    qapp.processEvents()


def scan(window):
    QTest.mouseClick(window.scan_button, Qt.MouseButton.LeftButton)
    assert window._scan_worker is not None
    wait_for_idle(window)


def click_column(window, physical_column):
    panel = window.sheet_overview_panel
    window.results_tabs.setCurrentWidget(panel)
    QApplication.processEvents()
    for row in range(panel.proxy_model.rowCount()):
        index = panel.proxy_model.index(row, 0)
        if panel.proxy_model.data(index, Qt.ItemDataRole.UserRole) == physical_column:
            panel.table.scrollTo(index)
            QApplication.processEvents()
            QTest.mouseClick(panel.table.viewport(), Qt.MouseButton.LeftButton, pos=panel.table.visualRect(index).center())
            wait_for_idle(window)
            return
    raise AssertionError("未找到字段")


def test_scan_populates_full_overview_without_changes_or_source_writes(window):
    source = window.data_manager.file_path
    digest = hashlib.sha256(source.read_bytes()).digest()
    before = window.data_manager.working_df.copy(deep=True)
    scan(window)
    panel = window.sheet_overview_panel
    result = panel.result
    assert window.results_tabs.currentWidget() is panel
    assert result.column_count == 4
    assert result.problem_column_count == 3
    assert result.normal_column_count == 1
    assert result.columns[0].outer_spaces == 2
    assert result.columns[2].hidden_chars == 10
    assert result.columns[3].empty == 100
    assert "字段数：4" in panel.summary_label.text()
    assert "扫描完成" in window.statusBar().currentMessage()
    assert panel.model.columnCount() == 15
    assert panel.table.isEnabled()
    assert panel.proxy_model.data(panel.proxy_model.index(0, 0)) == "全空字段"
    assert panel.proxy_model.data(panel.proxy_model.index(0, 8)) == 0
    assert not panel.model.flags(panel.model.index(0, 0)) & Qt.ItemFlag.ItemIsEditable
    assert window.data_manager.current_data.scan_cache.result is result
    assert not window.data_manager.undo_stack
    assert not window.export_button.isEnabled()
    pd.testing.assert_frame_equal(window.data_manager.working_df, before)
    assert hashlib.sha256(source.read_bytes()).digest() == digest


def test_sorted_overview_click_reuses_diagnosis_and_highlights_original_column(window):
    scan(window)
    window.sheet_overview_panel.table.sortByColumn(2, Qt.SortOrder.DescendingOrder)
    click_column(window, 0)
    assert window.selected_column == 0
    assert window.selected_column_label.text() == "当前选择列：商品编码"
    assert [index.column() for index in window.table.selectionModel().selectedColumns()] == [0]
    assert window.results_tabs.currentWidget() is window.diagnosis_panel
    assert window.diagnosis_panel.result.counts["前后空格"] == 2
    assert window.cleaning_panel.rule_boxes["outer_spaces"].isChecked()
    assert window.cleaning_panel.generate_button.isEnabled()
    assert window.diagnosis_panel.result.issues[0].excel_row == 2


def test_numeric_sorting_and_duplicate_header_selection(window):
    # 工作簿一般由 pandas 为重名表头加后缀；模型仍支持真正重名的列。
    frame = pd.DataFrame([[" A", "正常", "坏\n"], ["B", "不同", "坏\n"]], columns=["重名", "正常", "重名"])
    window.data_manager.current_data.original_df = frame.copy(deep=True)
    window.data_manager.current_data.working_df = frame
    window._display_sheet(window._preview)
    scan(window)
    panel = window.sheet_overview_panel
    panel.table.sortByColumn(2, Qt.SortOrder.AscendingOrder)
    assert [panel.proxy_model.data(panel.proxy_model.index(row, 2)) for row in range(3)] == [0, 1, 2]
    click_column(window, 2)
    assert window.selected_column == 2
    assert window.diagnosis_panel.result.counts["Tab / 换行"] == 2
    assert window.diagnosis_panel.result.counts["重复值"] == 2


def test_background_progress_keeps_event_loop_responsive_and_locks_mutating_actions(window):
    ticks = []
    progress = []
    timer = QTimer()
    timer.setInterval(5)
    timer.timeout.connect(lambda: ticks.append(1))
    real_detector = analyze_column

    def slower_detector(*args, **kwargs):
        time.sleep(0.03)
        return real_detector(*args, **kwargs)

    window._select_column(0)
    timer.start()
    try:
        with patch("app.core.sheet_analyzer.analyze_column", side_effect=slower_detector):
            window._start_sheet_scan()
            window._scan_worker.progress.connect(lambda *event: progress.append(event))
            assert not window.scan_button.isEnabled()
            assert not window.open_button.isEnabled()
            assert not window.sheet_combo.isEnabled()
            assert not window.table.isEnabled()
            assert not window.start_button.isEnabled()
            window._select_column(1)
            assert window.selected_column == 0
            wait_for_idle(window)
    finally:
        timer.stop()
    assert len(ticks) >= 3
    assert [event[0] for event in progress] == [1, 2, 3, 4]
    assert window.sheet_overview_panel.progress_bar.value() == 4
    assert window.scan_button.isEnabled()
    assert window.table.isEnabled()


def test_cancel_discards_partial_result_and_allows_retry(window):
    started = threading.Event()
    release = threading.Event()

    def blocked_detector(*args, **kwargs):
        started.set()
        assert release.wait(5)
        return analyze_column(*args, **kwargs)

    try:
        with patch("app.core.sheet_analyzer.analyze_column", side_effect=blocked_detector):
            window._start_sheet_scan()
            wait_until(started.is_set)
            QTest.mouseClick(window.sheet_overview_panel.cancel_button, Qt.MouseButton.LeftButton)
            assert not window.sheet_overview_panel.cancel_button.isEnabled()
            release.set()
            wait_for_idle(window)
    finally:
        release.set()
    assert window.data_manager.current_data.scan_cache is None
    assert window.sheet_overview_panel.result is None
    assert "扫描已取消" in window.sheet_overview_panel.state_label.text()
    scan(window)
    assert window.sheet_overview_panel.result is not None


def test_cancel_preserves_previous_complete_cache(window):
    scan(window)
    cached = window.data_manager.current_data.scan_cache
    with patch("app.core.sheet_analyzer.analyze_column", side_effect=lambda *a, **k: (time.sleep(0.03), analyze_column(*a, **k))[1]):
        window._start_sheet_scan()
        assert "上次完整扫描" in window.sheet_overview_panel.summary_label.text()
        window._cancel_sheet_scan()
        wait_for_idle(window)
    assert window.data_manager.current_data.scan_cache is cached
    assert window.sheet_overview_panel.result is cached.result
    assert "上次完整扫描" in window.sheet_overview_panel.state_label.text()


def test_late_cancel_before_queued_result_is_handled_does_not_publish(window):
    window._start_sheet_scan()
    assert window._scan_worker.wait(5000)
    # 线程已结束但 UI 尚未处理成功信号：取消也必须生效。
    window._cancel_sheet_scan()
    wait_for_idle(window)
    assert window.data_manager.current_data.scan_cache is None
    assert "扫描已取消" in window.statusBar().currentMessage()


def test_failed_scan_retains_last_complete_result_and_restores_controls(window):
    scan(window)
    cache = window.data_manager.current_data.scan_cache
    with patch("app.core.sheet_analyzer.analyze_column", side_effect=RuntimeError("测试异常")):
        with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok) as error:
            scan(window)
    error.assert_called_once()
    assert window.data_manager.current_data.scan_cache is cache
    assert window.sheet_overview_panel.result is cache.result
    assert "扫描失败" in window.sheet_overview_panel.state_label.text()
    assert window.scan_button.isEnabled()
    assert window.open_button.isEnabled()


def test_sheet_switch_reuses_valid_cache_and_new_file_resets_it(window, tmp_path):
    scan(window)
    first_cache = window.data_manager.current_data.scan_cache
    window.sheet_combo.setCurrentText("其他表")
    wait_for_idle(window)
    assert window.sheet_overview_panel.result is None
    scan(window)
    second_cache = window.data_manager.current_data.scan_cache
    with patch("app.ui.sheet_analysis_worker.analyze_sheet") as detector:
        window.sheet_combo.setCurrentText("待清洗")
        wait_for_idle(window)
        assert window.sheet_overview_panel.result is first_cache.result
        window.sheet_combo.setCurrentText("其他表")
        wait_for_idle(window)
        assert window.sheet_overview_panel.result is second_cache.result
    detector.assert_not_called()
    # 显式重新打开相同文件也应开始新会话。
    window._start_load(str(window.data_manager.file_path))
    wait_for_idle(window)
    assert window.data_manager.current_data.scan_cache is None
    assert window.sheet_overview_panel.result is None


def test_apply_undo_restore_mark_cache_stale_and_rescan_uses_working_data(window):
    digest = hashlib.sha256(window.data_manager.file_path.read_bytes()).digest()
    scan(window)
    click_column(window, 0)
    window._start_cleaning_preview()
    wait_for_idle(window)
    with patch("app.ui.main_window.confirm_apply", return_value=True):
        window._apply_cleaning()
    wait_for_idle(window)
    panel = window.sheet_overview_panel
    assert panel.state_label.text() == STALE_SCAN_MESSAGE
    assert not panel.table.isEnabled()
    assert "待更新" in window.results_tabs.tabText(window.results_tabs.indexOf(panel))
    scan(window)
    assert panel.result.columns[0].outer_spaces == 0
    assert panel.result.columns[0].quality_score == 100
    assert window.export_button.isEnabled()  # 原导出入口仍可用。
    with patch("app.ui.main_window.confirm_restore", return_value=True):
        window._restore_original()
    wait_for_idle(window)
    assert panel.state_label.text() == STALE_SCAN_MESSAGE
    scan(window)
    assert panel.result.columns[0].outer_spaces == 2
    window._undo_last_action()
    wait_for_idle(window)
    assert panel.state_label.text() == STALE_SCAN_MESSAGE
    scan(window)
    assert panel.result.columns[0].outer_spaces == 0
    window._undo_last_action()
    wait_for_idle(window)
    assert panel.state_label.text() == STALE_SCAN_MESSAGE
    assert hashlib.sha256(window.data_manager.file_path.read_bytes()).digest() == digest


def test_header_only_sheet_scans_without_selection(window):
    window.sheet_combo.setCurrentText("仅表头")
    wait_for_idle(window)
    assert window.selected_column is None
    scan(window)
    result = window.sheet_overview_panel.result
    assert result.row_count == 0
    assert result.column_count == 2
    assert result.problem_column_count == 0
    click_column(window, 1)
    assert window.selected_column == 1
    assert window.diagnosis_panel.result.row_count == 0


def test_close_during_scan_requests_cancel_and_waits_for_thread(window):
    with patch("app.core.sheet_analyzer.analyze_column", side_effect=lambda *a, **k: (time.sleep(0.05), analyze_column(*a, **k))[1]):
        window._start_sheet_scan()
        window.close()
        assert window._close_requested
        wait_for_idle(window)
    assert not window.isVisible()


def test_no_workbook_cannot_start_scan(qapp):
    window = MainWindow()
    try:
        assert not window.scan_button.isEnabled()
        window._start_sheet_scan()
        assert window._scan_worker is None
    finally:
        window.close()
        window.deleteLater()
        qapp.processEvents()
