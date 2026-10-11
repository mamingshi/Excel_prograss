"""重复辅助列预览、命名、兼容公式和只读保证。"""

from datetime import datetime

from openpyxl import Workbook
import pandas as pd
import pytest

from app.core.annotation_builder import (
    build_duplicate_formula, build_duplicate_values,
    generate_annotation_name, generate_duplicate_preview,
)
from tests.test_export_formulas import _calculate_excel


def test_requested_example_and_preview_only():
    frame = pd.DataFrame({"商品编码": ["P001", "P002", "P001", "P003"]})
    before = frame.copy(deep=True)
    preview = generate_duplicate_preview(frame, 0)
    assert preview.values.tolist() == ["重复", "正常", "重复", "正常"]
    assert preview.annotation.column_name == "商品编码_重复检查"
    assert preview.duplicate_count == 2
    assert preview.empty_count == 0
    pd.testing.assert_frame_equal(frame, before)


def test_missing_and_zero_keep_original_detector_semantics():
    values = pd.Series([None, "", float("nan"), pd.NA, pd.NaT, 0, "0", 1002, "1002"], dtype=object)
    assert build_duplicate_values(values).tolist() == ["空值"] * 5 + ["正常"] * 4


def test_excel_date_numeric_alias_is_rejected_instead_of_mislabelled():
    from openpyxl.utils.datetime import to_excel
    value = datetime(2026, 1, 1)
    with pytest.raises(ValueError, match="日期和数值"):
        build_duplicate_values(pd.Series([value, to_excel(value)], dtype=object))


def test_case_spaces_and_wildcards_are_compared_literally():
    values = pd.Series(["P001", "p001", "P001 ", "*", "?", "~", "P001"], dtype=object)
    assert build_duplicate_values(values).tolist() == ["重复", "正常", "正常", "正常", "正常", "正常", "重复"]


def test_numeric_equality_including_boolean_matches_existing_detector():
    values = pd.Series([1, 1.0, True, "1", False, 0, "0"], dtype=object)
    assert build_duplicate_values(values).tolist() == ["重复", "重复", "重复", "正常", "重复", "重复", "正常"]


def test_name_collision_starts_at_two_and_preserves_existing_fields():
    frame = pd.DataFrame({
        "商品编码": ["A", "A"], "商品编码_重复检查": ["保留", "保留"],
        "商品编码_重复检查_2": ["原有", "原有"],
    })
    preview = generate_duplicate_preview(frame, 0)
    assert preview.annotation.column_name == "商品编码_重复检查_3"
    assert frame["商品编码_重复检查"].tolist() == ["保留", "保留"]
    assert generate_annotation_name("Code", ["code_重复检查"]) == "Code_重复检查_2"


def test_nondefault_and_duplicate_indexes_do_not_change_excel_row_mapping():
    frame = pd.DataFrame({"左": [1, 2, 3], "编码": ["A", "B", "A"]}, index=[99, 99, -8])
    result = generate_duplicate_preview(frame, 1)
    assert result.values.index.tolist() == [99, 99, -8]
    assert result.values.tolist() == ["重复", "正常", "重复"]
    assert 'IF(B2="","空值",' in result.formula(0)
    assert 'IF(B4="","空值",' in result.formula(2)
    assert "$B$2:$B$4" in result.formula(0)


def test_formula_uses_actual_column_and_bounded_data_range():
    formula = build_duplicate_formula(3, 15, 87)
    assert formula.startswith('=IFERROR(IF(D15="","空值",')
    assert "$D$2:$D$87" in formula
    assert "SUMPRODUCT" in formula and "EXACT" in formula and "ISTEXT" in formula
    assert "XLOOKUP" not in formula and "_xlfn." not in formula and "LET(" not in formula
    assert "$AA$2:$AA$9" in build_duplicate_formula(26, 2, 9)
    # 用户截图中的同一坐标原为 298 字符，缩减后应控制在 220 字符以内。
    assert len(build_duplicate_formula(6, 6, 21)) <= 220


@pytest.mark.parametrize("arguments", [(-1, 2, 4), (16384, 2, 4), (0, 1, 4), (0, 5, 4), (0, 2, 1048577)])
def test_invalid_excel_coordinates_rejected(arguments):
    with pytest.raises(ValueError):
        build_duplicate_formula(*arguments)


def test_empty_preview_is_safe():
    preview = generate_duplicate_preview(pd.DataFrame(columns=["编码"]), 0)
    assert preview.row_count == 0
    assert preview.values.empty
    with pytest.raises(ValueError):
        preview.formula(0)


def test_real_excel_formula_matches_preview_for_types_case_and_literal_characters(tmp_path):
    datasets = [
        ["P001", "P002", "P001", "P003", None, "", 0, "0", "p001", "*", "?", "~", "=1+1", "商品编码"],
        [0, "0", 1, 1.0, True, "1", False, None, "", 1.00000000000001],
        ["TRUE", "FALSE", "1", "01", "1.0", "0", "00", "-0", "1e0", " 1 ",
         1, 1.0, True, 0, False, 2, "2", 2.5, "2.5"],
        ["A\u00a0", "A ", "A\u00a0", "A", "A\tB", "A\nB", "ABC  123", "ABC 123", "😀", "😀", "中", "中"],
        [datetime(2026, 1, 1), datetime(2026, 1, 1), datetime(2026, 1, 2), None],
        ["x" * 30000, "x" * 30000, "x" * 29999, None],
    ]
    book = Workbook()
    book.remove(book.active)
    expected, addresses = [], []
    for number, values in enumerate(datasets):
        sheet = book.create_sheet(f"样本{number}")
        sheet.append(["商品编码", "辅助列"])
        preview = generate_duplicate_preview(pd.DataFrame({"商品编码": pd.Series(values, dtype=object)}), 0)
        expected.extend(preview.values.tolist())
        for position, value in enumerate(values):
            row = position + 2
            cell = sheet.cell(row, 1, value)
            if isinstance(value, str):
                cell.data_type = "s"
            sheet.cell(row, 2, preview.formula(position))
            addresses.append((sheet.title, f"B{row}"))
    path = tmp_path / "重复检查实际计算.xlsx"
    error_sheet = book.create_sheet("源错误值")
    error_sheet.append(["商品编码", "辅助列"])
    error_values = ["#N/A", "P001", "P001", "#DIV/0!", 0, "0", False, 1, True, None]
    for row, value in enumerate(error_values, start=2):
        source = error_sheet.cell(row, 1, value)
        if isinstance(value, str):
            source.data_type = "e" if value.startswith("#") else "s"
        error_sheet.cell(row, 2, build_duplicate_formula(0, row, len(error_values) + 1))
        addresses.append((error_sheet.title, f"B{row}"))
    expected.extend(["空值", "重复", "重复", "空值", "重复", "正常", "重复", "重复", "重复", "空值"])
    book.save(path)
    book.close()
    actual = _calculate_excel(path, addresses, tmp_path)
    assert actual == expected
