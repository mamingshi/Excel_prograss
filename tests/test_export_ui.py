"""验证 V0.5 按钮状态、导出确认、后台任务及失败后的工作副本保留。"""

from unittest.mock import patch

from openpyxl import load_workbook, Workbook
import pandas as pd
import pytest
from PySide6.QtCore import QPoint, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QMessageBox

from app.core.exporter import NO_CHANGES_MESSAGE
from app.models.export_data import ExportMode
from app.ui.confirmation_dialogs import ExportConfirmationDialog
from app.core.file_loader import file_sha256
from app.ui.main_window import MainWindow
from tests.test_cleaning_ui import window, analyze, generate, wait_for_idle
from tests.test_working_data_ui import apply, undo


def prepare_applied(window):
    analyze(window)
    generate(window)
    apply(window)


def test_export_requires_applied_changes_and_restoring_disables_button(window):
    assert not window.export_button.isEnabled()
    with patch.object(QMessageBox, "information") as notice:
        window._export_results()
    assert notice.call_args.args[2] == NO_CHANGES_MESSAGE
    analyze(window)
    generate(window)
    assert not window.export_button.isEnabled()
    apply(window)
    assert window.export_button.isEnabled()
    with patch("app.ui.main_window.confirm_restore", return_value=True):
        window._restore_original()
    wait_for_idle(window)
    assert not window.export_button.isEnabled()
    undo(window)  # 撤销“恢复原始”，修改重新存在。
    assert window.export_button.isEnabled()


def test_no_file_displays_friendly_notice(qapp):
    window = MainWindow()
    try:
        assert not window.export_button.isEnabled()
        with patch.object(QMessageBox, "information") as notice:
            window._export_results()
        assert notice.call_args.args[2] == "请先打开 Excel 文件。"
    finally:
        window.close()


def test_successful_export_preserves_session_and_runs_in_background(window, tmp_path):
    prepare_applied(window)
    source = window.data_manager.file_path
    target = tmp_path / "清洗结果.xlsx"
    digest = file_sha256(source)
    working = window.data_manager.working_df.copy(deep=True)
    history = list(window.data_manager.undo_stack)
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")) as chooser,
          patch("app.ui.main_window.confirm_export", return_value=ExportMode.UPDATE_COLUMNS) as confirmation,
          patch("app.ui.main_window.show_export_success") as success):
        QTest.mouseClick(window.export_button, Qt.MouseButton.LeftButton)
        assert window._export_worker is not None
        assert window._is_busy()
        assert not window.export_button.isEnabled()
        assert not window.undo_button.isEnabled()
        assert not window.restore_button.isEnabled()
        assert not window.open_button.isEnabled()
        assert not window.sheet_combo.isEnabled()
        wait_for_idle(window)
    assert chooser.call_args.args[2].endswith("清洗预览测试_已清洗.xlsx")
    plan = confirmation.call_args.args[1]
    assert (plan.sheet_count, plan.cell_count) == (1, 6)
    success.assert_called_once()
    assert success.call_args.args[1].target_path == target.resolve()
    assert "导出成功" in window.statusBar().currentMessage()
    assert window.export_button.isEnabled() and window.undo_button.isEnabled()
    assert window.data_manager.undo_stack == history
    assert window.data_manager.is_modified  # 导出不会重置与源文件的差异。
    pd.testing.assert_frame_equal(window.data_manager.working_df, working)
    assert file_sha256(source) == digest
    exported = load_workbook(target)
    try:
        assert exported["问题数据"]["A2"].value == "P001"
        assert exported["正常表"]["A2"].value == "P001"
    finally:
        exported.close()


@pytest.mark.parametrize("cancel_stage", ["file_dialog", "confirmation"])
def test_cancelled_export_writes_nothing(window, tmp_path, cancel_stage):
    prepare_applied(window)
    target = tmp_path / "取消.xlsx"
    before = window.data_manager.working_df.copy(deep=True)
    selected = "" if cancel_stage == "file_dialog" else str(target)
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(selected, "")),
          patch("app.ui.main_window.confirm_export", return_value=None),
          patch("app.ui.main_window.show_export_success") as success):
        window._export_results()
    assert window._export_worker is None
    assert not target.exists()
    success.assert_not_called()
    pd.testing.assert_frame_equal(window.data_manager.working_df, before)


def test_actual_confirmation_shows_all_statistics_and_defaults_to_cancel(window, tmp_path):
    prepare_applied(window)
    target = tmp_path / "确认信息.xlsx"
    def inspect(dialog):
        assert "源文件：清洗预览测试.xlsx" in dialog.summary_label.text()
        assert "修改工作表：1" in dialog.summary_label.text()
        assert "修改单元格：6" in dialog.summary_label.text()
        assert str(target.resolve()) in dialog.summary_label.text()
        assert "原文件不会被修改" in dialog.safety_label.text()
        assert dialog.selected_mode() == ExportMode.ADD_FORMULA_COLUMNS
        assert dialog.cancel_button.isDefault()
        dialog.reject()
        return 0
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")),
          patch.object(ExportConfirmationDialog, "exec", inspect)):
        window._export_results()
    assert window._export_worker is None
    assert not target.exists()


@pytest.mark.parametrize("destination", ["source", "existing"])
def test_invalid_destination_is_reported_before_confirmation(window, tmp_path, destination):
    prepare_applied(window)
    target = window.data_manager.file_path if destination == "source" else tmp_path / "已存在.xlsx"
    if destination == "existing":
        target.write_bytes(b"do not overwrite")
    before = target.read_bytes()
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")),
          patch("app.ui.main_window.confirm_export") as confirmation,
          patch.object(window, "_show_error_dialog") as error):
        window._export_results()
    error.assert_called_once()
    confirmation.assert_not_called()
    assert target.read_bytes() == before
    assert window._export_worker is None


def test_save_failure_keeps_working_data_and_reenables_controls(window, tmp_path):
    prepare_applied(window)
    target = tmp_path / "保存失败.xlsx"
    working = window.data_manager.working_df.copy(deep=True)
    digest = file_sha256(window.data_manager.file_path)
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")),
          patch("app.ui.main_window.confirm_export", return_value=ExportMode.UPDATE_COLUMNS),
          patch.object(Workbook, "save", side_effect=PermissionError("文件占用")),
          patch.object(window, "_show_error_dialog") as error,
          patch("app.ui.main_window.show_export_success") as success):
        window._export_results()
        wait_for_idle(window)
    error.assert_called_once()
    assert "目录权限" in error.call_args.args[1]
    success.assert_not_called()
    assert not target.exists()
    assert window.export_button.isEnabled() and window.undo_button.isEnabled()
    assert window.open_button.isEnabled() and window.sheet_combo.isEnabled()
    pd.testing.assert_frame_equal(window.data_manager.working_df, working)
    assert file_sha256(window.data_manager.file_path) == digest


def test_changed_state_during_confirmation_rejects_old_export(window, tmp_path):
    prepare_applied(window)
    target = tmp_path / "过期快照.xlsx"
    def change_state(*args):
        window.data_manager.undo()
        return ExportMode.UPDATE_COLUMNS
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")),
          patch("app.ui.main_window.confirm_export", side_effect=change_state)):
        window._export_results()
    assert "工作副本已变化" in window.statusBar().currentMessage()
    assert window._export_worker is None
    assert not target.exists()


def test_success_dialog_reports_location_counts_and_folder_option(window, tmp_path):
    from app.models.export_data import ExportResult
    from app.ui.confirmation_dialogs import show_export_success
    target = tmp_path / "结果.xlsx"
    def inspect(dialog):
        assert str(target) in dialog.text()
        assert "修改工作表：2" in dialog.text()
        assert "修改单元格：37" in dialog.text()
        assert {button.text() for button in dialog.buttons()} == {"确定", "打开所在文件夹"}
        assert dialog.defaultButton().text() == "确定"
        return 0
    with patch.object(QMessageBox, "exec", inspect):
        show_export_success(window, ExportResult(target, 2, 37))

def test_mode_selector_defaults_to_formula_and_describes_both_options(window, tmp_path):
    from app.core.exporter import prepare_export
    prepare_applied(window)
    dialog = ExportConfirmationDialog(window, prepare_export(window.data_manager.get_export_snapshot()), tmp_path / "结果.xlsx")
    dialog.show()
    try:
        assert dialog.selected_mode() == ExportMode.ADD_FORMULA_COLUMNS
        assert "共 1 列" in dialog.mode_note_label.text()
        QTest.mouseClick(dialog.update_radio, Qt.MouseButton.LeftButton, pos=QPoint(8, dialog.update_radio.height() // 2))
        assert dialog.selected_mode() == ExportMode.UPDATE_COLUMNS
        assert "列数和顺序保持不变" in dialog.mode_note_label.text()
        QTest.mouseClick(dialog.formula_radio, Qt.MouseButton.LeftButton, pos=QPoint(8, dialog.formula_radio.height() // 2))
        assert dialog.selected_mode() == ExportMode.ADD_FORMULA_COLUMNS
        QTest.keyClick(dialog, Qt.Key.Key_Return)
        assert dialog.result() == ExportConfirmationDialog.DialogCode.Rejected
    finally:
        dialog.close()


@pytest.mark.parametrize("mode", [ExportMode.UPDATE_COLUMNS, ExportMode.ADD_FORMULA_COLUMNS])
def test_real_mode_choice_reaches_worker_and_exports_correct_structure(window, tmp_path, mode):
    prepare_applied(window)
    target = tmp_path / f"{mode.value}.xlsx"
    digest = file_sha256(window.data_manager.file_path)
    working = window.data_manager.working_df.copy(deep=True)
    def choose(dialog):
        if mode == ExportMode.UPDATE_COLUMNS:
            dialog.update_radio.setChecked(True)
        dialog.accept()
        return ExportConfirmationDialog.DialogCode.Accepted
    with (patch("app.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")),
          patch.object(ExportConfirmationDialog, "exec", choose),
          patch("app.ui.main_window.show_export_success") as success):
        window._export_results()
        wait_for_idle(window)
    result = success.call_args.args[1]
    assert result.mode == mode
    assert result.added_column_count == (1 if mode == ExportMode.ADD_FORMULA_COLUMNS else 0)
    book = load_workbook(target)
    try:
        sheet = book["问题数据"]
        if mode == ExportMode.ADD_FORMULA_COLUMNS:
            assert sheet["A2"].value == " P001 "
            assert sheet["B1"].value == "商品编码_清洗后"
            assert sheet["B2"].data_type == "f"
            assert sheet["C2"].value == "备注 0"
        else:
            assert sheet.max_column == 2
            assert sheet["A2"].value == "P001"
    finally:
        book.close()
    assert file_sha256(window.data_manager.file_path) == digest
    pd.testing.assert_frame_equal(window.data_manager.working_df, working)

