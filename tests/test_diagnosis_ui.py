"""验证真实 Excel 导入到后台检测及诊断界面的完整流程。"""

import hashlib
import time
from unittest.mock import patch

from openpyxl import Workbook
import pandas as pd
import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from app.ui.main_window import MainWindow


@pytest.fixture(scope="module")
def application():
    return QApplication.instance() or QApplication([])


def wait_for_idle(window):
    deadline = time.monotonic() + 10
    while window._worker is not None or window._analysis_worker is not None:
        QApplication.processEvents()
        QTest.qWait(10)
        assert time.monotonic() < deadline, "后台任务超时"
    QApplication.processEvents()


@pytest.fixture
def window(tmp_path, application):
    path = tmp_path / "诊断测试.xlsx"
    workbook = Workbook()
    workbook.active.title = "问题数据"
    workbook.active.append(["商品编码", "备注"])
    values = [" P001", "P002\u00a0", "P003\t", None, "DUP", "DUP", " \u00a0X\n ", 0, "0", "ABC 123"]
    for row, value in enumerate(values):
        workbook.active.append([value, f"备注{row}"])
    normal = workbook.create_sheet("正常表")
    normal.append(["商品编码"])
    normal.append(["P001"])
    normal.append(["P002"])
    workbook.create_sheet("仅字段").append(["商品编码"])
    workbook.save(path)
    workbook.close()
    widget = MainWindow()
    widget.show()
    application.processEvents()
    widget._start_load(str(path))
    wait_for_idle(widget)
    yield widget
    wait_for_idle(widget)
    widget.close()
    widget.deleteLater()
    application.processEvents()


def start_analysis(window):
    QTest.mouseClick(window.diagnosis_panel.start_button, Qt.MouseButton.LeftButton)
    assert window._analysis_worker is not None
    assert not window.open_button.isEnabled()
    assert not window.sheet_combo.isEnabled()
    assert not window.diagnosis_panel.start_button.isEnabled()
    wait_for_idle(window)


def test_detection_populates_counts_details_and_keeps_source_unchanged(window):
    panel = window.diagnosis_panel
    assert not panel.start_button.isEnabled()
    assert panel.row_count_label.text() == "数据行数：10"
    before = window._preview.dataframe.copy(deep=True)
    path = window._preview.file_path
    source_digest = hashlib.sha256(path.read_bytes()).digest()
    window._select_column(0)
    assert panel.column_label.text() == "当前列：商品编码"
    assert panel.start_button.isEnabled()
    start_analysis(window)
    assert panel.result.counts == {"前后空格": 2, "CHAR(160)": 2, "Tab / 换行": 2, "空值": 1, "重复值": 2,
                                   "文本数字混合": 1, "前导0风险": 0, "日期异常": 0, "全角字符": 0}
    assert panel.count_labels["CHAR(160)"].text() == "CHAR(160)：2 条"
    assert panel.details_model.rowCount() == 10
    assert panel.details_model.data(panel.details_model.index(0, 0)) == "2"
    assert panel.details_model.data(panel.details_model.index(0, 1)) == " P001"
    assert panel.details_model.headerData(0, Qt.Orientation.Horizontal) == "Excel行号"
    assert not panel.details_model.flags(panel.details_model.index(0, 0)) & Qt.ItemFlag.ItemIsEditable
    assert len([issue for issue in panel.result.issues if issue.excel_row == 8]) == 3
    assert panel.result.duplicate_groups[0].excel_rows == (6, 7)
    assert window.open_button.isEnabled()
    assert panel.start_button.isEnabled()
    pd.testing.assert_frame_equal(window._preview.dataframe, before)
    assert hashlib.sha256(path.read_bytes()).digest() == source_digest


def test_column_and_sheet_changes_clear_old_results(window):
    panel = window.diagnosis_panel
    window._select_column(0)
    start_analysis(window)
    window._select_column(0)
    assert panel.result is not None  # 再次点击同一列时保留结果。
    window._select_column(1)
    assert panel.result is None
    assert panel.details_model.rowCount() == 0
    assert panel.column_label.text() == "当前列：备注"
    start_analysis(window)
    assert not panel.result.issues
    window.sheet_combo.setCurrentText("正常表")
    wait_for_idle(window)
    assert panel.result is None
    assert not panel.start_button.isEnabled()
    assert panel.row_count_label.text() == "数据行数：2"


def test_new_file_clears_diagnostics(window, tmp_path):
    window._select_column(0)
    start_analysis(window)
    path = tmp_path / "第二份.xlsx"
    workbook = Workbook()
    workbook.active.append(["新字段"])
    workbook.active.append(["值"])
    workbook.save(path)
    workbook.close()
    window._start_load(str(path))
    wait_for_idle(window)
    assert window.diagnosis_panel.result is None
    assert window.diagnosis_panel.details_model.rowCount() == 0
    assert not window.diagnosis_panel.start_button.isEnabled()


def test_failed_sheet_switch_keeps_previous_diagnostics(window):
    window._select_column(0)
    start_analysis(window)
    result = window.diagnosis_panel.result
    window._preview.file_path.unlink()
    with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok):
        window.sheet_combo.setCurrentText("正常表")
        wait_for_idle(window)
    assert window.sheet_combo.currentText() == "问题数据"
    assert window.diagnosis_panel.result is result
    assert window.diagnosis_panel.start_button.isEnabled()


def test_header_only_sheet_can_be_analyzed(window):
    window.sheet_combo.setCurrentText("仅字段")
    wait_for_idle(window)
    window._select_column(0)
    start_analysis(window)
    assert window.diagnosis_panel.result.row_count == 0
    assert all(count == 0 for count in window.diagnosis_panel.result.counts.values())
    assert window.diagnosis_panel.details_model.rowCount() == 0


def test_detection_error_restores_controls_and_keeps_data(window):
    window._select_column(0)
    before = window._preview.dataframe.copy(deep=True)
    with patch("app.ui.analysis_worker.analyze_column", side_effect=RuntimeError("测试异常")):
        with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok) as dialog:
            start_analysis(window)
    dialog.assert_called_once()
    assert window.diagnosis_panel.start_button.isEnabled()
    assert window.table.isEnabled()
    pd.testing.assert_frame_equal(window._preview.dataframe, before)


def test_close_during_analysis_waits_for_background_task(window):
    window._select_column(0)
    window._start_analysis()
    window.close()
    assert window._close_requested
    wait_for_idle(window)
    assert not window.isVisible()
