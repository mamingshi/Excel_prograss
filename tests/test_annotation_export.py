"""真实 xlsx 验证辅助列末尾追加、两种导出模式、引用、结构与源文件保护。"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table
import pytest

from app.core.annotation_builder import build_duplicate_formula, generate_duplicate_preview
from app.core.cleaner import CleaningRules, generate_cleaning_preview
from app.core.data_manager import WorkbookDataManager
from app.core.exporter import ExcelExportError, export_workbook, prepare_export
from app.core.file_loader import ExcelFileLoader, file_sha256
from app.models.export_data import ExportMode
from tests.test_export_formulas import _calculate_excel


@pytest.fixture
def session(tmp_path):
    source = tmp_path / "订单原始.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["订单编号", "商品名称", "数量", "商品编码"])
    values = ["P001", "P002", "P001", None, 0, "0", " p003", "p003"]
    for position, value in enumerate(values):
        sheet.append([f"O{position + 1}", f" 商品{position + 1}", position + 1, value])
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:D9"
    sheet["D1"].fill = PatternFill("solid", fgColor="DDF0E5")
    sheet.column_dimensions["D"].width = 22
    other = book.create_sheet("统计")
    other.append(["合计"])
    other.append(["=SUM('订单'!C2:C9)"])
    book.save(source)
    book.close()
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(source))
    return source, manager


def annotate(manager):
    result = generate_duplicate_preview(manager.working_df, 3)
    return manager.apply_annotation(manager.prepare_annotation_preview(result))


def clean(manager, column):
    result = generate_cleaning_preview(manager.working_df.iloc[:, column], CleaningRules(outer_spaces=True))
    return manager.apply_preview(manager.prepare_preview(column, result))


@pytest.mark.parametrize("mode", list(ExportMode))
def test_annotation_only_export_appends_formula_and_preserves_original_cells(session, tmp_path, mode):
    source, manager = session
    digest = file_sha256(source)
    annotate(manager)
    target = tmp_path / f"辅助列_{mode}.xlsx"
    result = export_workbook(prepare_export(manager.get_export_snapshot(), mode), target)
    assert result.sheet_count == 1 and result.cell_count == 0
    assert result.annotation_column_count == 1
    assert result.added_column_count == 0
    book = load_workbook(target, data_only=False)
    sheet = book["订单"]
    assert sheet.max_row == 9 and sheet.max_column == 5
    assert sheet["E1"].value == "商品编码_重复检查"
    assert sheet["E1"].fill.patternType == "solid"
    assert sheet["E1"].fill.fgColor.rgb == "00DDF0E5"
    assert sheet["E2"].value == build_duplicate_formula(3, 2, 9)
    assert sheet["E9"].value == build_duplicate_formula(3, 9, 9)
    assert all(sheet.cell(row, 5).data_type == "f" for row in range(2, 10))
    assert [sheet.cell(row, 4).value for row in range(2, 10)] == ["P001", "P002", "P001", None, 0, "0", " p003", "p003"]
    assert sheet["D6"].data_type == "n" and sheet["D7"].data_type == "s"
    assert sheet.freeze_panes == "B2" and sheet.auto_filter.ref == "A1:D9"
    assert book["统计"]["A2"].value == "=SUM('订单'!C2:C9)"
    assert book.calculation.fullCalcOnLoad
    book.close()
    assert file_sha256(source) == digest
    assert len(manager.undo_stack) == 1


@pytest.mark.parametrize("mode", list(ExportMode))
def test_mixed_cleaning_and_annotation_reference_final_working_values(session, tmp_path, mode):
    source, manager = session
    annotate(manager)
    clean(manager, 1)
    clean(manager, 3)
    target = tmp_path / f"混合_{mode}.xlsx"
    plan = prepare_export(manager.get_export_snapshot(), mode)
    assert plan.column_count == 2 and plan.annotation_column_count == 1
    export_workbook(plan, target)
    book = load_workbook(target)
    sheet = book["订单"]
    if mode == ExportMode.UPDATE_COLUMNS:
        assert sheet["E1"].value == "商品编码_重复检查"
        assert sheet["E2"].value == build_duplicate_formula(3, 2, 9)
        assert sheet["D8"].value == "p003"
    else:
        assert sheet["C1"].value == "商品名称_清洗后"
        assert sheet["F1"].value == "商品编码_清洗后"
        assert sheet["G1"].value == "商品编码_重复检查"
        assert sheet["G2"].value == build_duplicate_formula(5, 2, 9)
        assert sheet["G9"].value == build_duplicate_formula(5, 9, 9)
        assert sheet["E8"].value == " p003"
        assert book["统计"]["A2"].value == "=SUM('订单'!D2:D9)"
    book.close()


@pytest.mark.parametrize("mode", list(ExportMode))
def test_real_excel_export_labels_match_final_working_copy(session, tmp_path, mode):
    source, manager = session
    digest = file_sha256(source)
    annotate(manager)
    clean(manager, 1)
    clean(manager, 3)
    target = tmp_path / f"真实计算_{mode}.xlsx"
    export_workbook(prepare_export(manager.get_export_snapshot(), mode), target)
    letter = "E" if mode == ExportMode.UPDATE_COLUMNS else "G"
    addresses = [("订单", f"{letter}{row}") for row in range(2, 10)]
    actual = _calculate_excel(target, addresses, tmp_path)
    assert actual == manager.working_df.iloc[:, -1].tolist()
    assert actual == ["重复", "正常", "重复", "空值", "正常", "正常", "重复", "重复"]
    assert file_sha256(source) == digest


def test_multiple_helpers_and_name_collision_are_exported_in_order(session, tmp_path):
    source, manager = session
    annotate(manager)
    annotate(manager)
    target = tmp_path / "两个辅助列.xlsx"
    result = export_workbook(prepare_export(manager.get_export_snapshot()), target)
    book = load_workbook(target)
    assert book["订单"]["E1"].value == "商品编码_重复检查"
    assert book["订单"]["F1"].value == "商品编码_重复检查_2"
    assert book["订单"]["F3"].value == build_duplicate_formula(3, 3, 9)
    assert result.annotation_column_count == 2
    book.close()


def test_undo_and_restore_do_not_export_removed_helpers(session, tmp_path):
    source, manager = session
    annotate(manager)
    manager.undo()
    with pytest.raises(ExcelExportError):
        prepare_export(manager.get_export_snapshot())
    annotate(manager)
    manager.restore_original()
    with pytest.raises(ExcelExportError):
        prepare_export(manager.get_export_snapshot())


def test_table_and_trailing_formatted_columns_are_not_shifted(session, tmp_path):
    source, manager = session
    book = load_workbook(source)
    sheet = book["订单"]
    sheet.add_table(Table(displayName="Orders", ref="A1:D9"))
    sheet["Z2"].fill = PatternFill("solid", fgColor="FFFF00")
    book.save(source)
    book.close()
    manager.open_workbook(ExcelFileLoader().load(source))
    annotate(manager)
    target = tmp_path / "保留结构.xlsx"
    export_workbook(prepare_export(manager.get_export_snapshot(), ExportMode.ADD_FORMULA_COLUMNS), target)
    book = load_workbook(target)
    assert book["订单"]["E1"].value == "商品编码_重复检查"
    assert book["订单"].tables["Orders"].ref == "A1:D9"
    assert book["订单"]["Z2"].fill.fgColor.rgb == "00FFFF00"
    book.close()


def test_merged_target_is_rejected_without_partial_output(session, tmp_path):
    source, manager = session
    book = load_workbook(source)
    book["订单"].merge_cells("E1:F1")
    book.save(source)
    book.close()
    manager.open_workbook(ExcelFileLoader().load(source))
    digest = file_sha256(source)
    annotate(manager)
    target = tmp_path / "不能覆盖合并.xlsx"
    with pytest.raises(ExcelExportError, match="合并区域"):
        export_workbook(prepare_export(manager.get_export_snapshot()), target)
    assert not target.exists()
    assert not list(tmp_path.glob(".excel-cleaner-*"))
    assert file_sha256(source) == digest


def test_unregistered_added_column_and_tampered_labels_are_rejected(session):
    _, manager = session
    manager.working_df["未知列"] = "未知"
    with pytest.raises(ExcelExportError, match="结构"):
        prepare_export(manager.get_export_snapshot())
    manager.open_workbook(ExcelFileLoader().load(manager.file_path))
    annotate(manager)
    snapshot = manager.get_export_snapshot()
    snapshot.sheets[0].working_df.iloc[0, -1] = "正常"
    with pytest.raises(ExcelExportError, match="标签"):
        prepare_export(snapshot)


def test_existing_target_is_never_overwritten(session, tmp_path):
    _, manager = session
    annotate(manager)
    target = tmp_path / "已有文件.xlsx"
    target.write_bytes(b"existing")
    with pytest.raises(ExcelExportError, match="已经存在"):
        export_workbook(prepare_export(manager.get_export_snapshot()), target)
    assert target.read_bytes() == b"existing"


def test_actual_aa_column_and_duplicate_indexes(tmp_path):
    source = tmp_path / "宽表.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append([f"字段{i}" for i in range(26)] + ["商品编码"])
    for index, code in enumerate(["P001", "P002", "P001"]):
        sheet.append([index] * 26 + [code])
    book.save(source)
    book.close()
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(source))
    # 行号只按位置，非默认 index 不改变 Excel 公式行。
    manager.current_data.original_df.index = [88, 88, -9]
    manager.current_data.working_df.index = [88, 88, -9]
    preview = generate_duplicate_preview(manager.working_df, 26)
    manager.apply_annotation(manager.prepare_annotation_preview(preview))
    target = tmp_path / "宽表输出.xlsx"
    export_workbook(prepare_export(manager.get_export_snapshot()), target)
    book = load_workbook(target)
    assert book.active["AB2"].value == build_duplicate_formula(26, 2, 4)
    assert book.active["AB4"].value == build_duplicate_formula(26, 4, 4)
    book.close()


def test_multiple_sheets_with_only_annotations_count_as_export_changes(session, tmp_path):
    source, manager = session
    annotate(manager)
    manager.add_sheet(ExcelFileLoader().load(source, "统计"))
    preview = generate_duplicate_preview(manager.working_df, 0)
    manager.apply_annotation(manager.prepare_annotation_preview(preview))
    plan = prepare_export(manager.get_export_snapshot())
    assert plan.sheet_count == 2 and plan.annotation_column_count == 2
    target = tmp_path / "多表辅助列.xlsx"
    export_workbook(plan, target)
    book = load_workbook(target)
    assert book["订单"]["E1"].value == "商品编码_重复检查"
    assert book["统计"]["B1"].value == "合计_重复检查"
    assert book["统计"]["A2"].value == "=SUM('订单'!C2:C9)"
    book.close()
