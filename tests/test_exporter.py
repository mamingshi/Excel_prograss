"""使用真实临时工作簿验证差异导出、元数据保留和事务安全。"""

from dataclasses import replace
from copy import copy
import os
from unittest.mock import patch
import warnings
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
import pytest

from app.core.cleaner import CleaningRules, generate_cleaning_preview
from app.core.data_manager import WorkbookDataManager, DataStateError
from app.core.exporter import (
    ExcelExportError, NO_CHANGES_MESSAGE, dataframe_column_to_excel_column,
    dataframe_index_to_excel_row, export_workbook, get_changes, prepare_export, suggest_export_path,
)
from app.core.file_loader import ExcelFileLoader, file_sha256
from app.models.export_data import ExportSnapshot, SheetExportSnapshot


RULES = CleaningRules(nbsp=True, tab=True, newline=True, carriage_return=True, outer_spaces=True)


def apply_column(manager, column=0):
    preview = generate_cleaning_preview(manager.working_df.iloc[:, column], RULES)
    return manager.apply_preview(manager.prepare_preview(column, preview))


@pytest.fixture
def session(tmp_path):
    source = tmp_path / "客户订单.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["商品编码", "数量", "说明", "公式"])
    for values in [
        [" P001 ", 0, "原值保持", "=SUM(B2:B4)"],
        ["P002\u00a0", 1002, "正常", "=B3*2"],
        ["ABC\t123", 12.5, "正常", None],
        [None, 0, "空编码", None],
        ["0", 1, "字符串零", None],
        [" =1+1 ", 2, "应保持文本", None],
        ["DUP", 3, "重复保留", None],
        ["DUP", 4, "重复保留", None],
        ["正常", " B010 ", "另一列可清洗", None],
        [" ", None, "清洗为空", None],
    ]:
        sheet.append(values)
    sheet["A2"].font = Font(name="微软雅黑", bold=True, color="008080")
    sheet["A2"].fill = PatternFill("solid", fgColor="CCFFCC")
    sheet["A2"].border = Border(left=Side(style="thin", color="FF0000"))
    sheet["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    sheet["A2"].number_format = "@"
    sheet["B4"].number_format = "0.00"
    sheet.column_dimensions["A"].width = 24
    sheet.row_dimensions[2].height = 28
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:D11"
    sheet.print_area = "A1:D11"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_margins.left = 0.3
    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="2000")
    validation.add("B2:B11")
    sheet.add_data_validation(validation)
    second = book.create_sheet("客户")
    second.append(["客户编码", "备注"])
    second.append([" X001 ", "保持"])
    second.append(["X002", "保持"])
    hidden = book.create_sheet("说明")
    hidden.append(["不要重建此表", "=1+1"])
    hidden["E7"] = "合并区域"
    hidden.merge_cells("E7:F7")
    hidden.sheet_state = "hidden"
    hidden["A1"].font = Font(bold=True)
    book.save(source)
    book.close()
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(source))
    return source, manager


def target_and_plan(session):
    source, manager = session
    apply_column(manager)
    return source.with_name("结果.xlsx"), prepare_export(manager.get_export_snapshot())


def test_changed_values_numeric_types_headers_and_source_hash(session):
    source, manager = session
    digest = file_sha256(source)
    original = manager.original_df.copy(deep=True)
    target, plan = target_and_plan(session)
    result = export_workbook(plan, target)
    assert (result.sheet_count, result.cell_count) == (1, 5)
    assert file_sha256(source) == digest
    pd.testing.assert_frame_equal(manager.original_df, original)
    book = load_workbook(target, data_only=False)
    try:
        sheet = book["订单"]
        assert [sheet.cell(1, column).value for column in range(1, 5)] == original.columns.tolist()
        assert [sheet[f"A{row}"].value for row in (2, 3, 4)] == ["P001", "P002", "ABC 123"]
        assert sheet["A5"].value is None
        assert sheet["A6"].value == "0" and sheet["A6"].data_type == "s"
        assert sheet["A7"].value == "=1+1" and sheet["A7"].data_type == "s"
        assert sheet["A11"].value is None
        assert [sheet[f"B{row}"].value for row in (2, 3, 4)] == [0, 1002, 12.5]
        assert [type(sheet[f"B{row}"].value) for row in (2, 3, 4)] == [int, int, float]
        assert all(sheet[f"B{row}"].data_type == "n" for row in (2, 3, 4))
        assert sheet["B11"].value is None
        assert sheet["B10"].value == " B010 "
        assert sheet["A8"].value == sheet["A9"].value == "DUP"
        assert sheet["D2"].value == "=SUM(B2:B4)" and sheet["D2"].data_type == "f"
        assert sheet["D3"].value == "=B3*2"
    finally:
        book.close()


def test_styles_metadata_and_untouched_sheets_preserved(session):
    source, manager = session
    target, plan = target_and_plan(session)
    export_workbook(plan, target)
    before, after = load_workbook(source), load_workbook(target)
    try:
        assert after.sheetnames == before.sheetnames
        first, saved = before["订单"], after["订单"]
        for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
            assert copy(getattr(first["A2"], attribute)) == copy(getattr(saved["A2"], attribute))
        assert saved["B4"].number_format == "0.00"
        assert saved.column_dimensions["A"].width == 24
        assert saved.row_dimensions[2].height == 28
        assert saved.freeze_panes == "B2"
        assert saved.auto_filter.ref == "A1:D11"
        assert str(saved.print_area) == str(first.print_area)
        assert saved.print_title_rows == first.print_title_rows
        assert saved.page_setup == first.page_setup
        assert saved.page_margins == first.page_margins
        assert saved.data_validations == first.data_validations
        for name in ("客户", "说明"):
            original, copied = before[name], after[name]
            assert copied.sheet_state == original.sheet_state
            assert str(copied.merged_cells) == str(original.merged_cells)
            assert list(copied.values) == list(original.values)
            for row in original:
                for cell in row:
                    assert copied[cell.coordinate]._style == cell._style
    finally:
        before.close()
        after.close()


def test_all_changed_sheets_export_even_when_viewing_unmodified_sheet(session):
    source, manager = session
    apply_column(manager)
    manager.add_sheet(ExcelFileLoader().load(source, "客户"))
    apply_column(manager)
    manager.add_sheet(ExcelFileLoader().load(source, "说明"))
    target = source.with_name("多表.xlsx")
    result = export_workbook(prepare_export(manager.get_export_snapshot()), target)
    assert (result.sheet_count, result.cell_count) == (2, 6)
    book = load_workbook(target)
    try:
        assert book["订单"]["A2"].value == "P001"
        assert book["客户"]["A2"].value == "X001"
        assert book["说明"]["A1"].value == "不要重建此表"
    finally:
        book.close()


def test_undo_exports_only_current_differences(session):
    source, manager = session
    apply_column(manager, 0)
    apply_column(manager, 1)
    manager.undo()
    # 即使日志中出现过 B 列清洗，已经撤销的值不应导出。
    plan = prepare_export(manager.get_export_snapshot())
    assert all(change.excel_column == 1 for change in plan.changes)
    target = source.with_name("撤销后.xlsx")
    export_workbook(plan, target)
    book = load_workbook(target)
    try:
        assert book["订单"]["B10"].value == " B010 "
    finally:
        book.close()


@pytest.mark.parametrize("operation", ["undo", "restore_original"])
def test_no_changes_after_undo_or_restore_generates_no_file(session, operation):
    source, manager = session
    apply_column(manager)
    getattr(manager, operation)()
    assert not get_changes(manager.get_export_snapshot())
    with pytest.raises(ExcelExportError, match=NO_CHANGES_MESSAGE):
        prepare_export(manager.get_export_snapshot())
    assert list(source.parent.glob("*.xlsx")) == [source]


def test_positional_mapping_non_default_index_duplicate_and_empty_headers(tmp_path):
    source = tmp_path / "重名字段.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["字段", "字段", None])
    sheet.append(["保持", " A2 ", " C2 "])
    sheet.append(["保持", " A3 ", " C3 "])
    book.save(source)
    book.close()
    preview = ExcelFileLoader().load(source)
    preview.dataframe.index = [100, 7]
    manager = WorkbookDataManager()
    manager.open_workbook(preview)
    apply_column(manager, 1)
    apply_column(manager, 2)
    plan = prepare_export(manager.get_export_snapshot())
    assert [(change.excel_row, change.excel_column) for change in plan.changes] == [(2, 2), (3, 2), (2, 3), (3, 3)]
    assert plan.changes[0].column_name == "字段.1"
    target = source.with_name("结果.xlsx")
    export_workbook(plan, target)
    saved = load_workbook(target)
    try:
        assert saved.active["A1"].value == saved.active["B1"].value == "字段"
        assert saved.active["C1"].value is None
        assert saved.active["B2"].value == "A2"
        assert saved.active["C3"].value == "C3"
    finally:
        saved.close()


def test_nan_none_and_normal_values_do_not_create_false_differences(tmp_path):
    original = pd.DataFrame({"A": [None, float("nan"), pd.NA, "", "0", 0, 1002, 12.5, " P001 "]}, dtype=object)
    working = original.copy(deep=True)
    working.iat[8, 0] = "P001"
    snapshot = ExportSnapshot(tmp_path / "原始.xlsx", None, (SheetExportSnapshot("S", original, working),), ())
    changes = get_changes(snapshot)
    assert len(changes) == 1
    assert changes[0].excel_row == 10
    pd.testing.assert_frame_equal(original.iloc[:8], working.iloc[:8])
    assert dataframe_index_to_excel_row(0) == 2
    assert dataframe_column_to_excel_column(0) == 1
    for function in (dataframe_index_to_excel_row, dataframe_column_to_excel_column):
        with pytest.raises(ExcelExportError):
            function(-1)


def test_cached_formula_target_is_rejected_without_any_output(tmp_path):
    source = tmp_path / "公式.xlsx"
    book = Workbook()
    book.active.append(["编码", "备注"])
    book.active.append(['=" P001 "', "正常"])
    book.save(source)
    book.close()
    # 构造 Excel 已缓存的文本公式结果，模拟读取结果可被误清洗的场景。
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    xml = ElementTree.fromstring(parts["xl/worksheets/sheet1.xml"])
    cell = xml.find(f".//{namespace}c[@r='A2']")
    cell.set("t", "str")
    cell.find(namespace + "v").text = " P001 "
    parts["xl/worksheets/sheet1.xml"] = ElementTree.tostring(xml)
    with ZipFile(source, "w") as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(source))
    assert manager.working_df.iat[0, 0] == " P001 "
    apply_column(manager)
    digest = file_sha256(source)
    target = source.with_name("结果.xlsx")
    with pytest.raises(ExcelExportError, match="该单元格包含 Excel 公式"):
        export_workbook(prepare_export(manager.get_export_snapshot()), target)
    assert file_sha256(source) == digest
    assert not target.exists()
    assert not list(source.parent.glob(".excel-cleaner-*"))


@pytest.mark.parametrize("destination", ["source", "existing", "missing_parent", "xlsm"])
def test_unsafe_destinations_rejected_and_existing_files_unchanged(session, destination):
    source, manager = session
    target, plan = target_and_plan(session)
    digest = file_sha256(source)
    if destination == "source":
        target = source
    elif destination == "existing":
        target.write_bytes(b"existing file must remain")
    elif destination == "missing_parent":
        target = source.parent / "不存在" / "结果.xlsx"
    else:
        target = target.with_suffix(".xlsm")
    before = target.read_bytes() if target.is_file() else None
    with pytest.raises(ExcelExportError):
        export_workbook(plan, target)
    if before is not None:
        assert target.read_bytes() == before
    assert file_sha256(source) == digest
    assert not list(source.parent.glob(".excel-cleaner-*"))


def test_hard_link_alias_to_source_is_rejected(session):
    source, manager = session
    target, plan = target_and_plan(session)
    os.link(source, target)
    digest = file_sha256(source)
    with pytest.raises(ExcelExportError, match="不能覆盖原"):
        export_workbook(plan, target)
    assert file_sha256(source) == digest


def test_default_filename_avoids_existing_targets(session):
    source, _ = session
    assert suggest_export_path(source).name == "客户订单_已清洗.xlsx"
    source.with_name("客户订单_已清洗.xlsx").touch()
    source.with_name("客户订单_已清洗_1.xlsx").touch()
    assert suggest_export_path(source).name == "客户订单_已清洗_2.xlsx"


@pytest.mark.parametrize("stage", ["copy", "load", "save", "permission", "publish"])
def test_failures_leave_source_unchanged_and_remove_temporary_file(session, stage):
    source, manager = session
    target, plan = target_and_plan(session)
    digest = file_sha256(source)
    if stage == "copy":
        mock_target, error = "app.core.exporter.shutil.copyfile", OSError("复制失败")
    elif stage == "load":
        mock_target, error = "app.core.exporter.load_workbook", ValueError("加载失败")
    elif stage == "save":
        mock_target, error = "openpyxl.workbook.workbook.Workbook.save", OSError("保存失败")
    elif stage == "permission":
        mock_target, error = "app.core.exporter.tempfile.mkstemp", PermissionError("无写权限/占用")
    else:
        mock_target, error = "app.core.exporter._publish_new_file", PermissionError("文件占用")
    with patch(mock_target, side_effect=error):
        with pytest.raises(ExcelExportError):
            export_workbook(plan, target)
    assert not target.exists()
    assert file_sha256(source) == digest
    assert not list(source.parent.glob(".excel-cleaner-*"))


def test_destination_created_during_save_is_not_overwritten(session):
    source, manager = session
    target, plan = target_and_plan(session)
    real_save = Workbook.save
    def race_save(book, filename):
        real_save(book, filename)
        target.write_bytes(b"other application created this")
    with patch.object(Workbook, "save", race_save):
        with pytest.raises(ExcelExportError, match="已经存在"):
            export_workbook(plan, target)
    assert target.read_bytes() == b"other application created this"
    assert not list(source.parent.glob(".excel-cleaner-*"))


def test_missing_source_and_changed_source_are_rejected(session):
    source, manager = session
    target, plan = target_and_plan(session)
    original_bytes = source.read_bytes()
    source.unlink()
    with pytest.raises(ExcelExportError, match="源文件不存在"):
        export_workbook(plan, target)
    source.write_bytes(original_bytes)
    book = load_workbook(source)
    book["订单"]["C2"] = "外部编辑"
    book.save(source)
    book.close()
    digest = file_sha256(source)
    with pytest.raises(ExcelExportError, match="发生变化"):
        export_workbook(plan, target)
    assert file_sha256(source) == digest
    assert not target.exists()
    with pytest.raises(DataStateError, match="发生变化"):
        manager.add_sheet(ExcelFileLoader().load(source, "客户"))


@pytest.mark.parametrize("mismatch", ["working_columns", "working_rows", "source_columns", "sheet_missing"])
def test_structure_mismatches_rejected(session, mismatch):
    source, manager = session
    apply_column(manager)
    snapshot = manager.get_export_snapshot()
    first = snapshot.sheets[0]
    if mismatch == "working_columns":
        first.working_df.columns = ["错误字段", "数量", "说明", "公式"]
    elif mismatch == "working_rows":
        snapshot = replace(snapshot, sheets=(replace(first, working_df=first.working_df.iloc[:-1]),))
    elif mismatch == "source_columns":
        first.original_df.columns = first.working_df.columns = ["错误字段", "数量", "说明", "公式"]
    else:
        snapshot = replace(snapshot, sheets=(replace(first, sheet_name="不存在"),))
    digest = file_sha256(source)
    target = source.with_name("结构异常.xlsx")
    with pytest.raises(ExcelExportError):
        export_workbook(prepare_export(snapshot), target)
    assert file_sha256(source) == digest
    assert not target.exists()
    assert not list(source.parent.glob(".excel-cleaner-*"))


def test_unsupported_extension_warning_cancels_export(session):
    source, manager = session
    target, plan = target_and_plan(session)
    real_load = load_workbook
    def load_with_warning(*args, **kwargs):
        warnings.warn("Unknown extension is not supported and will be removed", UserWarning)
        return real_load(*args, **kwargs)
    with patch("app.core.exporter.load_workbook", side_effect=load_with_warning):
        with pytest.raises(ExcelExportError, match="无法保留"):
            export_workbook(plan, target)
    assert not target.exists()
    assert not list(source.parent.glob(".excel-cleaner-*"))


def test_none_nan_snapshots_export_without_writing_missing_values(session):
    source, manager = session
    preview = ExcelFileLoader().load(source)
    preview.dataframe.iat[3, 0] = None
    preview.dataframe.iat[9, 1] = float("nan")
    manager.open_workbook(preview)
    target, plan = target_and_plan(session)
    export_workbook(plan, target)
    saved = load_workbook(target)
    try:
        assert saved["订单"]["A5"].value is None
        assert saved["订单"]["B11"].value is None
    finally:
        saved.close()


def test_unmodified_carriage_returns_and_remaining_cr_rules_are_preserved(session):
    source, manager = session
    book = load_workbook(source)
    book["订单"]["C2"] = "普通文本\r\n下一行"
    book["订单"]["A3"] = " P002\rX "
    book["说明"]["C1"] = "隐藏表\r回车"
    book.save(source)
    book.close()
    # 使用合法字符引用构造含 CR 的源数据；不依赖测试环境是否安装 lxml。
    with ZipFile(source) as archive:
        parts = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with ZipFile(source, "w") as archive:
        for info, content in parts:
            namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
            values = {}
            if info.filename == "xl/worksheets/sheet1.xml":
                values = {"C2": "普通文本\r\n下一行", "A3": " P002\rX ", "C3": "只含\n换行"}
            elif info.filename == "xl/worksheets/sheet3.xml":
                values = {"C1": "隐藏表\r回车"}
            if values:
                xml = ElementTree.fromstring(content)
                for coordinate, value in values.items():
                    cell = xml.find(f".//{namespace}c[@r='{coordinate}']")
                    cell.find(f"{namespace}is/{namespace}t").text = value
                content = ElementTree.tostring(xml)
            if info.filename.endswith(".xml"):
                content = content.replace(b"\r", b"&#13;")
            archive.writestr(info, content)
    manager.open_workbook(ExcelFileLoader().load(source))
    rules = CleaningRules(outer_spaces=True)
    preview = generate_cleaning_preview(manager.working_df.iloc[:, 0], rules)
    manager.apply_preview(manager.prepare_preview(0, preview))
    digest = file_sha256(source)
    target = source.with_name("保留回车.xlsx")
    export_workbook(prepare_export(manager.get_export_snapshot()), target)
    saved = load_workbook(target)
    try:
        assert saved["订单"]["C2"].value == "普通文本\r\n下一行"
        assert saved["订单"]["C3"].value == "只含\n换行"
        assert saved["订单"]["A3"].value == "P002\rX"
        assert saved["说明"]["C1"].value == "隐藏表\r回车"
        assert file_sha256(source) == digest
    finally:
        saved.close()
