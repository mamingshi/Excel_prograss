"""导出公式的构造限制与可选真实 Excel 计算验证（不增加运行依赖）。"""

import json
import os
from pathlib import Path
import subprocess

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from app.core.cleaner import CleaningRules, apply_cleaning_rules
from app.core.export_formula_builder import build_export_formula
from app.core.exporter import export_workbook, prepare_export
from app.core.file_loader import file_sha256
from app.models.export_data import ExportMode
from tests.test_exporter import session, apply_column


@pytest.mark.parametrize("reference", ["A0", "XFE2", "B2)", "Sheet1!A2"])
def test_export_formula_rejects_invalid_reference(reference):
    with pytest.raises(ValueError):
        build_export_formula((CleaningRules(outer_spaces=True),), reference)


def test_formula_preserves_types_and_uses_trim_only_to_find_boundaries():
    formula = build_export_formula((CleaningRules(outer_spaces=True),), "B7")
    assert formula.startswith('=IF(ISBLANK(B7),"",IF(ISTEXT(B7),')
    assert "MID(B7," in formula
    assert "FIND(CHAR(1),SUBSTITUTE(" in formula
    assert not formula.endswith("TRIM(B7),B7))")


def test_repeated_rules_are_compacted_without_changing_first_strategy():
    rules = (CleaningRules(tab=True, outer_spaces=True),
             CleaningRules(tab=True, outer_spaces=True, hidden_strategy="delete"))
    assert build_export_formula(rules) == build_export_formula((rules[0],))


def _calculate_excel(path: Path, addresses: list[tuple[str, str]], directory: Path):
    """用独立隐藏实例只读打开测试副本，关闭时不保存，不接触用户已打开的 Excel。"""
    if os.name != "nt":
        pytest.skip("真实 Excel 计算仅在 Windows 安装 Excel 时运行")
    import winreg
    try:
        key = winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"Excel.Application")
        key.Close()
    except FileNotFoundError:
        pytest.skip("未安装 Excel；保留公式结构及真实工作簿导出测试")
    config = directory / "excel-config.json"
    result = directory / "excel-result.json"
    script = directory / "calculate-test-formulas.ps1"
    config.write_text(json.dumps({"path": str(path), "cells": addresses}, ensure_ascii=False), encoding="utf-8")
    script.write_text(r'''
param([string]$ConfigPath, [string]$ResultPath)
$ErrorActionPreference = "Stop"
$config = Get-Content -LiteralPath $ConfigPath -Encoding UTF8 -Raw | ConvertFrom-Json
$excelTest = $null
$bookTest = $null
try {
    $excelTest = New-Object -ComObject Excel.Application
    $excelTest.Visible = $false
    $excelTest.DisplayAlerts = $false
    $excelTest.AskToUpdateLinks = $false
    $bookTest = $excelTest.Workbooks.Open($config.path, 0, $true)
    $excelTest.CalculateFullRebuild()
    $values = @()
    foreach ($cell in $config.cells) {
        $values += ,($bookTest.Worksheets.Item($cell[0]).Range($cell[1]).Value2)
    }
    $text = ConvertTo-Json -InputObject $values -Depth 5 -Compress
    [System.IO.File]::WriteAllText($ResultPath, $text, (New-Object System.Text.UTF8Encoding($false)))
} finally {
    if ($null -ne $bookTest) {
        $bookTest.Close($false)
        [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($bookTest)
    }
    if ($null -ne $excelTest) {
        $excelTest.Quit()
        [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($excelTest)
    }
    [GC]::Collect()
    [GC]::WaitForPendingFinalizers()
}
''', encoding="utf-8-sig")
    subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script),
         "-ConfigPath", str(config), "-ResultPath", str(result)],
        check=True, timeout=60, capture_output=True,
        creationflags=subprocess.CREATE_NO_WINDOW,
    )
    return json.loads(result.read_text(encoding="utf-8"))


def test_real_excel_calculation_matches_python_cleaning(tmp_path):
    values = [" P001 ", "P002\u00a0", "ABC\t123", " ABC  123 ", "", " ", "   ", "0", 0,
              1002, 12.5, None, "=1+1", "一\u00a0个", "A\r\nB", " \tP003\u00a0\n ",
              "DUP", " 😀  中间 😀 ", " " + "a" * 30000 + " "]
    all_rules = dict(nbsp=True, tab=True, newline=True, carriage_return=True, outer_spaces=True)
    sequences = [
        (CleaningRules(**all_rules),),
        (CleaningRules(**all_rules, hidden_strategy="delete"),),
        (CleaningRules(outer_spaces=True),),
        (CleaningRules(nbsp=True, tab=True),),
        (CleaningRules(outer_spaces=True), CleaningRules(nbsp=True, tab=True, hidden_strategy="delete")),
    ]
    book = Workbook()
    sheet = book.active
    sheet.title = "公式计算"
    sheet.append(["原始", "结果"])
    expected, addresses = [], []
    for sequence in sequences:
        for value in values:
            row = len(expected) + 2
            original = sheet.cell(row, 1)
            original.value = value
            if isinstance(value, str):
                original.data_type = "s"
            sheet.cell(row, 2).value = build_export_formula(sequence, f"A{row}")
            result = value
            for rules in sequence:
                result = apply_cleaning_rules(result, rules)
            expected.append("" if result is None else result)
            addresses.append((sheet.title, f"B{row}"))
    path = tmp_path / "公式真实计算.xlsx"
    book.save(path)
    book.close()
    # 测试输入中的 CR 使用合法 XML 字符引用，不让写入器改变测试样本。
    from zipfile import ZipFile
    from xml.etree import ElementTree
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(path, "w") as archive:
        for info, data in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = ElementTree.fromstring(data)
                for position in range(len(expected)):
                    value = values[position % len(values)]
                    if isinstance(value, str) and value:
                        cell = xml.find(f".//{namespace}c[@r='A{position + 2}']")
                        text = cell.find(f"{namespace}is/{namespace}t")
                        text.text = value
                        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                data = ElementTree.tostring(xml).replace(b"\r", b"&#13;")
            archive.writestr(info, data)
    actual = _calculate_excel(path, addresses, tmp_path)
    assert actual == expected
    for value, calculated in zip(expected, actual):
        if isinstance(value, str):
            assert isinstance(calculated, str)
        elif isinstance(value, (float, int)):
            assert isinstance(calculated, (float, int))


def test_real_excel_export_results_and_shifted_original_formulas(session, tmp_path):
    source, manager = session
    apply_column(manager)
    apply_column(manager, 1)
    target = tmp_path / "真实导出计算.xlsx"
    export_workbook(prepare_export(manager.get_export_snapshot(), ExportMode.ADD_FORMULA_COLUMNS), target)
    addresses = [("订单", f"{column}{row}") for column in ("B", "D") for row in range(2, 12)]
    addresses += [("订单", "F2"), ("订单", "F3"), ("订单", "A11")]
    digest = file_sha256(source)
    actual = _calculate_excel(target, addresses, tmp_path)
    expected = []
    for column in (0, 1):
        expected.extend("" if pd.isna(value) else value for value in manager.working_df.iloc[:, column])
    expected += [1014.5, 2004, " "]
    assert actual == expected
    assert file_sha256(source) == digest
