"""简洁 / 严格公式的区别、引用、类型保护与实际 Excel 计算。"""

import pytest
from openpyxl.formula import Tokenizer

from app.core.cleaner import CleaningRules
from app.core.manual_formula_builder import build_manual_formula


@pytest.mark.parametrize("rules, fragments", [
    (CleaningRules(nbsp=True), ('UNICHAR(160)," "',)),
    (CleaningRules(tab=True, newline=True, carriage_return=True, hidden_strategy="delete"),
     ('CHAR(9),""', 'CHAR(10),""', 'CHAR(13),""')),
    (CleaningRules(nbsp=True, tab=True, outer_spaces=True),
     ("UNICHAR(160)", 'CHAR(9)," "', "TRIM(")),
])
def test_paste_ready_formulas_protect_types_and_keep_selected_character_strategy(rules, fragments):
    formula = build_manual_formula(rules, "$AA2")
    assert "_xlfn." not in formula
    assert Tokenizer(formula).items
    assert 'IF(ISBLANK($AA2),"",IF(ISTEXT($AA2),' in formula
    for fragment in fragments:
        assert fragment in formula


@pytest.mark.parametrize("reference, expected", [
    ("D2", "D2"), ("aa2", "AA2"), (" $D2 ", "$D2"),
    ("$XFD$1048576", "$XFD$1048576"),
])
def test_source_reference_is_validated_and_normalized(reference, expected):
    formula = build_manual_formula(CleaningRules(tab=True), reference)
    assert expected in formula
    assert 'CHAR(9)," "' in formula
    assert formula.startswith(f'=IF(ISBLANK({expected}),"",IF(ISTEXT({expected}),')


@pytest.mark.parametrize("reference", ["", "A0", "A1048577", "XFE2", "A2)", "Sheet1!A2", "=A2"])
@pytest.mark.parametrize("mode", ["simple", "strict"])
def test_invalid_references_are_not_copied_as_formulas(reference, mode):
    with pytest.raises(ValueError, match="单元格引用"):
        build_manual_formula(CleaningRules(outer_spaces=True), reference, mode=mode)


@pytest.mark.parametrize("rules", [
    CleaningRules(fullwidth=True),
    CleaningRules(numeric_conversion=True),
    CleaningRules(date_normalize=True, date_confirmed=True),
    CleaningRules(nbsp=True, numeric_conversion=True),
])
def test_advanced_rules_require_explicit_separate_handling(rules):
    with pytest.raises(ValueError, match="暂不支持"):
        build_manual_formula(rules, "D2")


def test_outer_space_formula_preserves_internal_spaces_and_types():
    formula = build_manual_formula(CleaningRules(outer_spaces=True), "D2", mode="strict")
    assert "MID(D2," in formula
    assert formula.startswith('=IF(ISBLANK(D2),"",IF(ISTEXT(D2),')
    assert formula.endswith(",D2))")


def test_default_nbsp_and_outer_spaces_formula_is_short_and_preserves_type_guards():
    rules = CleaningRules(nbsp=True, outer_spaces=True)
    formula = build_manual_formula(rules, "C2")
    assert formula == '=IF(ISBLANK(C2),"",IF(ISTEXT(C2),TRIM(SUBSTITUTE(C2,UNICHAR(160)," ")),C2))'
    assert len(formula) < 100
    assert len(formula) < len(build_manual_formula(rules, "C2", mode="strict"))


def test_simple_mode_does_not_add_trim_when_not_selected():
    formula = build_manual_formula(CleaningRules(nbsp=True), "D2")
    assert "TRIM(" not in formula
    assert 'SUBSTITUTE(D2,UNICHAR(160)," ")' in formula


def test_unknown_formula_mode_is_rejected():
    with pytest.raises(ValueError, match="公式模式"):
        build_manual_formula(CleaningRules(), mode="unknown")


def test_real_excel_simple_and_strict_results_preserve_types_and_differ_only_as_expected(tmp_path):
    from openpyxl import Workbook
    from tests.test_export_formulas import _calculate_excel

    # 每项分别给出原值、简洁结果、严格结果，直接验证实际计算而非仅检查公式字符串。
    cases = [
        (None, "", ""),
        ("", "", ""),
        (0, 0, 0),
        ("0", "0", "0"),
        (1002, 1002, 1002),
        ("00123", "00123", "00123"),
        ("  ABC  123  ", "ABC 123", "ABC  123"),
        (" A\u00a0\u00a0B ", "A B", "A  B"),
        ("   ", "", ""),
        ("P001", "P001", "P001"),
        (" =1+1 ", "=1+1", "=1+1"),
        (True, True, True),
        ("  中文  😀  ", "中文 😀", "中文  😀"),
        ("A\tB", "A\tB", "A\tB"),
    ]
    book = Workbook()
    sheet = book.active
    sheet.title = "两种公式"
    sheet.append(["简洁结果", "严格结果", "源值", "未勾选首尾空格"])
    rules = CleaningRules(nbsp=True, outer_spaces=True)
    addresses, expected = [], []
    for row, (value, simple, strict) in enumerate(cases, 2):
        original = sheet.cell(row, 3, value)
        if isinstance(value, str):
            original.data_type = "s"
        reference = f"C{row}"
        for column, mode, result in ((1, "simple", simple), (2, "strict", strict)):
            formula = build_manual_formula(rules, reference, mode=mode)
            # 粘贴时使用裸函数名；测试写入 xlsx 时补上文件格式所需的前缀。
            sheet.cell(row, column, formula.replace("UNICHAR(", "_xlfn.UNICHAR("))
            addresses.append((sheet.title, f"{'AB'[column - 1]}{row}"))
            expected.append(result)
        formula = build_manual_formula(CleaningRules(nbsp=True), reference)
        sheet.cell(row, 4, formula.replace("UNICHAR(", "_xlfn.UNICHAR("))
        addresses.append((sheet.title, f"D{row}"))
        expected.append(value.replace("\u00a0", " ") if isinstance(value, str) else ("" if value is None else value))
    path = tmp_path / "简洁与严格公式.xlsx"
    book.save(path)
    book.close()
    # openpyxl 会遗漏纯空格文本的 xml:space；补齐输入标记，避免 Excel 先把样本当空白。
    from zipfile import ZipFile
    from xml.etree import ElementTree

    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for info, data in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(data)
                namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                for text in root.iter(namespace + "t"):
                    if text.text and not text.text.strip():
                        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                data = ElementTree.tostring(root)
            archive.writestr(info, data)
    actual = _calculate_excel(path, addresses, tmp_path)
    assert actual == expected
    for result, calculated in zip(expected, actual):
        if isinstance(result, str):
            assert isinstance(calculated, str)
        elif isinstance(result, bool):
            assert isinstance(calculated, bool)
        else:
            assert isinstance(calculated, (int, float)) and not isinstance(calculated, bool)
