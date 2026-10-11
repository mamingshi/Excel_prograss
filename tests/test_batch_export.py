"""批量应用后沿用两种导出模式，检查辅助列、公式引用及源文件保护。"""

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from app.core.annotation_builder import generate_duplicate_preview
from app.core.batch_cleaner import generate_batch_preview
from app.core.data_manager import WorkbookDataManager
from app.core.exporter import export_workbook, prepare_export
from app.core.file_loader import ExcelFileLoader, file_sha256
from app.models.cleaning_rule_config import ColumnCleaningConfig
from app.models.export_data import ExportMode


@pytest.mark.parametrize("mode", list(ExportMode))
def test_batch_export_preserves_source_and_refreshes_existing_annotations(tmp_path, mode):
    source = tmp_path / "批量源.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet.append(["商品编码", "金额", "订单日期"])
    sheet.append([" P001 ", "1002", "2026/10/11"])
    sheet.append(["P001", 0, "2026-10-12"])
    sheet.append(["P002", 12.5, "2026年10月13日"])
    workbook.save(source)
    workbook.close()
    digest = file_sha256(source)
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(source))
    original = manager.original_df.copy(deep=True)
    annotation = generate_duplicate_preview(manager.working_df, 0)
    manager.apply_annotation(manager.prepare_annotation_preview(annotation))
    before = manager.working_df.copy(deep=True)
    configs = (
        ColumnCleaningConfig(0, "商品编码", True, ("outer_spaces",)),
        ColumnCleaningConfig(1, "金额", True, ("numeric_conversion",)),
        ColumnCleaningConfig(2, "订单日期", True, ("date_normalize",), date_confirmed=True),
    )
    snapshot = manager.prepare_sheet_scan()
    preview = generate_batch_preview(snapshot.dataframe, configs, "订单", original_column_count=3)
    manager.apply_batch_preview(manager.prepare_batch_preview(snapshot, preview))
    assert manager.working_df.iloc[:, 3].tolist() == ["重复", "重复", "正常"]
    assert len(manager.undo_stack) == 2
    plan = prepare_export(manager.get_export_snapshot(), mode)
    assert plan.cell_count == preview.changed_cell_count == 4
    target = tmp_path / f"导出_{mode.name}.xlsx"
    result = export_workbook(plan, target)
    exported = load_workbook(target)
    try:
        sheet = exported["订单"]
        if mode == ExportMode.UPDATE_COLUMNS:
            assert sheet["A2"].value == "P001"
            assert sheet["B2"].value == 1002 and sheet["B2"].data_type == "n"
            assert sheet["C2"].value == "2026-10-11"
            assert sheet["D1"].value == "商品编码_重复检查"
            assert "A2" in sheet["D2"].value and "$A$2:$A$4" in sheet["D2"].value
        else:
            assert [cell.value for cell in sheet[1]] == [
                "商品编码", "商品编码_清洗后", "金额", "金额_清洗后",
                "订单日期", "订单日期_清洗后", "商品编码_重复检查",
            ]
            assert sheet["A2"].value == " P001 "
            assert sheet["B2"].data_type == "f" and "A2" in sheet["B2"].value
            assert sheet["C2"].value == "1002"
            assert sheet["D2"].value == 1002 and sheet["D2"].data_type == "n"
            assert sheet["F2"].value == "2026-10-11"
            assert "$B$2:$B$4" in sheet["G2"].value and "B2" in sheet["G2"].value
    finally:
        exported.close()
    assert result.cell_count == 4
    assert file_sha256(source) == digest
    pd.testing.assert_frame_equal(manager.original_df, original)
    manager.undo()
    pd.testing.assert_frame_equal(manager.working_df, before)
    assert len(manager.current_data.annotations) == 1
