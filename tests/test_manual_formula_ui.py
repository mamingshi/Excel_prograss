"""真实 Qt 交互验证独立公式、复制、选列及工作状态和源文件保护。"""

from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.worksheet.table import Table
import pandas as pd
import pytest
from PySide6.QtCore import QTimer, Qt
from PySide6.QtTest import QTest

from app.core.file_loader import file_sha256
from app.ui.main_window import MainWindow
from app.ui.manual_formula_dialog import ManualFormulaDialog
from tests.test_sheet_scan_ui import wait_for_idle


@pytest.fixture
def window(tmp_path, qapp):
    path = tmp_path / "含表格对象.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["ID", "金额", "备注", "商品编码"])
    sheet.append([1, 0, "正常", " P001\u00a0\t "])
    sheet.append([2, 10, "正常", "P002"])
    sheet.add_table(Table(displayName="Products", ref="A1:D3"))
    book.create_sheet("第二张").append(["编码"])
    book.save(path)
    book.close()
    widget = MainWindow()
    widget.show()
    widget._start_load(str(path))
    wait_for_idle(widget)
    yield widget
    wait_for_idle(widget)
    widget.close()
    widget.deleteLater()
    qapp.processEvents()


def visit_formula(window, qapp, inspect, button=None):
    """进入真实模态窗口，检查后关闭；把 Qt 回调中的断言带回测试。"""
    errors = []

    def visit():
        dialog = qapp.activeModalWidget()
        try:
            assert isinstance(dialog, ManualFormulaDialog)
            inspect(dialog)
        except BaseException as error:
            errors.append(error)
        finally:
            if dialog is not None:
                dialog.reject()

    QTimer.singleShot(0, visit)
    QTest.mouseClick(button or window.formula_button, Qt.MouseButton.LeftButton)
    if errors:
        raise errors[0]


def test_tool_is_available_without_opening_or_exporting_a_file(qapp):
    window = MainWindow()
    window.show()
    try:
        assert window.formula_button.isEnabled()
        assert not window.export_button.isEnabled()

        def inspect(dialog):
            assert "无需打开或导出" in dialog.context_label.text()
            assert dialog.mode_combo.currentData() == "simple"
            dialog.reference_edit.setText("AA2")
            dialog.rule_boxes["tab"].setChecked(True)
            QTest.mouseClick(dialog.copy_button, Qt.MouseButton.LeftButton)
            assert qapp.clipboard().text() == dialog.formula_edit.toPlainText()
            assert 'SUBSTITUTE(AA2,CHAR(9)," ")' in qapp.clipboard().text()

        visit_formula(window, qapp, inspect)
        assert window.data_manager.current_sheet is None
        assert not window.data_manager.undo_stack
    finally:
        window.close()
        window.deleteLater()


def test_excel_table_and_selected_column_need_no_detection_or_applied_changes(window, qapp):
    window._select_column(3)
    assert window.cleaning_panel.analysis is None
    assert not window.export_button.isEnabled()
    original = window.data_manager.original_df.copy(deep=True)
    working = window.data_manager.working_df.copy(deep=True)
    digest = file_sha256(window.data_manager.file_path)
    token = window.data_manager.export_state_token
    configs = window.batch_panel.configs

    def inspect(dialog):
        assert dialog.reference_edit.text() == "D2"
        assert "商品编码" in dialog.context_label.text()
        dialog.rule_boxes["nbsp"].setChecked(True)
        dialog.rule_boxes["outer_spaces"].setChecked(True)
        assert dialog.mode_combo.currentData() == "simple"
        assert len(dialog.formula_edit.toPlainText()) < 100
        dialog.mode_combo.setCurrentIndex(1)
        assert "MID(" in dialog.formula_edit.toPlainText()
        dialog.mode_combo.setCurrentIndex(0)
        assert "UNICHAR(160)" in dialog.formula_edit.toPlainText()
        assert "_xlfn." not in dialog.formula_edit.toPlainText()
        QTest.mouseClick(dialog.copy_button, Qt.MouseButton.LeftButton)
        assert qapp.clipboard().text().startswith("=")

    with patch("app.ui.main_window.QFileDialog.getSaveFileName", side_effect=AssertionError("不应导出")):
        visit_formula(window, qapp, inspect)
    pd.testing.assert_frame_equal(window.data_manager.original_df, original)
    pd.testing.assert_frame_equal(window.data_manager.working_df, working)
    assert file_sha256(window.data_manager.file_path) == digest
    assert window.data_manager.export_state_token == token
    assert window.batch_panel.configs == configs
    assert not window.data_manager.undo_stack
    assert not window.cleaning_panel.rule_boxes["nbsp"].isChecked()


def test_single_column_rule_and_strategy_are_prefilled_without_changing_them(window, qapp):
    window._select_column(3)
    window._start_analysis()
    wait_for_idle(window)
    window.cleaning_panel.strategy_combo.setCurrentIndex(1)
    before = window.cleaning_panel.current_rules()

    def inspect(dialog):
        assert dialog.reference_edit.text() == "D2"
        assert dialog.rule_boxes["nbsp"].isChecked()
        assert dialog.rule_boxes["tab"].isChecked()
        assert dialog.strategy_combo.currentData() == "delete"
        assert 'CHAR(9),""' in dialog.formula_edit.toPlainText()
        dialog.rule_boxes["tab"].setChecked(False)

    visit_formula(window, qapp, inspect, window.cleaning_panel.formula_button)
    assert window.cleaning_panel.current_rules() == before


def test_batch_formula_uses_its_selected_field_and_own_rules_without_scan(window, qapp):
    window._select_column(0)
    window._open_batch_workspace()
    panel = window.batch_panel
    assert not panel.scanned
    panel.config_model.update_config(3, enabled=True, rules=("tab", "outer_spaces"))
    panel.config_table.selectRow(3)
    configs = panel.configs

    def inspect(dialog):
        assert dialog.reference_edit.text() == "D2"
        assert "商品编码" in dialog.context_label.text()
        assert dialog.rule_boxes["tab"].isChecked()
        assert dialog.rule_boxes["outer_spaces"].isChecked()
        assert not dialog.rule_boxes["nbsp"].isChecked()
        dialog.rule_boxes["nbsp"].setChecked(True)

    visit_formula(window, qapp, inspect)
    assert panel.configs == configs
    assert window.selected_column == 0
    assert not window.data_manager.undo_stack


def test_invalid_reference_clears_old_formula_and_recovery_enables_copy(qapp):
    dialog = ManualFormulaDialog(selected_rules=("tab",), cell_reference="D2")
    try:
        assert dialog.copy_button.isEnabled()
        dialog.reference_edit.setText("D0")
        assert not dialog.copy_button.isEnabled()
        assert not dialog.formula_edit.toPlainText()
        assert "单元格引用" in dialog.status_label.text()
        dialog.mode_combo.setCurrentIndex(1)
        assert not dialog.copy_button.isEnabled() and not dialog.formula_edit.toPlainText()
        dialog.mode_combo.setCurrentIndex(0)
        assert not dialog.copy_button.isEnabled() and not dialog.formula_edit.toPlainText()
        qapp.clipboard().setText("保留剪贴板")
        dialog._copy_formula()
        assert qapp.clipboard().text() == "保留剪贴板"
        dialog.reference_edit.setText("$AA2")
        assert dialog.copy_button.isEnabled()
        assert "$AA2" in dialog.formula_edit.toPlainText()
    finally:
        dialog.close()
        dialog.deleteLater()


def test_advanced_rule_is_not_silently_omitted_from_copied_formula(qapp):
    dialog = ManualFormulaDialog(selected_rules=("nbsp", "fullwidth"))
    try:
        assert not dialog.copy_button.isEnabled()
        assert not dialog.formula_edit.toPlainText()
        assert "全角转半角" in dialog.status_label.text()
        dialog.mode_combo.setCurrentIndex(1)
        assert not dialog.copy_button.isEnabled() and not dialog.formula_edit.toPlainText()
        dialog.rule_boxes["fullwidth"].setChecked(False)
        assert dialog.copy_button.isEnabled()
        assert "UNICHAR(160)" in dialog.formula_edit.toPlainText()
    finally:
        dialog.close()
        dialog.deleteLater()


def test_changes_to_rules_and_strategy_update_exact_clipboard_content(qapp):
    dialog = ManualFormulaDialog(cell_reference="D2", selected_rules=("tab",))
    dialog.show()
    try:
        first = dialog.formula_edit.toPlainText()
        dialog.strategy_combo.setCurrentIndex(1)
        assert first != dialog.formula_edit.toPlainText()
        assert 'CHAR(9),""' in dialog.formula_edit.toPlainText()
        dialog.rule_boxes["newline"].setChecked(True)
        assert 'CHAR(10),""' in dialog.formula_edit.toPlainText()
        dialog.reference_edit.setText("AA7")
        QTest.mouseClick(dialog.copy_button, Qt.MouseButton.LeftButton)
        assert qapp.clipboard().text() == dialog.formula_edit.toPlainText()
        assert "AA7" in qapp.clipboard().text() and "D2" not in qapp.clipboard().text()
    finally:
        dialog.close()
        dialog.deleteLater()


def test_tool_remains_available_after_apply_and_undo(window, qapp):
    window._select_column(3)
    window._start_analysis()
    wait_for_idle(window)
    window._start_cleaning_preview()
    wait_for_idle(window)
    with patch("app.ui.main_window.confirm_apply", return_value=True):
        window._apply_cleaning()
    wait_for_idle(window)
    assert window.formula_button.isEnabled() and window.data_manager.undo_stack
    history = list(window.data_manager.undo_stack)
    visit_formula(window, qapp, lambda dialog: dialog.rule_boxes["tab"].setChecked(True))
    assert window.data_manager.undo_stack == history
    window._undo_last_action()
    wait_for_idle(window)
    assert window.formula_button.isEnabled() and not window.data_manager.undo_stack


def test_switching_modes_updates_clipboard_without_resetting_reference_rules_or_strategy(qapp):
    dialog = ManualFormulaDialog(
        cell_reference="$AA7", selected_rules=("nbsp", "tab", "outer_spaces"), hidden_strategy="delete",
    )
    dialog.show()
    try:
        assert dialog.mode_combo.currentData() == "simple"
        simple = dialog.formula_edit.toPlainText()
        assert "合并为一个" in dialog.mode_note_label.text()
        checked = {name: box.isChecked() for name, box in dialog.rule_boxes.items()}
        dialog.mode_combo.setCurrentIndex(1)
        strict = dialog.formula_edit.toPlainText()
        assert len(simple) < len(strict) and "MID(" in strict
        assert "保留中间" in dialog.mode_note_label.text()
        assert dialog.reference_edit.text() == "$AA7"
        assert dialog.strategy_combo.currentData() == "delete"
        assert {name: box.isChecked() for name, box in dialog.rule_boxes.items()} == checked
        QTest.mouseClick(dialog.copy_button, Qt.MouseButton.LeftButton)
        assert qapp.clipboard().text() == strict
        assert "已复制严格版" in dialog.status_label.text()
        dialog.mode_combo.setCurrentIndex(0)
        assert dialog.formula_edit.toPlainText() == simple
        QTest.mouseClick(dialog.copy_button, Qt.MouseButton.LeftButton)
        assert qapp.clipboard().text() == simple
        assert "已复制简洁版" in dialog.status_label.text()
    finally:
        dialog.close()
        dialog.deleteLater()


def test_mode_note_follows_outer_spaces_option_and_new_windows_default_to_simple(qapp):
    dialog = ManualFormulaDialog(selected_rules=("nbsp",))
    try:
        assert dialog.mode_combo.currentData() == "simple"
        assert "未勾选" in dialog.mode_note_label.text()
        assert "TRIM(" not in dialog.formula_edit.toPlainText()
        dialog.rule_boxes["outer_spaces"].setChecked(True)
        assert "合并为一个" in dialog.mode_note_label.text()
        assert "TRIM(" in dialog.formula_edit.toPlainText()
        dialog.mode_combo.setCurrentIndex(1)
        assert "只去除首尾" in dialog.mode_note_label.text()
        another = ManualFormulaDialog(selected_rules=("nbsp", "outer_spaces"))
        try:
            assert another.mode_combo.currentData() == "simple"
        finally:
            another.close()
            another.deleteLater()
    finally:
        dialog.close()
        dialog.deleteLater()
