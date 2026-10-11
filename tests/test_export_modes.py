"""两种模式、相邻公式列、插列引用、规则回滚及安全拒绝的真实工作簿测试。"""

from dataclasses import replace
from copy import copy

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
import pandas as pd
import pytest

from app.core.cleaner import CleaningRules, generate_cleaning_preview
from app.core.column_inserter import ColumnMap, ColumnInsertionError, translate_formula
from app.core.exporter import ExcelExportError, export_workbook, prepare_export
from app.core.file_loader import ExcelFileLoader, file_sha256
from app.models.export_data import ExportMode
from tests.test_exporter import session, apply_column


def plan_for(manager):
    return prepare_export(manager.get_export_snapshot(), ExportMode.ADD_FORMULA_COLUMNS)


def apply_rules(manager, rules, column=0):
    result = generate_cleaning_preview(manager.working_df.iloc[:, column], rules)
    manager.apply_preview(manager.prepare_preview(column, result))


def test_adjacent_formula_column_keeps_dirty_values_and_preserves_metadata(session):
    source, manager = session
    apply_column(manager)
    digest = file_sha256(source)
    before_working = manager.working_df.copy(deep=True)
    target = source.with_name("对照.xlsx")
    result = export_workbook(plan_for(manager), target)
    assert result.mode == ExportMode.ADD_FORMULA_COLUMNS
    assert result.added_column_count == 1 and result.cell_count == 5
    book, original = load_workbook(target), load_workbook(source)
    try:
        sheet = book["订单"]
        assert [cell.value for cell in sheet[1]] == ["商品编码", "商品编码_清洗后", "数量", "说明", "公式"]
        assert sheet["A2"].value == " P001 "
        assert sheet["A3"].value == "P002\u00a0"
        assert all(sheet.cell(row, 2).data_type == "f" for row in range(2, 12))
        assert "A2" in sheet["B2"].value and "_xlfn.UNICHAR(160)" in sheet["B2"].value
        assert sheet["C2"].value == 0 and sheet["C3"].value == 1002 and sheet["C4"].value == 12.5
        assert sheet["E2"].value == "=SUM(C2:C4)"
        assert sheet["E3"].value == "=C3*2"
        for attribute in ("font", "fill", "border", "alignment", "number_format"):
            assert copy(getattr(sheet["B2"], attribute)) == copy(getattr(original["订单"]["A2"], attribute))
        assert sheet.column_dimensions["A"].width == sheet.column_dimensions["B"].width == 24
        assert sheet.row_dimensions[2].height == 28
        assert sheet.freeze_panes == "C2"
        assert sheet.auto_filter.ref == "A1:E11"
        assert str(sheet.print_area) == "'订单'!$A$1:$E$11"
        assert str(sheet.data_validations.dataValidation[0].sqref) == "C2:C11"
        assert sheet.print_title_rows == original["订单"].print_title_rows
        assert book["说明"].sheet_state == "hidden"
        assert str(book["说明"].merged_cells) == "E7:F7"
        assert list(book["说明"].values) == list(original["说明"].values)
        assert book.calculation.calcMode == "auto"
        assert book.calculation.fullCalcOnLoad and book.calculation.forceFullCalc
    finally:
        book.close()
        original.close()
    assert file_sha256(source) == digest
    pd.testing.assert_frame_equal(manager.working_df, before_working)


def test_multiple_columns_and_sheets_each_get_one_final_column(session):
    source, manager = session
    apply_column(manager)
    apply_column(manager, 1)
    manager.add_sheet(ExcelFileLoader().load(source, "客户"))
    apply_column(manager)
    result = export_workbook(plan_for(manager), source.with_name("多列.xlsx"))
    assert result.added_column_count == 3 and result.sheet_count == 2
    book = load_workbook(result.target_path)
    try:
        sheet = book["订单"]
        assert [cell.value for cell in sheet[1]] == [
            "商品编码", "商品编码_清洗后", "数量", "数量_清洗后", "说明", "公式"
        ]
        assert "A3" in sheet["B3"].value
        assert "C3" in sheet["D3"].value
        assert sheet["F2"].value == "=SUM(C2:C4)"
        assert sheet["F3"].value == "=C3*2"
        assert sheet["C10"].value == " B010 "
        assert book["客户"]["A2"].value == " X001 "
        assert book["客户"]["B2"].data_type == "f"
        assert book["客户"]["C2"].value == "保持"
    finally:
        book.close()


def test_undo_and_restore_roll_back_formula_rules(session):
    source, manager = session
    first = CleaningRules(outer_spaces=True)
    second = CleaningRules(nbsp=True, outer_spaces=True)
    apply_rules(manager, first)
    apply_rules(manager, second)
    assert manager.current_data.column_rules[0] == (first, second)
    snapshot = manager.get_export_snapshot()
    manager.undo()
    assert snapshot.sheets[0].column_rules[0] == (first, second)
    assert manager.current_data.column_rules[0] == (first,)
    target = source.with_name("撤销规则.xlsx")
    export_workbook(plan_for(manager), target)
    book = load_workbook(target)
    try:
        assert "UNICHAR(160)" not in book["订单"]["B3"].value
    finally:
        book.close()
    manager.restore_original()
    assert not manager.current_data.column_rules
    with pytest.raises(ExcelExportError, match="无需导出"):
        plan_for(manager)
    manager.undo()
    assert manager.current_data.column_rules[0] == (first,)
    manager.undo()
    assert not manager.current_data.column_rules
    assert not manager.is_modified


def test_sequential_rules_and_mixed_hidden_strategies_are_preserved(session):
    source, manager = session
    first = CleaningRules(outer_spaces=True)
    second = CleaningRules(tab=True, hidden_strategy="delete", nbsp=True)
    third = CleaningRules(outer_spaces=True)
    apply_rules(manager, first)
    apply_rules(manager, second)
    apply_rules(manager, third)
    target = source.with_name("多步.xlsx")
    export_workbook(plan_for(manager), target)
    book = load_workbook(target)
    try:
        formula = book["订单"]["B2"].value
        assert 'CHAR(9),""' in formula
        assert '_xlfn.UNICHAR(160)," "' in formula
    finally:
        book.close()


def test_header_collision_creates_unique_name(session):
    source, manager = session
    book = load_workbook(source)
    book["订单"]["C1"] = "商品编码_清洗后"
    book.save(source)
    book.close()
    manager.open_workbook(ExcelFileLoader().load(source))
    apply_column(manager)
    target = source.with_name("重名.xlsx")
    export_workbook(plan_for(manager), target)
    book = load_workbook(target)
    try:
        assert book["订单"]["B1"].value == "商品编码_清洗后_1"
        assert book["订单"]["D1"].value == "商品编码_清洗后"
    finally:
        book.close()


def test_other_sheet_formulas_names_validation_and_conditional_formatting_follow_original_cells(session):
    source, manager = session
    book = load_workbook(source)
    book["客户"]["C2"] = "='订单'!$B$3"
    book["订单"].print_title_cols = "A:B"
    book["订单"].auto_filter.add_filter_column(1, ["1002"])
    book.defined_names.add(DefinedName("数量列表", attr_text="'订单'!$B$2:$B$11"))
    validation = DataValidation(type="custom", formula1="'订单'!$B$2>=0")
    validation.add("A2:A3")
    book["客户"].add_data_validation(validation)
    book["订单"].conditional_formatting.add("B2:B11", FormulaRule(formula=["B2>100"]))
    book.save(source)
    book.close()
    manager.open_workbook(ExcelFileLoader().load(source))
    apply_column(manager)
    target = source.with_name("引用.xlsx")
    export_workbook(plan_for(manager), target)
    saved = load_workbook(target)
    try:
        assert saved["客户"]["C2"].value == "='订单'!$C$3"
        assert saved.defined_names["数量列表"].attr_text == "'订单'!$C$2:$C$11"
        assert saved["客户"].data_validations.dataValidation[0].formula1 == "'订单'!$C$2>=0"
        sheet = saved["订单"]
        assert sheet.print_title_cols == "$A:$C"
        assert sheet.auto_filter.filterColumn[0].colId == 2
        assert str(next(iter(sheet.conditional_formatting)).sqref) == "C2:C11"
        assert next(iter(sheet.conditional_formatting._cf_rules.values()))[0].formula == ["C2>100"]
    finally:
        saved.close()


@pytest.mark.parametrize("feature", ["crossing_formula", "dynamic_formula", "merged", "table"])
def test_complex_dependencies_refuse_new_mode_but_keep_original_mode_available(session, feature):
    source, manager = session
    book = load_workbook(source)
    sheet = book["订单"]
    if feature == "crossing_formula":
        sheet["D2"] = "=COUNTA(A2:C2)"
    elif feature == "dynamic_formula":
        sheet["D2"] = '=INDIRECT("B2")'
    elif feature == "merged":
        sheet.merge_cells("A10:B10")
    else:
        sheet.add_table(Table(displayName="订单数据", ref="A1:D11"))
    book.save(source)
    book.close()
    manager.open_workbook(ExcelFileLoader().load(source))
    apply_column(manager)
    digest = file_sha256(source)
    target = source.with_name("拒绝.xlsx")
    with pytest.raises(ExcelExportError, match="原列更新"):
        export_workbook(plan_for(manager), target)
    assert not target.exists() and file_sha256(source) == digest
    assert not list(source.parent.glob(".excel-cleaner-*"))
    result = export_workbook(prepare_export(manager.get_export_snapshot()), source.with_name("原列.xlsx"))
    assert result.mode == ExportMode.UPDATE_COLUMNS


def test_missing_or_inconsistent_rules_do_not_generate_guessed_formulas(session):
    source, manager = session
    apply_column(manager)
    snapshot = manager.get_export_snapshot()
    original = snapshot.sheets[0]
    for rules in ({}, {0: (CleaningRules(outer_spaces=True),)}):
        modified = replace(snapshot, sheets=(replace(original, column_rules=rules),))
        target = source.with_name("不猜测.xlsx")
        with pytest.raises(ExcelExportError, match="原列更新"):
            export_workbook(prepare_export(modified, ExportMode.ADD_FORMULA_COLUMNS), target)
        assert not target.exists()


@pytest.mark.parametrize("formula,expected", [
    ("=$B$2+B3", "=$C$2+C3"),
    ("='订单'!B2+SUM(B3:B4)", "='订单'!C2+SUM(C3:C4)"),
    ('=IF(B2="B2",B3,0)', '=IF(C2="B2",C3,0)'),
    ("=SUM($B:$B)", "=SUM($C:$C)"),
])
def test_a1_reference_mapping_does_not_shift_text_literals(formula, expected):
    assert translate_formula(formula, "订单", {"订单": ColumnMap((1,))}) == expected


def test_range_crossing_insertion_is_not_silently_expanded():
    with pytest.raises(ColumnInsertionError):
        translate_formula("=SUM(A2:C2)", "订单", {"订单": ColumnMap((1,))})

@pytest.mark.parametrize("formula", ["=SUM(XFD:XFD)", "=@B2", "=SUM(B2#)"])
def test_unrepresentable_references_cancel_instead_of_silently_shifting(formula):
    with pytest.raises(ColumnInsertionError, match="原列更新"):
        translate_formula(formula, "订单", {"订单": ColumnMap((1,))})


def test_relative_global_name_is_not_guessed(session):
    source, manager = session
    book = load_workbook(source)
    book.defined_names.add(DefinedName("相对位置", attr_text="B2"))
    book.save(source)
    book.close()
    manager.open_workbook(ExcelFileLoader().load(source))
    apply_column(manager)
    target = source.with_name("相对名称.xlsx")
    with pytest.raises(ExcelExportError, match="相对定义名称"):
        export_workbook(plan_for(manager), target)
    assert not target.exists()


def test_comparison_columns_use_positions_with_duplicate_headers_and_custom_index(tmp_path):
    from openpyxl import Workbook
    from app.core.data_manager import WorkbookDataManager
    source = tmp_path / "重复字段.xlsx"
    book = Workbook()
    book.active.append(["编码", "编码", None])
    book.active.append([" A2 ", " B2 ", "保持"])
    book.active.append([" A3 ", " B3 ", "0"])
    book.save(source)
    book.close()
    preview = ExcelFileLoader().load(source)
    preview.dataframe.index = [100, 7]
    manager = WorkbookDataManager()
    manager.open_workbook(preview)
    apply_column(manager, 0)
    apply_column(manager, 1)
    target = source.with_name("重复字段对照.xlsx")
    export_workbook(plan_for(manager), target)
    saved = load_workbook(target)
    try:
        sheet = saved.active
        assert [cell.value for cell in sheet[1]] == ["编码", "编码_清洗后", "编码", "编码.1_清洗后", None]
        assert sheet["A2"].value == " A2 " and sheet["C3"].value == " B3 "
        assert "A2" in sheet["B2"].value and "C3" in sheet["D3"].value
        assert sheet["E3"].value == "0"
    finally:
        saved.close()

