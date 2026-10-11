"""V0.8 真实工作簿、辅助列、Undo、两种导出模式及 Qt 交互。"""

from datetime import datetime
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest

from app.core.annotation_builder import generate_annotation_preview, build_duplicate_formula
from app.core.cleaner import CleaningRules, generate_cleaning_preview
from app.core.data_manager import WorkbookDataManager
from app.core.exporter import prepare_export, export_workbook
from app.core.file_loader import ExcelFileLoader, file_sha256
from app.models.export_data import ExportMode
from app.ui.main_window import MainWindow
from tests.test_sheet_scan_ui import wait_for_idle
from tests.test_export_formulas import _calculate_excel


@pytest.fixture
def session(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["编码", "订单日期", "备注", "计算"])
    codes = [1001, "1002", "00123", "Ｐ００１", "ABC", "P001"]
    dates = ["2026-10-10", "2026/10/10", "2026年10月10日", "01/02/03", "abc", datetime(2026, 10, 11)]
    for row, (code, day) in enumerate(zip(codes, dates), start=2):
        sheet.append([code, day, f"备注{row}", "=1+1" if row == 2 else None])
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = "A1:D7"
    source = tmp_path / "类型测试.xlsx"
    book.save(source)
    book.close()
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(source))
    return source, manager


def clean(manager, column, rules):
    result = generate_cleaning_preview(manager.working_df.iloc[:, column], rules)
    prepared = manager.prepare_preview(column, result)
    manager.apply_preview(prepared)
    return result


def annotate(manager, column, kind, confirmed=False):
    result = generate_annotation_preview(manager.working_df, column, kind, date_confirmed=confirmed)
    manager.apply_annotation(manager.prepare_annotation_preview(result))
    return result


def test_helpers_labels_naming_refresh_and_reverse_undo(session):
    source, manager = session
    digest = file_sha256(source)
    before = manager.original_df.copy(deep=True)
    first = annotate(manager, 0, "numeric_type")
    assert first.values.tolist() == ["数字", "文本数字", "文本数字（前导0）", "文本", "文本", "文本"]
    second = annotate(manager, 0, "leading_zero")
    assert second.values.tolist() == ["正常", "正常", "前导0", "正常", "正常", "正常"]
    third = annotate(manager, 1, "date")
    assert third.values.tolist() == ["正常", "格式不统一", "格式不统一", "日期歧义", "无法识别", "正常"]
    again = annotate(manager, 0, "leading_zero")
    assert again.annotation.column_name == "编码_前导0风险_2"
    clean(manager, 0, CleaningRules(numeric_conversion=True))
    assert manager.working_df.iloc[1, 4] == "数字"  # 已有类型辅助列随源列更新
    manager.undo()
    assert manager.working_df.iloc[1, 4] == "文本数字"
    for _ in range(4):
        manager.undo()
    pd.testing.assert_frame_equal(manager.working_df, before)
    assert manager.current_data.annotations == ()
    assert file_sha256(source) == digest


def test_restore_and_undo_restore_preserve_new_annotation_metadata(session):
    _, manager = session
    annotate(manager, 0, "numeric_type")
    annotate(manager, 1, "date", True)
    clean(manager, 1, CleaningRules(date_normalize=True, date_confirmed=True))
    before = manager.working_df.copy(deep=True)
    specs = manager.current_data.annotations
    manager.restore_original()
    pd.testing.assert_frame_equal(manager.working_df, manager.original_df)
    assert manager.current_data.annotations == ()
    manager.undo()
    pd.testing.assert_frame_equal(manager.working_df, before)
    assert manager.current_data.annotations == specs
    manager.undo()
    assert manager.working_df.iloc[1, 1] == "2026/10/10"


@pytest.mark.parametrize("mode", list(ExportMode))
def test_both_exports_preserve_source_and_write_python_results_and_exact_duplicate_formula(session, tmp_path, mode):
    source, manager = session
    digest = file_sha256(source)
    clean(manager, 0, CleaningRules(fullwidth=True, numeric_conversion=True))
    clean(manager, 1, CleaningRules(date_normalize=True, date_confirmed=True))
    for column, kind in ((0, "numeric_type"), (0, "leading_zero"), (1, "date"), (0, "duplicate")):
        annotate(manager, column, kind)
    plan = prepare_export(manager.get_export_snapshot(), mode)
    target = tmp_path / (mode.value + ".xlsx")
    result = export_workbook(plan, target)
    assert result.annotation_column_count == 4
    book = load_workbook(target, data_only=False)
    sheet = book["订单"]
    if mode == ExportMode.UPDATE_COLUMNS:
        assert sheet["A3"].value == 1002 and sheet["A3"].data_type == "n"
        assert sheet["A4"].value == "00123"
        assert sheet["B3"].value == "2026-10-10"
        assert sheet["B5"].value == "01/02/03"
        assert sheet["D2"].value == "=1+1"
        first, source_column = 5, 0
    else:
        assert sheet["A3"].value == "1002"
        assert sheet["B3"].value == 1002 and sheet["B3"].data_type == "n"
        assert sheet["B4"].value == "00123" and sheet["B4"].data_type == "s"
        assert sheet["D3"].value == "2026-10-10"
        assert sheet["D5"].value == "01/02/03"
        assert sheet["F2"].value == "=1+1"
        first, source_column = 7, 1
    assert sheet.cell(1, first).value == "编码_类型检查"
    assert sheet.cell(3, first).value == "数字"
    assert sheet.cell(4, first + 1).value == "前导0"
    assert sheet.cell(5, first + 2).value == "日期歧义"
    assert sheet.cell(6, first + 2).value == "无法识别"
    assert sheet.cell(2, first + 3).value == build_duplicate_formula(source_column, 2, 7)
    assert sheet.freeze_panes == ("B2" if mode == ExportMode.UPDATE_COLUMNS else "C2")
    book.close()
    assert file_sha256(source) == digest
    assert manager.working_df.iloc[2, 0] == "00123"


def test_fullwidth_equals_is_exported_as_literal_text_in_both_modes(tmp_path):
    book = Workbook()
    book.active.append(["编码"])
    book.active.append(["＝1+1"])
    path = tmp_path / "等号.xlsx"
    book.save(path)
    book.close()
    manager = WorkbookDataManager()
    manager.open_workbook(ExcelFileLoader().load(path))
    clean(manager, 0, CleaningRules(fullwidth=True))
    for mode in ExportMode:
        target = tmp_path / f"等号_{mode.value}.xlsx"
        export_workbook(prepare_export(manager.get_export_snapshot(), mode), target)
        output = load_workbook(target, data_only=False)
        cell = output.active.cell(2, 1 if mode == ExportMode.UPDATE_COLUMNS else 2)
        assert cell.value == "=1+1" and cell.data_type == "s"
        output.close()


def test_string_dtype_numeric_conversion_can_apply_and_undo_with_exact_dtype(session):
    _, manager = session
    frame = pd.DataFrame({"编码": pd.Series(["1002", "00123", pd.NA], dtype="string")})
    manager.current_data.original_df = frame.copy(deep=True)
    manager.current_data.working_df = frame.copy(deep=True)
    clean(manager, 0, CleaningRules(numeric_conversion=True))
    assert manager.working_df.dtypes.iloc[0] == object
    assert manager.working_df.iloc[0, 0] == 1002
    manager.undo()
    pd.testing.assert_frame_equal(manager.working_df, frame)


def test_date_annotation_requires_evidence_or_explicit_confirmation():
    data = pd.DataFrame({"日期编码": ["ABC", "DEF"]})
    with pytest.raises(ValueError, match="证据不足"):
        generate_annotation_preview(data, 0, "date")
    preview = generate_annotation_preview(data, 0, "date", date_confirmed=True)
    assert preview.values.tolist() == ["无法识别", "无法识别"]
    assert preview.formula(0) == "检测结果（导出为值）"


@pytest.fixture
def window(session, qapp):
    source, _ = session
    widget = MainWindow()
    widget.show()
    widget._start_load(str(source))
    wait_for_idle(widget)
    yield widget
    wait_for_idle(widget)
    widget.close()
    widget.deleteLater()
    qapp.processEvents()


def analyze(window, column):
    window._select_column(column)
    window._start_analysis()
    wait_for_idle(window)
    window._set_cleaning_sidebar_visible(True)


@pytest.mark.parametrize("kind,column", [("numeric_type", 0), ("leading_zero", 0), ("date", 1)])
def test_real_annotation_buttons_preview_apply_undo_without_source_write(window, kind, column):
    digest = file_sha256(window.data_manager.file_path)
    analyze(window, column)
    button = window.cleaning_panel.type_panel.annotation_buttons[kind]
    assert button.isEnabled()
    window.cleaning_scroll.ensureWidgetVisible(button)
    QTest.qWait(20)
    QTest.mouseClick(button, Qt.MouseButton.LeftButton)
    wait_for_idle(window)
    preview = window.annotation_preview_panel.result
    assert preview is not None and preview.annotation.kind == kind
    assert window.data_manager.working_df.shape[1] == 4
    assert window.annotation_preview_panel.model.data(window.annotation_preview_panel.model.index(0, 3)) == "检测结果（导出为值）"
    with patch("app.ui.main_window.confirm_annotation", return_value=True):
        window._apply_annotation()
    assert window.data_manager.working_df.shape[1] == 5
    assert not window.cleaning_panel.type_panel.annotation_buttons[kind].isEnabled()
    window._undo_last_action()
    wait_for_idle(window)
    assert window.data_manager.working_df.shape[1] == 4
    assert file_sha256(window.data_manager.file_path) == digest


def test_ui_requires_date_confirmation_retains_risk_rows_and_supports_cancel_and_apply(window):
    analyze(window, 1)
    panel = window.cleaning_panel.type_panel
    box = panel.rule_boxes["date_normalize"]
    assert not box.isEnabled()
    panel.date_confirmed_box.setChecked(True)
    assert box.isEnabled()
    box.setChecked(True)
    window._start_cleaning_preview()
    wait_for_idle(window)
    result = window.cleaning_preview_panel.result
    assert result.changed_count == 3
    assert window.cleaning_preview_panel.model.rowCount() == 5  # 三条变化与两条保留原因
    assert "保留" in window.cleaning_preview_panel.state_label.text()
    before = window.data_manager.working_df.copy(deep=True)
    with patch("app.ui.main_window.confirm_apply", return_value=False):
        window._apply_cleaning()
    pd.testing.assert_frame_equal(window.data_manager.working_df, before)
    with patch("app.ui.main_window.confirm_apply", return_value=True):
        window._apply_cleaning()
    wait_for_idle(window)
    assert window.data_manager.working_df.iloc[1, 1] == "2026-10-10"
    assert window.data_manager.working_df.iloc[3, 1] == "01/02/03"
    assert window.data_manager.working_df.iloc[4, 1] == "abc"
    window._undo_last_action()
    wait_for_idle(window)
    pd.testing.assert_frame_equal(window.data_manager.working_df, before)


def test_ui_numeric_preview_explains_leading_zero_and_selection_resets_confirmation(window):
    analyze(window, 0)
    panel = window.cleaning_panel.type_panel
    assert not panel.annotation_buttons["date"].isEnabled()
    panel.rule_boxes["numeric_conversion"].setChecked(True)
    window._start_cleaning_preview()
    wait_for_idle(window)
    assert any("前导0" in note for record in window.cleaning_preview_panel.result.records for note in record.notes)
    assert window.data_manager.working_df.iloc[1, 0] == "1002"
    analyze(window, 1)
    panel.date_confirmed_box.setChecked(True)
    panel.rule_boxes["date_normalize"].setChecked(True)
    window._select_column(0)
    assert not panel.date_confirmed_box.isChecked()
    assert not panel.rule_boxes["date_normalize"].isChecked()


def test_scan_ui_includes_four_new_counts_and_matches_detail(window):
    window._start_sheet_scan()
    wait_for_idle(window)
    model = window.sheet_overview_panel.model
    assert model.columnCount() == 15
    assert model.HEADERS[-4:] == ("文本数字混合", "前导0风险", "日期异常", "全角字符")
    result = window.sheet_overview_panel.result
    assert result.columns[0].numeric_mixture == 2
    assert result.columns[0].leading_zero == 1
    assert result.columns[0].fullwidth == 1
    assert result.columns[1].date_anomalies == 4


@pytest.mark.parametrize("mode", list(ExportMode))
def test_real_excel_calculates_duplicate_helpers_after_type_conversions(session, tmp_path, mode):
    _, manager = session
    clean(manager, 0, CleaningRules(fullwidth=True, numeric_conversion=True))
    annotate(manager, 0, "numeric_type")
    duplicate = annotate(manager, 0, "duplicate")
    target = tmp_path / f"实际类型计算_{mode.value}.xlsx"
    export_workbook(prepare_export(manager.get_export_snapshot(), mode), target)
    helper_letter = "G" if mode == ExportMode.ADD_FORMULA_COLUMNS else "F"
    addresses = [("订单", f"{helper_letter}{row}") for row in range(2, 8)]
    assert _calculate_excel(target, addresses, tmp_path) == duplicate.values.tolist()


@pytest.mark.parametrize("rule,column", [("fullwidth", 0), ("numeric_conversion", 0), ("date_normalize", 1)])
def test_changing_advanced_rule_invalidates_old_preview_and_disables_apply(window, rule, column):
    analyze(window, column)
    panel = window.cleaning_panel.type_panel
    if rule == "date_normalize":
        panel.date_confirmed_box.setChecked(True)
    window._start_cleaning_preview()
    wait_for_idle(window)
    assert window.cleaning_preview_panel.result is not None
    before = window.data_manager.working_df.copy(deep=True)
    box = panel.rule_boxes[rule]
    box.setChecked(not box.isChecked())
    assert window.cleaning_preview_panel.result is None
    assert window._prepared_preview is None
    assert not window.cleaning_preview_panel.apply_button.isEnabled()
    assert not window.data_manager.undo_stack
    pd.testing.assert_frame_equal(window.data_manager.working_df, before)
