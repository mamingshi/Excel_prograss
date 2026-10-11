"""真实工作簿上的清洗预览、规则交互、类型保持和源文件保护。"""

import hashlib
import time
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook
import pandas as pd
import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from app.ui.main_window import MainWindow


def wait_for_idle(window):
    deadline = time.monotonic() + 10
    while window._is_busy():
        QApplication.processEvents()
        QTest.qWait(10)
        assert time.monotonic() < deadline, "后台任务超时"
    QApplication.processEvents()


@pytest.fixture
def window(tmp_path, qapp):
    path = tmp_path / "清洗预览测试.xlsx"
    workbook = Workbook()
    workbook.active.title = "问题数据"
    workbook.active.append(["商品编码", "备注"])
    values = [" P001 ", "P002\u00a0", "ABC\t123", "ABC\n123", "A\rB", "ABC 123",
              1002, 0, None, " \tP003\u00a0\n ", "DUP", "DUP", "0", ""]
    for row, value in enumerate(values):
        workbook.active.append([value, f"备注 {row}"])
    normal = workbook.create_sheet("正常表")
    normal.append(["编码"])
    normal.append(["P001"])
    normal.append(["ABC 123"])
    workbook.create_sheet("仅字段").append(["编码"])
    workbook.save(path)
    workbook.close()
    # 使用合法 XML 字符引用保留 CR，避免测试写入器的原始 CR 被 XML 规范化为 LF。
    with ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with ZipFile(path, "w") as archive:
        for info, content in entries:
            if info.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"\r", b"&#13;")
            archive.writestr(info, content)
    widget = MainWindow()
    widget.show()
    qapp.processEvents()
    widget._start_load(str(path))
    wait_for_idle(widget)
    yield widget
    wait_for_idle(widget)
    widget.close()
    widget.deleteLater()
    qapp.processEvents()


def analyze(window, column=0):
    window._select_column(column)
    QTest.mouseClick(window.diagnosis_panel.start_button, Qt.MouseButton.LeftButton)
    wait_for_idle(window)
    assert window.cleaning_panel.generate_button.isEnabled()


def click_page(window, page):
    """通过实际页签点击验证导航，不对隐藏控件模拟点击。"""
    bar = window.results_tabs.tabBar()
    position = bar.tabRect(window.results_tabs.indexOf(page)).center()
    QTest.mouseClick(bar, Qt.MouseButton.LeftButton, pos=position)
    QApplication.processEvents()


def generate(window):
    click_page(window, window.cleaning_preview_panel)
    if window.cleaning_scroll.isHidden():
        QTest.mouseClick(window.cleaning_toggle_button, Qt.MouseButton.LeftButton)
    QTest.mouseClick(window.cleaning_panel.generate_button, Qt.MouseButton.LeftButton)
    assert window._cleaning_worker is not None
    assert not window.cleaning_panel.generate_button.isEnabled()
    assert not window.diagnosis_panel.start_button.isEnabled()
    assert not window.open_button.isEnabled()
    assert not window.sheet_combo.isEnabled()
    wait_for_idle(window)


def test_recommendations_require_detection_and_identify_each_character(window):
    panel = window.cleaning_panel
    assert not panel.generate_button.isEnabled()
    window._select_column(0)
    assert not panel.generate_button.isEnabled()
    window._start_cleaning_preview()
    assert window.cleaning_preview_panel.result is None
    analyze(window)
    assert all(box.isChecked() for box in panel.rule_boxes.values())
    assert panel.current_rules().hidden_strategy == "space"
    assert "空值：2 条" in panel.empty_label.text()
    assert "重复值：2 条" in panel.duplicate_label.text()
    assert "CHAR(13)" in panel.formula_edit.toPlainText()
    assert "不会修改原始数据" in panel.safety_label.text()


def test_preview_stats_differences_and_source_are_preserved(window):
    path = window._preview.file_path
    source_digest = hashlib.sha256(path.read_bytes()).digest()
    original = window._preview.dataframe.copy(deep=True)
    analyze(window)
    generate(window)
    panel = window.cleaning_preview_panel
    result = panel.result
    assert (result.row_count, result.changed_count, result.unchanged_count) == (14, 6, 8)
    assert result.processed_counts == {"前后空格": 3, "CHAR(160)": 2, "Tab / 换行": 4}
    assert panel.model.rowCount() == 6
    assert [record.excel_row for record in panel._visible_records] == [2, 3, 4, 5, 6, 11]
    assert panel.model.data(panel.model.index(1, 1)) == '"P002[NBSP]"'
    assert panel.model.data(panel.model.index(1, 2)) == '"P002"'
    assert panel.model.data(panel.model.index(1, 1), Qt.ItemDataRole.UserRole) == "P002\u00a0"
    assert result.preview_series.iloc[2] == "ABC 123"
    assert result.preview_series.iloc[4] == "A B"
    assert result.preview_series.iloc[6] == 1002
    assert type(result.preview_series.iloc[6]) is int
    assert result.preview_series.iloc[7] == 0
    assert result.preview_series.iloc[10:12].tolist() == ["DUP", "DUP"]
    assert window.results_tabs.currentWidget() is panel
    assert not panel.model.flags(panel.model.index(0, 1)) & Qt.ItemFlag.ItemIsEditable
    pd.testing.assert_frame_equal(window._preview.dataframe, original)
    assert hashlib.sha256(path.read_bytes()).digest() == source_digest


def test_show_unchanged_rows_and_preserve_numeric_text_distinction(window):
    analyze(window)
    generate(window)
    panel = window.cleaning_preview_panel
    panel.show_unchanged_box.setChecked(True)
    assert panel.model.rowCount() == 14
    assert panel.model.data(panel.model.index(7, 1)) == "0"
    assert panel.model.data(panel.model.index(12, 1)) == '"0"'
    assert panel.model.data(panel.model.index(8, 4)) == "否"
    assert panel.model.data(panel.model.index(8, 3)) == "无"
    panel.show_unchanged_box.setChecked(False)
    assert panel.model.rowCount() == 6
    assert panel.result.changed_count == 6


def test_rule_changes_clear_old_preview_and_apply_only_selected_rules(window):
    analyze(window)
    generate(window)
    for box in window.cleaning_panel.rule_boxes.values():
        box.setChecked(False)
    assert window.cleaning_preview_panel.result is None
    assert window.cleaning_panel.formula_edit.toPlainText() == "=A2"
    window.cleaning_panel.rule_boxes["nbsp"].setChecked(True)
    assert window.cleaning_panel.formula_edit.toPlainText() == '=SUBSTITUTE(A2,UNICHAR(160)," ")'
    generate(window)
    result = window.cleaning_preview_panel.result
    assert result.changed_count == 2
    assert result.records[1].cleaned_value == "P002 "
    assert result.records[2].cleaned_value == "ABC\t123"
    assert result.processed_counts == {"前后空格": 0, "CHAR(160)": 2, "Tab / 换行": 0}


def test_explicit_delete_strategy_changes_hidden_rules_and_formula(window):
    analyze(window)
    window.cleaning_panel.strategy_combo.setCurrentIndex(1)
    assert window.cleaning_panel.rule_boxes["tab"].text() == "删除 Tab"
    assert 'CHAR(9),""' in window.cleaning_panel.formula_edit.toPlainText()
    assert 'UNICHAR(160)," "' in window.cleaning_panel.formula_edit.toPlainText()
    generate(window)
    result = window.cleaning_preview_panel.result
    assert result.records[2].cleaned_value == "ABC123"
    assert result.records[4].cleaned_value == "AB"
    assert result.records[9].cleaned_value == "P003"


def test_repeated_preview_always_starts_from_original_column(window):
    analyze(window)
    window.cleaning_panel.strategy_combo.setCurrentIndex(1)
    generate(window)
    assert window.cleaning_preview_panel.result.records[2].cleaned_value == "ABC123"
    window.cleaning_panel.strategy_combo.setCurrentIndex(0)
    assert window.cleaning_preview_panel.result is None
    generate(window)
    assert window.cleaning_preview_panel.result.records[2].cleaned_value == "ABC 123"
    assert window._preview.dataframe.iat[2, 0] == "ABC\t123"


def test_column_change_resets_preview_and_updates_formula_reference(window):
    analyze(window)
    generate(window)
    window._select_column(0)
    assert window.cleaning_preview_panel.result is not None
    window._select_column(1)
    assert window.cleaning_preview_panel.result is None
    assert not window.cleaning_panel.generate_button.isEnabled()
    analyze(window, column=1)
    assert window.cleaning_panel.formula_edit.toPlainText() == "=B2"
    window.cleaning_panel.rule_boxes["outer_spaces"].setChecked(True)
    assert window.cleaning_panel.formula_edit.toPlainText() == "=TRIM(B2)"
    generate(window)
    assert window.cleaning_preview_panel.result.changed_count == 0


def test_sheet_and_new_file_clear_results(window, tmp_path):
    analyze(window)
    generate(window)
    window.sheet_combo.setCurrentText("正常表")
    wait_for_idle(window)
    assert window.results_tabs.currentWidget() is window.raw_data_page
    assert window.cleaning_preview_panel.result is None
    assert window.cleaning_panel.analysis is None
    analyze(window)
    generate(window)
    assert window.cleaning_preview_panel.result.changed_count == 0
    assert window.cleaning_preview_panel.model.rowCount() == 0
    window.cleaning_preview_panel.show_unchanged_box.setChecked(True)
    assert window.cleaning_preview_panel.model.rowCount() == 2
    path = tmp_path / "另一份.xlsx"
    workbook = Workbook()
    workbook.active.append(["新字段"])
    workbook.active.append([" P001 "])
    workbook.save(path)
    workbook.close()
    window._start_load(str(path))
    wait_for_idle(window)
    assert window.results_tabs.currentWidget() is window.raw_data_page
    assert window.cleaning_preview_panel.result is None
    assert not window.cleaning_panel.generate_button.isEnabled()


def test_failed_sheet_load_preserves_existing_preview(window):
    analyze(window)
    generate(window)
    previous = window.cleaning_preview_panel.result
    window._preview.file_path.unlink()
    with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok):
        window.sheet_combo.setCurrentText("正常表")
        wait_for_idle(window)
    assert window.cleaning_preview_panel.result is previous
    assert window.results_tabs.currentWidget() is window.cleaning_preview_panel
    assert window.cleaning_panel.generate_button.isEnabled()


def test_reanalysis_clears_old_preview_and_failure_blocks_generation(window):
    analyze(window)
    generate(window)
    with patch("app.ui.analysis_worker.analyze_column", side_effect=RuntimeError("测试异常")):
        with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok):
            window._start_analysis()
            wait_for_idle(window)
    assert window.cleaning_preview_panel.result is None
    assert not window.cleaning_panel.generate_button.isEnabled()


def test_cleaning_worker_error_restores_controls_and_keeps_data(window):
    analyze(window)
    original = window._preview.dataframe.copy(deep=True)
    with patch("app.ui.cleaning_worker.generate_cleaning_preview", side_effect=RuntimeError("测试异常")):
        with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok) as dialog:
            generate(window)
    dialog.assert_called_once()
    assert window.cleaning_preview_panel.result is None
    assert window.cleaning_panel.generate_button.isEnabled()
    assert window.table.isEnabled()
    pd.testing.assert_frame_equal(window._preview.dataframe, original)


def test_header_only_sheet_and_close_during_generation(window):
    window.sheet_combo.setCurrentText("仅字段")
    wait_for_idle(window)
    analyze(window)
    generate(window)
    assert window.cleaning_preview_panel.result.row_count == 0
    window._start_cleaning_preview()
    window.close()
    assert window._close_requested
    wait_for_idle(window)
    assert not window.isVisible()


def test_page_navigation_keeps_column_diagnostics_and_preview(window):
    assert window.results_tabs.currentWidget() is window.raw_data_page
    assert window.cleaning_scroll.isHidden()
    header = window.table.horizontalHeader()
    position = header.viewport().rect().center()
    position.setX(header.sectionViewportPosition(0) + header.sectionSize(0) // 2)
    QTest.mouseClick(header.viewport(), Qt.MouseButton.LeftButton, pos=position)
    assert window.selected_column == 0
    assert window.results_tabs.currentWidget() is window.raw_data_page
    assert window.start_button.isVisible()
    QTest.mouseClick(window.start_button, Qt.MouseButton.LeftButton)
    wait_for_idle(window)
    assert window.results_tabs.currentWidget() is window.diagnosis_panel
    assert window.cleaning_scroll.isHidden()
    diagnosis = window.diagnosis_panel.result
    generate(window)
    preview = window.cleaning_preview_panel.result
    click_page(window, window.raw_data_page)
    assert window.table.isVisible()
    assert window.selected_column == 0
    assert len(window.table.selectionModel().selectedColumns()) == 1
    click_page(window, window.diagnosis_panel)
    assert window.diagnosis_panel.result is diagnosis
    click_page(window, window.cleaning_preview_panel)
    assert window.cleaning_preview_panel.result is preview
    assert not window.cleaning_scroll.isHidden()


def test_sidebar_collapse_remembers_each_page_and_preserves_rules(window):
    analyze(window)
    generate(window)
    rules = window.cleaning_panel.current_rules()
    preview = window.cleaning_preview_panel.result
    narrow_width = window.cleaning_preview_panel.table.width()
    QTest.mouseClick(window.cleaning_toggle_button, Qt.MouseButton.LeftButton)
    QApplication.processEvents()
    assert window.cleaning_scroll.isHidden()
    assert window.cleaning_preview_panel.table.width() > narrow_width
    click_page(window, window.raw_data_page)
    QTest.mouseClick(window.cleaning_toggle_button, Qt.MouseButton.LeftButton)
    assert not window.cleaning_scroll.isHidden()
    click_page(window, window.cleaning_preview_panel)
    assert window.cleaning_scroll.isHidden()
    click_page(window, window.raw_data_page)
    assert not window.cleaning_scroll.isHidden()
    click_page(window, window.cleaning_preview_panel)
    QTest.mouseClick(window.cleaning_toggle_button, Qt.MouseButton.LeftButton)
    assert not window.cleaning_scroll.isHidden()
    assert window.cleaning_panel.current_rules() == rules
    assert window.cleaning_preview_panel.result is preview


def test_fit_columns_applies_to_current_table(window):
    analyze(window)
    window.table.setColumnWidth(0, 1800)
    window.diagnosis_panel.details_table.setColumnWidth(1, 60)
    QTest.mouseClick(window.fit_columns_button, Qt.MouseButton.LeftButton)
    assert window.diagnosis_panel.details_table.columnWidth(1) > 60
    assert window.table.columnWidth(0) == 1800
    generate(window)
    window.cleaning_preview_panel.table.setColumnWidth(1, 60)
    QTest.mouseClick(window.fit_columns_button, Qt.MouseButton.LeftButton)
    assert window.cleaning_preview_panel.table.columnWidth(1) > 60
    assert window.table.columnWidth(0) == 1800
    click_page(window, window.raw_data_page)
    QTest.mouseClick(window.fit_columns_button, Qt.MouseButton.LeftButton)
    assert window.table.columnWidth(0) < 1800
