"""使用真实临时工作簿验证文件读取及 Qt 交互。"""

import hashlib
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QAbstractItemView, QFileDialog, QMessageBox

from app.core.file_loader import ExcelFileLoader, ExcelLoadError
from app.ui.main_window import MainWindow


def create_workbook(path: Path) -> None:
    workbook = Workbook()
    products = workbook.active
    products.title = "商品"
    products.append(["商品编码", "商品名称", "数量"])
    products.append(["00123", "苹果", 5])
    products.append(["NA", " 含空格 ", None])
    orders = workbook.create_sheet("订单")
    orders.append(["订单号", "金额"])
    for number in range(150):
        orders.append([f"D{number:03}", number * 10])
    workbook.create_sheet("空表")
    workbook.create_sheet("仅字段").append(["商品编码", "备注"])
    workbook.save(path)
    workbook.close()


class FileLoaderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / "测试数据.xlsx"
        create_workbook(self.path)
        self.loader = ExcelFileLoader()

    def tearDown(self) -> None:
        self.directory.cleanup()

    def test_default_sheet_preserves_text_and_source_file(self) -> None:
        original_digest = hashlib.sha256(self.path.read_bytes()).digest()
        preview = self.loader.load(self.path)
        self.assertEqual(preview.sheet_names, ("商品", "订单", "空表", "仅字段"))
        self.assertEqual(preview.sheet_name, "商品")
        self.assertEqual(preview.dataframe.shape, (2, 3))
        self.assertEqual(list(preview.dataframe.columns), ["商品编码", "商品名称", "数量"])
        self.assertEqual(preview.dataframe.iat[0, 0], "00123")
        self.assertEqual(preview.dataframe.iat[1, 0], "NA")
        self.assertEqual(preview.dataframe.iat[1, 1], " 含空格 ")
        self.assertEqual(preview.dataframe.iat[1, 2], "")
        self.assertEqual(original_digest, hashlib.sha256(self.path.read_bytes()).digest())
        # Windows 下重命名成功也能验证读取完成后已释放文件句柄。
        renamed_path = self.path.with_name("已释放.xlsx")
        self.path.rename(renamed_path)
        renamed_path.rename(self.path)

    def test_switch_sheet(self) -> None:
        preview = self.loader.load(self.path, "订单")
        self.assertEqual(preview.sheet_name, "订单")
        self.assertEqual(preview.dataframe.shape, (150, 2))
        self.assertEqual(preview.dataframe.iat[149, 0], "D149")

    def test_empty_and_header_only_sheets(self) -> None:
        self.assertEqual(self.loader.load(self.path, "空表").dataframe.shape, (0, 0))
        self.assertEqual(self.loader.load(self.path, "仅字段").dataframe.shape, (0, 2))

    def test_invalid_file_and_missing_sheet_have_friendly_errors(self) -> None:
        invalid_path = self.path.with_name("损坏.xlsx")
        invalid_path.write_text("这不是工作簿", encoding="utf-8")
        cases = (
            (invalid_path, None, "不是有效"),
            (self.path.with_name("不存在.xlsx"), None, "不存在"),
            (self.path.with_suffix(".xls"), None, "仅支持"),
            (self.path, "不存在的表", "工作表已不存在"),
        )
        for file_path, sheet_name, expected_message in cases:
            with self.subTest(file=file_path, sheet=sheet_name):
                with self.assertRaisesRegex(ExcelLoadError, expected_message):
                    self.loader.load(file_path, sheet_name)

    def test_permission_error_is_translated(self) -> None:
        with patch("app.core.file_loader.pd.ExcelFile", side_effect=PermissionError("denied")):
            with self.assertRaisesRegex(ExcelLoadError, "无法访问") as caught:
                self.loader.load(self.path)
        self.assertEqual(caught.exception.detail, "denied")


class MainWindowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.application = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.path = Path(self.directory.name) / "测试数据.xlsx"
        create_workbook(self.path)
        self.window = MainWindow()
        self.window.show()
        self.application.processEvents()

    def tearDown(self) -> None:
        self.wait_for_load()
        self.window.close()
        self.window.deleteLater()
        self.application.processEvents()
        self.directory.cleanup()

    def wait_for_load(self) -> None:
        deadline = time.monotonic() + 10
        while self.window._worker is not None:
            self.application.processEvents()
            QTest.qWait(10)
            if time.monotonic() > deadline:
                self.fail("后台读取未在 10 秒内结束")
        self.application.processEvents()

    def open_workbook(self) -> None:
        with patch.object(QFileDialog, "getOpenFileName", return_value=(str(self.path), "")):
            QTest.mouseClick(self.window.open_button, Qt.MouseButton.LeftButton)
        self.assertFalse(self.window.open_button.isEnabled())
        self.wait_for_load()

    def test_initial_state_and_cancel(self) -> None:
        self.assertEqual(self.window.windowTitle(), "Excel 数据清洗助手")
        self.assertEqual((self.window.width(), self.window.height()), (1200, 750))
        self.assertEqual(self.window.selected_column_label.text(), "当前选择列：未选择")
        self.assertFalse(self.window.sheet_combo.isEnabled())
        self.assertEqual(self.window.row_count_label.text(), "总行数：0")
        with patch.object(QFileDialog, "getOpenFileName", return_value=("", "")):
            QTest.mouseClick(self.window.open_button, Qt.MouseButton.LeftButton)
        self.assertIsNone(self.window._preview)
        self.assertIsNone(self.window._worker)

    def test_import_preview_and_read_only(self) -> None:
        self.open_workbook()
        model = self.window.table.model()
        self.assertEqual((model.rowCount(), model.columnCount()), (2, 3))
        self.assertEqual(self.window.sheet_combo.count(), 4)
        self.assertEqual(model.headerData(0, Qt.Orientation.Horizontal), "商品编码")
        self.assertEqual(model.headerData(0, Qt.Orientation.Vertical), "1")
        index = model.index(0, 0)
        self.assertEqual(model.data(index), "00123")
        self.assertFalse(model.flags(index) & Qt.ItemFlag.ItemIsEditable)
        self.assertFalse(model.setData(index, "更改"))
        self.assertEqual(model.data(index), "00123")
        self.assertEqual(self.window.table.editTriggers(), QAbstractItemView.EditTrigger.NoEditTriggers)
        self.assertEqual(self.window.file_label.text(), "当前文件名：测试数据.xlsx")
        self.assertEqual(self.window.sheet_label.text(), "当前工作表：商品")
        self.assertEqual(self.window.row_count_label.text(), "总行数：2")
        self.assertEqual(self.window.column_count_label.text(), "总列数：3")

    def click_header(self, column: int) -> None:
        header = self.window.table.horizontalHeader()
        position = header.viewport().rect().center()
        position.setX(header.sectionViewportPosition(column) + header.sectionSize(column) // 2)
        QTest.mouseClick(header.viewport(), Qt.MouseButton.LeftButton, pos=position)

    def test_header_selection_and_sheet_switch_reset(self) -> None:
        self.open_workbook()
        self.click_header(0)
        self.assertEqual(self.window.selected_column, 0)
        self.assertEqual(self.window.selected_column_label.text(), "当前选择列：商品编码")
        self.assertEqual(len(self.window.table.selectionModel().selectedColumns()), 1)
        self.window.sheet_combo.setCurrentText("订单")
        self.wait_for_load()
        self.assertIsNone(self.window.selected_column)
        self.assertEqual(self.window.selected_column_label.text(), "当前选择列：未选择")
        self.assertEqual(self.window.sheet_label.text(), "当前工作表：订单")
        self.assertEqual(self.window.row_count_label.text(), "总行数：150")
        self.assertGreater(self.window.table.verticalScrollBar().maximum(), 0)
        self.window.table.verticalScrollBar().setValue(self.window.table.verticalScrollBar().maximum())
        self.assertEqual(self.window.table.model().data(self.window.table.model().index(149, 0)), "D149")

    def test_empty_and_header_only_preview(self) -> None:
        self.open_workbook()
        self.window.sheet_combo.setCurrentText("空表")
        self.wait_for_load()
        self.assertEqual(self.window.row_count_label.text(), "总行数：0")
        self.assertEqual(self.window.column_count_label.text(), "总列数：0")
        self.window.sheet_combo.setCurrentText("仅字段")
        self.wait_for_load()
        self.click_header(0)
        self.assertEqual(self.window.selected_column_label.text(), "当前选择列：商品编码")
        self.assertEqual(self.window.column_count_label.text(), "总列数：2")

    def test_failed_switch_restores_previous_preview(self) -> None:
        self.open_workbook()
        self.click_header(0)
        self.path.unlink()
        with patch.object(QMessageBox, "exec", return_value=QMessageBox.StandardButton.Ok) as dialog:
            self.window.sheet_combo.setCurrentText("订单")
            self.wait_for_load()
        dialog.assert_called_once()
        self.assertEqual(self.window.sheet_combo.currentText(), "商品")
        self.assertEqual(self.window.table.model().rowCount(), 2)
        self.assertEqual(self.window.selected_column_label.text(), "当前选择列：商品编码")
        self.assertTrue(self.window.open_button.isEnabled())
        self.assertTrue(self.window.sheet_combo.isEnabled())

    def test_new_file_clears_selection_and_updates_metadata(self) -> None:
        self.open_workbook()
        self.click_header(1)
        second_path = self.path.with_name("第二份.xlsx")
        workbook = Workbook()
        workbook.active.title = "新表"
        workbook.active.append(["字段"])
        workbook.active.append(["值"])
        workbook.save(second_path)
        workbook.close()
        self.window._start_load(str(second_path))
        self.wait_for_load()
        self.assertIsNone(self.window.selected_column)
        self.assertEqual(self.window.file_label.text(), "当前文件名：第二份.xlsx")
        self.assertEqual(self.window.sheet_combo.count(), 1)
        self.assertEqual(self.window.sheet_label.text(), "当前工作表：新表")

    def test_close_during_load_waits_for_worker(self) -> None:
        self.window._start_load(str(self.path))
        self.window.close()
        self.assertTrue(self.window._close_requested)
        self.wait_for_load()
        self.assertFalse(self.window.isVisible())


if __name__ == "__main__":
    unittest.main()
