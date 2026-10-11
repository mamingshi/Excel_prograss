"""主窗口：文件预览、整表扫描、列诊断和清洗预览的任务调度。"""

from dataclasses import replace
from pathlib import Path
from openpyxl.utils.cell import get_column_letter

from PySide6.QtCore import QItemSelectionModel, QModelIndex, QSignalBlocker, QSize, Qt, Slot
from PySide6.QtGui import QCloseEvent
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QComboBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSplitter,
    QStackedWidget,
    QTabWidget,
    QTableView,
    QVBoxLayout,
    QWidget,
)

from app.core.file_loader import ExcelPreview
from app.core.batch_cleaner import BatchCleaningPreview, recommend_batch_configs
from app.core.cleaner import CleaningPreview
from app.core.annotation_builder import AnnotationPreview
from app.core.detector import ColumnAnalysis
from app.core.sheet_analyzer import SheetAnalysis
from app.core.data_manager import DataStateError, PreparedAnnotationPreview, PreparedBatchPreview, PreparedPreview, SheetScanSnapshot, WorkbookDataManager
from app.core.exporter import (
    ExcelExportError, NO_CHANGES_MESSAGE, prepare_export, suggest_export_path, validate_export_paths,
)
from app.models.cleaning_action import AnnotationAction, BatchCleaningAction, CleaningAction
from app.models.dataframe_model import DataFrameModel
from app.models.export_data import ExportResult
from app.ui.cell_preview_dialog import CellPreviewDialog
from app.ui.analysis_worker import ColumnAnalysisWorker
from app.ui.sheet_analysis_worker import SheetAnalysisWorker
from app.ui.sheet_overview_panel import SheetOverviewPanel
from app.ui.diagnosis_panel import DiagnosisPanel
from app.ui.cleaning_panel import CleaningPanel
from app.ui.cleaning_preview_panel import CleaningPreviewPanel
from app.ui.cleaning_worker import CleaningPreviewWorker
from app.ui.batch_worker import BatchPreviewWorker
from app.ui.batch_cleaning_panel import BatchCleaningPanel
from app.ui.annotation_worker import AnnotationPreviewWorker
from app.ui.annotation_preview_panel import AnnotationPreviewPanel
from app.ui.load_worker import ExcelLoadWorker
from app.ui.export_worker import ExcelExportWorker
from app.ui.confirmation_dialogs import confirm_annotation, confirm_apply, confirm_batch_apply, confirm_restore, confirm_export, show_export_success
from app.ui.theme import apply_theme, icon, primary_button, style_table


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        apply_theme(QApplication.instance())
        self.setWindowTitle("Excel 数据清洗助手")
        self.setWindowIcon(icon("app"))
        self.resize(1200, 750)
        self._preview: ExcelPreview | None = None
        self.data_manager = WorkbookDataManager()
        self._prepared_preview: PreparedPreview | None = None
        self._prepared_annotation: PreparedAnnotationPreview | None = None
        self._annotation_worker: AnnotationPreviewWorker | None = None
        self._loading_new_workbook = False
        self._analysis_completion_message: str | None = None
        self._worker: ExcelLoadWorker | None = None
        self._analysis_worker: ColumnAnalysisWorker | None = None
        self._batch_worker: BatchPreviewWorker | None = None
        self._batch_snapshot: SheetScanSnapshot | None = None
        self._prepared_batch: PreparedBatchPreview | None = None
        self._batch_completion_message: str | None = None
        self._scan_for_batch = False
        self._scan_worker: SheetAnalysisWorker | None = None
        self._scan_snapshot: SheetScanSnapshot | None = None
        self._scan_completion_message: str | None = None
        self._scan_cancel_requested = False
        self._cleaning_worker: CleaningPreviewWorker | None = None
        self._export_worker: ExcelExportWorker | None = None
        self._export_completion_message: str | None = None
        self._close_requested = False
        self._cleaning_sidebar_width = 325
        self.selected_column: int | None = None
        self._model = DataFrameModel(self)
        self._build_ui()
        self._connect_signals()
        self._on_page_changed(self.results_tabs.currentIndex())

    def _build_ui(self) -> None:
        central_widget = QWidget(self)
        central_widget.setObjectName("workspaceRoot")
        layout = QVBoxLayout(central_widget)
        layout.setContentsMargins(12, 10, 12, 10)
        layout.setSpacing(8)

        # 标题、工具栏和工作区分别布局，避免窗口变高时出现大块空白。
        heading = QHBoxLayout()
        logo = QLabel()
        logo.setPixmap(icon("app").pixmap(QSize(32, 32)))
        logo.setFixedSize(32, 32)
        heading.addWidget(logo)
        heading.addSpacing(2)
        title = QLabel("Excel 数据清洗助手")
        title.setObjectName("appTitle")
        heading.addWidget(title)
        heading.addStretch()
        self.scan_button = QPushButton("扫描当前工作表")
        self.scan_button.setIcon(icon("scan"))
        self.scan_button.setToolTip("只读分析当前工作副本的全部字段，查看数据质量总览。")
        heading.addWidget(self.scan_button)
        self.undo_button = QPushButton("撤销")
        self.undo_button.setIcon(icon("undo"))
        self.undo_button.setToolTip("撤销本工作簿最近一次操作；跨 Sheet 操作会自动切换工作表。")
        self.restore_button = QPushButton("恢复原始")
        self.restore_button.setIcon(icon("restore"))
        self.restore_button.setToolTip("将当前工作表的工作副本恢复为原始快照。")
        heading.addWidget(self.undo_button)
        heading.addWidget(self.restore_button)
        self.export_button = QPushButton("导出清洗结果")
        primary_button(self.export_button, "export")
        self.export_button.setToolTip("选择原列更新或新增公式对照列，将所有 Sheet 的已应用清洗及辅助列另存为新 .xlsx 文件。")
        heading.addWidget(self.export_button)
        for text, tone in (("V0.9", "neutral"), ("另存为新文件", "success")):
            badge = QLabel(text)
            badge.setProperty("role", "badge")
            badge.setProperty("tone", tone)
            heading.addWidget(badge, 0, Qt.AlignmentFlag.AlignVCenter)
        layout.addLayout(heading)

        toolbar_card = QFrame()
        toolbar_card.setProperty("role", "card")
        toolbar_card.setSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Fixed)
        toolbar = QHBoxLayout(toolbar_card)
        toolbar.setContentsMargins(12, 10, 12, 10)
        toolbar.setSpacing(10)

        self.open_button = QPushButton("打开 Excel")
        primary_button(self.open_button, "folder")
        self.fit_columns_button = QPushButton("适应列宽")
        self.fit_columns_button.setIcon(icon("fit"))
        self.fit_columns_button.setEnabled(False)
        self.fit_columns_button.setToolTip("调整当前页面的表格列宽，超出窗口时可横向滚动。")
        self.diagnosis_panel = DiagnosisPanel(self)
        # 复用诊断组件的检测动作，将入口放到所有页面都能访问的工具栏。
        self.start_button = self.diagnosis_panel.start_button
        self.sheet_combo = QComboBox()
        self.sheet_combo.setMinimumWidth(180)
        self.sheet_combo.setMaximumWidth(260)
        self.sheet_combo.setEnabled(False)
        toolbar.addWidget(self.open_button)
        toolbar.addWidget(self.fit_columns_button)
        toolbar.addWidget(self.start_button)
        toolbar.addSpacing(6)
        toolbar.addWidget(QLabel("工作表："))
        toolbar.addWidget(self.sheet_combo)
        toolbar.addStretch()
        self.selected_column_label = QLabel("当前选择列：未选择")
        self.selected_column_label.setTextFormat(Qt.TextFormat.PlainText)
        self.selected_column_label.setObjectName("selectedColumn")
        self.selected_column_label.setProperty("role", "badge")
        self.selected_column_label.setWordWrap(True)
        self.selected_column_label.setMinimumWidth(200)
        self.selected_column_label.setMaximumWidth(300)
        toolbar.addWidget(self.selected_column_label, 0, Qt.AlignmentFlag.AlignVCenter)
        self.cleaning_toggle_button = QPushButton("展开清洗设置 ›")
        self.cleaning_toggle_button.setCheckable(True)
        self.cleaning_toggle_button.setToolTip("展开或收起推荐规则和公式面板，为表格留出更多空间。")
        toolbar.addWidget(self.cleaning_toggle_button)
        layout.addWidget(toolbar_card)

        self.table = QTableView()
        style_table(self.table)
        self.table.setModel(self._model)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectColumns)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setAlternatingRowColors(True)
        self.table.setWordWrap(False)
        # 按像素滚动，单列比窗口更宽时也能查看该列右侧的文本。
        self.table.setHorizontalScrollMode(QAbstractItemView.ScrollMode.ScrollPerPixel)
        self.table.verticalHeader().setVisible(True)
        header = self.table.horizontalHeader()
        header.setSectionsClickable(True)
        header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        header.setStretchLastSection(False)
        header.setDefaultSectionSize(150)
        header.setMinimumSectionSize(60)
        # 保持 Excel 原始顺序，后续如需排序应独立设计。
        self.table.setSortingEnabled(False)

        preview_card = QFrame()
        self.raw_data_page = preview_card
        preview_card.setObjectName("rawDataPage")
        preview_card.setProperty("role", "card")
        preview_layout = QVBoxLayout(preview_card)
        preview_layout.setContentsMargins(12, 10, 12, 12)
        preview_layout.setSpacing(8)
        preview_heading = QHBoxLayout()
        preview_title = QLabel("工作副本 · 点击表头选择待检测列")
        preview_title.setProperty("role", "sectionTitle")
        preview_heading.addWidget(preview_title)
        preview_heading.addStretch()
        preview_hint = QLabel("拖动表头调宽 · 双击单元格查看全文")
        preview_hint.setProperty("role", "muted")
        preview_hint.setToolTip("拖动表头分隔线调整列宽；双击表头适应列宽；双击单元格查看完整内容。")
        preview_heading.addWidget(preview_hint)
        preview_layout.addLayout(preview_heading)
        self.preview_stack = QStackedWidget()
        empty_page = QWidget()
        empty_layout = QVBoxLayout(empty_page)
        empty_layout.setSpacing(8)
        empty_layout.addStretch()
        empty_title = QLabel("从一份 Excel 开始")
        empty_title.setProperty("role", "sectionTitle")
        empty_title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        empty_layout.addWidget(empty_title)
        empty_hint = QLabel("打开 .xlsx 文件，选择列后即可检测数据质量")
        empty_hint.setProperty("role", "muted")
        empty_hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        empty_layout.addWidget(empty_hint)
        self.empty_open_button = QPushButton("打开 Excel")
        primary_button(self.empty_open_button, "folder")
        empty_layout.addWidget(self.empty_open_button, 0, Qt.AlignmentFlag.AlignHCenter)
        empty_layout.addStretch()
        self.preview_stack.addWidget(empty_page)
        self.preview_stack.addWidget(self.table)
        preview_layout.addWidget(self.preview_stack, 1)

        self.cleaning_preview_panel = CleaningPreviewPanel(self)
        self.results_tabs = QTabWidget()
        self.sheet_overview_panel = SheetOverviewPanel(self)
        self.results_tabs.addTab(self.raw_data_page, "工作数据")
        self.results_tabs.addTab(self.sheet_overview_panel, "数据质量总览")
        self.results_tabs.addTab(self.diagnosis_panel, "数据诊断")
        self.results_tabs.addTab(self.cleaning_preview_panel, "清洗预览")
        self.annotation_preview_panel = AnnotationPreviewPanel(self)
        self.results_tabs.addTab(self.annotation_preview_panel, "辅助列预览")
        self.batch_panel = BatchCleaningPanel(self)
        self.results_tabs.addTab(self.batch_panel, "批量清洗")
        # 各张主表共享整块工作区，切换页面只改变展示，不重置数据或结果。
        self._sidebar_preferences = {
            self.raw_data_page: False,
            self.sheet_overview_panel: False,
            self.diagnosis_panel: False,
            self.cleaning_preview_panel: True,
            self.annotation_preview_panel: False,
            self.batch_panel: False,
        }
        self.cleaning_panel = CleaningPanel(self)
        # 侧栏可滚动，小窗口下规则和公式仍然能完整访问。
        self.cleaning_scroll = QScrollArea()
        self.cleaning_scroll.setObjectName("cleaningScroll")
        self.cleaning_scroll.setWidgetResizable(True)
        self.cleaning_scroll.viewport().setObjectName("cleaningViewport")
        self.cleaning_scroll.setFrameShape(QFrame.Shape.NoFrame)
        self.cleaning_scroll.setMinimumWidth(315)
        self.cleaning_scroll.setWidget(self.cleaning_panel)
        self.workspace_splitter = QSplitter(Qt.Orientation.Horizontal)
        self.workspace_splitter.setHandleWidth(12)
        self.workspace_splitter.addWidget(self.results_tabs)
        self.workspace_splitter.addWidget(self.cleaning_scroll)
        self.workspace_splitter.setChildrenCollapsible(False)
        self.workspace_splitter.setStretchFactor(0, 1)
        self.workspace_splitter.setStretchFactor(1, 0)
        self.workspace_splitter.setSizes([835, 325])
        self.cleaning_scroll.hide()
        layout.addWidget(self.workspace_splitter, 1)
        self.setCentralWidget(central_widget)

        self.file_label = QLabel("当前文件名：未打开")
        self.sheet_label = QLabel("当前工作表：未选择")
        self.row_count_label = QLabel("总行数：0")
        self.column_count_label = QLabel("总列数：0")
        self.modification_label = QLabel("当前状态：未修改")
        # 临时操作提示占用左侧，计数及修改状态始终保留在右侧。
        self.statusBar().addWidget(self.file_label, 1)
        self.statusBar().addWidget(self.sheet_label, 1)
        self.statusBar().addPermanentWidget(self.row_count_label)
        self.statusBar().addPermanentWidget(self.column_count_label)
        self.statusBar().addPermanentWidget(self.modification_label)

    def _connect_signals(self) -> None:
        self.batch_panel.scan_requested.connect(self._request_batch_scan)
        self.batch_panel.preview_requested.connect(self._start_batch_preview)
        self.batch_panel.apply_requested.connect(self._apply_batch_cleaning)
        self.batch_panel.cancel_requested.connect(self._cancel_batch_task)
        self.batch_panel.configuration_changed.connect(self._clear_batch_preview)
        self.batch_panel.duplicate_requested.connect(self._request_batch_duplicate)
        self.sheet_overview_panel.batch_requested.connect(self._open_batch_workspace)
        self.open_button.clicked.connect(self._open_excel)
        self.scan_button.clicked.connect(self._start_sheet_scan)
        self.sheet_overview_panel.cancel_requested.connect(self._cancel_sheet_scan)
        self.sheet_overview_panel.column_requested.connect(self._select_overview_column)
        self.empty_open_button.clicked.connect(self._open_excel)
        self.fit_columns_button.clicked.connect(self._fit_columns)
        self.cleaning_toggle_button.toggled.connect(self._set_cleaning_sidebar_visible)
        self.results_tabs.currentChanged.connect(self._on_page_changed)
        self.undo_button.clicked.connect(self._undo_last_action)
        self.restore_button.clicked.connect(self._restore_original)
        self.export_button.clicked.connect(self._export_results)
        self.sheet_combo.currentTextChanged.connect(self._switch_sheet)
        self.table.horizontalHeader().sectionClicked.connect(self._select_column)
        self.table.horizontalHeader().sectionDoubleClicked.connect(self.table.resizeColumnToContents)
        self.table.doubleClicked.connect(self._show_cell_preview)
        self.diagnosis_panel.analysis_requested.connect(self._start_analysis)
        self.cleaning_panel.preview_requested.connect(self._start_cleaning_preview)
        self.cleaning_panel.annotation_requested.connect(self._start_annotation_preview)
        self.cleaning_panel.type_annotation_requested.connect(lambda kind: self._start_annotation_preview(kind))
        self.cleaning_panel.annotation_invalidated.connect(self._clear_annotation_preview)
        self.annotation_preview_panel.apply_requested.connect(self._apply_annotation)
        self.cleaning_panel.preview_invalidated.connect(self._clear_cleaning_preview)
        self.cleaning_preview_panel.apply_requested.connect(self._apply_cleaning)

    def _is_busy(self) -> bool:
        return any(worker is not None for worker in (
            self._worker, self._analysis_worker, self._cleaning_worker, self._export_worker,
            self._scan_worker, self._annotation_worker, self._batch_worker,
        ))

    def _update_controls(self) -> None:
        busy = self._is_busy()
        self.open_button.setEnabled(not busy)
        self.scan_button.setEnabled(not busy and self._preview is not None)
        self.sheet_overview_panel.set_busy(busy, scanning=self._scan_worker is not None)
        self.empty_open_button.setEnabled(not busy)
        self.sheet_combo.setEnabled(not busy and self._preview is not None)
        self.fit_columns_button.setEnabled(
            not busy and self._preview is not None
            and self._current_table().model().columnCount() > 0
        )
        # 后台任务期间固定数据和当前列，防止结果返回到另一份数据上。
        self.table.setEnabled(not busy)
        self.batch_panel.set_busy(
            busy, task_active=self._batch_worker is not None
            or (self._scan_for_batch and self._scan_worker is not None),
        )
        self.batch_panel.apply_button.setEnabled(not busy and self._can_apply_batch())
        self.cleaning_toggle_button.setEnabled(self.results_tabs.currentWidget() is not self.batch_panel)
        self.diagnosis_panel.set_busy(busy)
        self.cleaning_panel.set_busy(busy)
        self.cleaning_preview_panel.apply_button.setEnabled(not busy and self._can_apply_preview())
        self.annotation_preview_panel.apply_button.setEnabled(not busy and self._can_apply_annotation())
        self.undo_button.setEnabled(not busy and bool(self.data_manager.undo_stack))
        self.export_button.setEnabled(not busy and self._preview is not None and self.data_manager.is_modified)
        self.restore_button.setEnabled(
            not busy and self.data_manager.current_sheet is not None
            and self.data_manager.current_data.is_modified
        )

    def _can_apply_preview(self) -> bool:
        preview = self._prepared_preview
        return (
            preview is not None and self.selected_column == preview.column_index
            and self.cleaning_preview_panel.result is preview.result
            and preview.result.changed_count > 0
            and preview.result.rules == self.cleaning_panel.current_rules()
            and self.data_manager.is_preview_current(preview)
        )

    @Slot()
    def _clear_cleaning_preview(self) -> None:
        self._prepared_preview = None
        self.cleaning_preview_panel.clear_preview()
        self._update_controls()

    def _update_modification_state(self) -> None:
        state = self.data_manager.current_data if self.data_manager.current_sheet is not None else None
        text = f"已修改（{state.operation_count} 次操作）" if state and state.is_modified else "未修改"
        self.modification_label.setText("当前状态：" + text)
        self.modification_label.setToolTip("底部状态对应当前工作表；标题 * 表示工作簿中至少一个 Sheet 有内存修改。")
        self.setWindowTitle("Excel 数据清洗助手" + (" *" if self.data_manager.is_modified else ""))
        self._update_controls()

    def _current_table(self) -> QTableView:
        page = self.results_tabs.currentWidget()
        if page is self.batch_panel:
            return self.batch_panel.current_table
        if page is self.sheet_overview_panel:
            return self.sheet_overview_panel.table
        if page is self.diagnosis_panel:
            return self.diagnosis_panel.details_table
        if page is self.annotation_preview_panel:
            return self.annotation_preview_panel.table
        if page is self.cleaning_preview_panel:
            return self.cleaning_preview_panel.table
        return self.table

    @Slot(int)
    def _on_page_changed(self, _index: int) -> None:
        page = self.results_tabs.currentWidget()
        self._set_cleaning_sidebar_visible(self._sidebar_preferences[page])
        self._update_controls()

    @Slot(bool)
    def _set_cleaning_sidebar_visible(self, visible: bool) -> None:
        """记住各页面的展开状态和侧栏宽度；折叠不丢失规则或检测结果。"""
        self._sidebar_preferences[self.results_tabs.currentWidget()] = visible
        was_visible = not self.cleaning_scroll.isHidden()
        if was_visible and not visible:
            self._cleaning_sidebar_width = self.workspace_splitter.sizes()[1]
        self.cleaning_scroll.setVisible(visible)
        if visible and not was_visible:
            available = self.workspace_splitter.width() - self.workspace_splitter.handleWidth()
            self.workspace_splitter.setSizes([
                max(1, available - self._cleaning_sidebar_width),
                self._cleaning_sidebar_width,
            ])
        with QSignalBlocker(self.cleaning_toggle_button):
            self.cleaning_toggle_button.setChecked(visible)
        self.cleaning_toggle_button.setText("收起清洗设置 ‹" if visible else "展开清洗设置 ›")

    @Slot()
    def _clear_annotation_preview(self) -> None:
        self._prepared_annotation = None
        self.annotation_preview_panel.clear_preview()
        self._update_controls()

    def _can_apply_annotation(self) -> bool:
        preview = self._prepared_annotation
        return (
            preview is not None and preview.column_index == self.selected_column
            and self.annotation_preview_panel.result is preview.result
            and self.data_manager.is_preview_current(preview)
        )

    @Slot()
    def _start_annotation_preview(self, kind: str = "duplicate", *, use_sheet_scan: bool = False) -> None:
        if (self._is_busy() or self._preview is None or self.selected_column is None
                or self.selected_column >= self.data_manager.original_df.shape[1]):
            return
        analysis = self.cleaning_panel.analysis
        duplicate_count = analysis.counts["重复值"] if analysis is not None else 0
        if use_sheet_scan and analysis is None:
            state = self.data_manager.current_data
            cache = state.scan_cache
            if cache is not None and cache.revision == state.revision:
                duplicate_count = cache.result.columns[self.selected_column].duplicates
        if (kind == "duplicate" and duplicate_count == 0) or (kind != "duplicate" and analysis is None):
            return
        worker = AnnotationPreviewWorker(
            self.data_manager.working_df.copy(deep=False), self.selected_column, self,
            kind=kind, date_confirmed=self.cleaning_panel.type_panel.date_confirmed_box.isChecked(),
        )
        self._annotation_worker = worker
        self._clear_annotation_preview()
        worker.succeeded.connect(self._show_annotation_preview)
        worker.failed.connect(self._show_annotation_error)
        worker.finished.connect(self._finish_annotation_preview)
        worker.finished.connect(worker.deleteLater)
        self._update_controls()
        self.statusBar().showMessage("正在生成辅助检查列预览……")
        worker.start()

    @Slot(object)
    def _show_annotation_preview(self, result: AnnotationPreview) -> None:
        if self._close_requested:
            return
        try:
            self._prepared_annotation = self.data_manager.prepare_annotation_preview(result)
        except DataStateError as error:
            self._show_annotation_error(str(error))
            return
        self.annotation_preview_panel.show_result(result)
        self.results_tabs.setCurrentWidget(self.annotation_preview_panel)
        self._update_controls()

    @Slot(str)
    def _show_annotation_error(self, detail: str) -> None:
        if not self._close_requested:
            self._show_error_dialog("辅助列预览失败", "无法生成辅助列预览，工作数据保持不变。", detail)

    @Slot()
    def _finish_annotation_preview(self) -> None:
        self._annotation_worker = None
        self._update_controls()
        self.statusBar().clearMessage()
        if self._close_requested:
            self.close()

    @Slot()
    def _apply_annotation(self) -> None:
        if self._is_busy() or not self._can_apply_annotation():
            return
        preview = self._prepared_annotation
        if not confirm_annotation(self, preview.sheet_name, preview.result):
            return
        if (self._is_busy() or self._prepared_annotation is not preview
                or not self._can_apply_annotation()):
            self.statusBar().showMessage("辅助列预览已变化，请重新生成。", 6000)
            return
        try:
            action = self.data_manager.apply_annotation(preview)
        except Exception as error:
            self._show_error_dialog("无法应用辅助列", "添加失败，工作副本保持不变。", str(error))
            return
        self._display_sheet(self.data_manager.select_sheet(action.sheet_name))
        self._select_column(action.column_index)
        if self._model.rowCount():
            self.table.scrollTo(self._model.index(0, action.column_index))
        self.statusBar().showMessage(f"已在最右侧添加 {action.column_name}，未删除任何记录，可以撤销。", 10000)

    @Slot()
    def _start_cleaning_preview(self) -> None:
        if self._is_busy() or self._preview is None or self.selected_column is None:
            return
        if (self.cleaning_panel.analysis is None
                or self.selected_column >= self.data_manager.original_df.shape[1]):
            return
        column = self.data_manager.working_df.iloc[:, self.selected_column].copy(deep=True)
        worker = CleaningPreviewWorker(column, self.cleaning_panel.current_rules(), self)
        self._cleaning_worker = worker
        self._clear_cleaning_preview()
        worker.succeeded.connect(self._show_cleaning_preview)
        worker.failed.connect(self._show_cleaning_error)
        worker.finished.connect(self._finish_cleaning_preview)
        worker.finished.connect(worker.deleteLater)
        self._update_controls()
        self.statusBar().showMessage("正在生成清洗预览，请稍候……")
        worker.start()

    @Slot(object)
    def _show_cleaning_preview(self, result: CleaningPreview) -> None:
        try:
            self._prepared_preview = self.data_manager.prepare_preview(self.selected_column, result)
        except DataStateError as error:
            self._show_cleaning_error(str(error))
            return
        self.cleaning_preview_panel.show_result(result)
        self.results_tabs.setCurrentWidget(self.cleaning_preview_panel)
        self._update_controls()

    @Slot()
    def _apply_cleaning(self) -> None:
        if self._is_busy():
            return
        if not self._can_apply_preview():
            self.statusBar().showMessage("请先选择列并生成有变化的清洗预览。", 6000)
            return
        preview = self._prepared_preview
        if not confirm_apply(self, preview.sheet_name, preview.result.column_name, preview.result):
            return
        # 模态确认返回后再次校验，确保应用的仍是用户确认过的预览。
        if self._is_busy() or self._prepared_preview is not preview or not self._can_apply_preview():
            self.statusBar().showMessage("预览已变化，请重新生成后再应用。", 6000)
            return
        try:
            action = self.data_manager.apply_preview(preview)
        except Exception as error:
            self._show_error_dialog("无法应用清洗", "应用失败，工作副本未被修改。", str(error))
            return
        if action is not None:
            self._refresh_after_change(
                action.column_index, f"已清洗{action.column_name}列，共修改 {len(action.affected_indices)} 行"
            )

    @Slot()
    def _undo_last_action(self) -> None:
        if self._is_busy() or not self.data_manager.undo_stack:
            return
        previous_sheet = self.data_manager.current_sheet
        previous_column = self.selected_column
        try:
            action = self.data_manager.undo()
        except Exception as error:
            self._show_error_dialog("撤销失败", "无法恢复此操作，请查看错误详情。", str(error))
            return
        if action is None:
            return
        if isinstance(action, AnnotationAction):
            column = action.source_column_index
            message = f"已撤销{action.sheet_name}工作表的{action.column_name}辅助列"
        elif isinstance(action, BatchCleaningAction):
            column = None
            message = f"已撤销{action.sheet_name}工作表的整批清洗"
        elif isinstance(action, CleaningAction):
            column = action.column_index
            message = f"已撤销{action.sheet_name}工作表的{action.column_name}列清洗"
        else:
            column = previous_column if previous_sheet == action.sheet_name else None
            message = f"已撤销{action.sheet_name}工作表的恢复原始操作"
        self._refresh_after_change(column, message)

    @Slot()
    def _restore_original(self) -> None:
        if (self._is_busy() or self.data_manager.current_sheet is None
                or not self.data_manager.current_data.is_modified):
            return
        sheet = self.data_manager.current_sheet
        if not confirm_restore(self, sheet):
            return
        if self._is_busy() or self.data_manager.current_sheet != sheet:
            return
        column = self.selected_column
        try:
            action = self.data_manager.restore_original()
        except Exception as error:
            self._show_error_dialog("恢复失败", "无法恢复当前工作表，请查看错误详情。", str(error))
            return
        if action is not None:
            self._refresh_after_change(column, f"已恢复{sheet}工作表的原始数据")

    @Slot()
    def _export_results(self) -> None:
        if self._is_busy():
            return
        if self.data_manager.file_path is None:
            QMessageBox.information(self, "导出提示", "请先打开 Excel 文件。")
            return
        try:
            plan = prepare_export(self.data_manager.get_export_snapshot())
        except ExcelExportError as error:
            if str(error) == NO_CHANGES_MESSAGE:
                QMessageBox.information(self, "导出提示", NO_CHANGES_MESSAGE)
            else:
                self._show_error_dialog("无法导出", str(error), error.detail)
            return
        except Exception as error:
            self._show_error_dialog("无法导出", "无法准备导出快照，源文件保持不变。", str(error))
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "导出清洗结果", str(suggest_export_path(plan.snapshot.source_path)),
            "Excel 工作簿 (*.xlsx)", options=QFileDialog.Option.DontConfirmOverwrite,
        )
        if not file_path:
            return
        target = Path(file_path)
        if not target.suffix:
            target = target.with_suffix(".xlsx")
        try:
            _, target = validate_export_paths(plan.snapshot.source_path, target)
        except ExcelExportError as error:
            self._show_error_dialog("无法导出", str(error), error.detail)
            return
        except OSError as error:
            self._show_error_dialog("无法导出", "无法访问源文件或保存位置，请检查权限。", str(error))
            return
        mode = confirm_export(self, plan, target)
        if mode is None:
            return
        plan = replace(plan, mode=mode)
        if self._is_busy() or self.data_manager.export_state_token != plan.snapshot.state_token:
            self.statusBar().showMessage("工作副本已变化，请重新确认导出。", 6000)
            return
        worker = ExcelExportWorker(plan, target, self)
        self._export_worker = worker
        self._export_completion_message = None
        worker.succeeded.connect(self._show_export_result)
        worker.failed.connect(self._show_export_error)
        worker.finished.connect(self._finish_export)
        worker.finished.connect(worker.deleteLater)
        self._update_controls()
        self.statusBar().showMessage("正在导出清洗结果，请稍候……")
        worker.start()

    @Slot(object)
    def _show_export_result(self, result: ExportResult) -> None:
        self._export_completion_message = (
            f"导出成功：{result.target_path.name}（{result.mode.label}，"
            f"{result.sheet_count} 个工作表，{result.cell_count} 个单元格，{result.annotation_column_count} 个辅助列）"
        )
        if not self._close_requested:
            show_export_success(self, result)

    @Slot(str, str)
    def _show_export_error(self, message: str, detail: str) -> None:
        self._export_completion_message = "导出失败，源文件和工作副本均保持不变。"
        if not self._close_requested:
            self._show_error_dialog("无法导出清洗结果", message, detail)

    @Slot()
    def _finish_export(self) -> None:
        self._export_worker = None
        self._update_controls()
        if self._export_completion_message:
            self.statusBar().showMessage(self._export_completion_message, 10000)
        else:
            self.statusBar().clearMessage()
        if self._close_requested:
            self.close()

    def _refresh_after_change(self, column: int | None, message: str) -> None:
        """先更新工作表并清除旧结果，再后台检测；状态消息在检测结束后保留。"""
        self._display_sheet(self.data_manager.select_sheet(self.data_manager.current_sheet))
        if column is not None and 0 <= column < self._model.columnCount():
            self._select_column(column)
            self._start_analysis(completion_message=message)
        else:
            self.statusBar().showMessage(message, 10000)

    @Slot(str)
    def _show_cleaning_error(self, detail: str) -> None:
        if not self._close_requested:
            self._show_error_dialog("清洗预览失败", "无法生成清洗预览，原始数据保持不变。", detail)

    @Slot()
    def _finish_cleaning_preview(self) -> None:
        self._cleaning_worker = None
        self._update_controls()
        self.statusBar().clearMessage()
        if self._close_requested:
            self.close()

    def _refresh_batch_workspace(self) -> None:
        self._prepared_batch = None
        state = self.data_manager.current_data
        cache = state.scan_cache
        analysis = cache.result if cache and cache.revision == state.revision else None
        configs = recommend_batch_configs(
            state.working_df.columns, analysis, original_column_count=state.original_df.shape[1],
        )
        self.batch_panel.set_sheet(self.data_manager.current_sheet, configs, scanned=analysis is not None)

    @Slot()
    def _open_batch_workspace(self) -> None:
        if self._preview is not None:
            self.results_tabs.setCurrentWidget(self.batch_panel)

    @Slot()
    def _request_batch_scan(self) -> None:
        if self._is_busy() or self._preview is None:
            return
        self._scan_for_batch = True
        self.batch_panel.begin_task("正在扫描并生成逐列建议……", self._model.columnCount())
        self._start_sheet_scan()
        if self._scan_worker is None:
            self._scan_for_batch = False

    @Slot()
    def _clear_batch_preview(self) -> None:
        self._prepared_batch = None
        self.batch_panel.clear_preview()
        self._update_controls()

    def _can_apply_batch(self) -> bool:
        preview = self._prepared_batch
        return (
            preview is not None and self.batch_panel.result is preview.result
            and preview.result.changed_cell_count > 0
            and preview.result.configs == tuple(config for config in self.batch_panel.configs if config.enabled)
            and self.data_manager.is_preview_current(preview)
        )

    @Slot()
    def _start_batch_preview(self) -> None:
        if self._is_busy() or self._preview is None or not self.batch_panel.scanned:
            return
        self._clear_batch_preview()
        self._batch_snapshot = self.data_manager.prepare_sheet_scan()
        worker = BatchPreviewWorker(
            self._batch_snapshot, self.batch_panel.configs, self.data_manager.original_df.shape[1], self,
        )
        self._batch_worker = worker
        self._batch_completion_message = None
        self.batch_panel.begin_task(
            "正在生成批量预览……", sum(config.enabled for config in self.batch_panel.configs),
        )
        worker.progress.connect(self.batch_panel.show_progress)
        worker.succeeded.connect(self._show_batch_preview)
        worker.failed.connect(self._show_batch_error)
        worker.cancelled.connect(self._show_batch_cancelled)
        worker.finished.connect(self._finish_batch_preview)
        worker.finished.connect(worker.deleteLater)
        self._update_controls()
        self.statusBar().showMessage("正在生成批量预览……")
        worker.start()

    @Slot(object)
    def _show_batch_preview(self, result: BatchCleaningPreview) -> None:
        if self._close_requested or self._batch_worker.isInterruptionRequested():
            return
        try:
            self._prepared_batch = self.data_manager.prepare_batch_preview(self._batch_snapshot, result)
        except Exception as error:
            self._show_batch_error(str(error))
            return
        self.batch_panel.show_result(result)
        self._batch_completion_message = (
            f"批量预览完成：预计修改 {result.changed_cell_count} 个单元格，涉及 {result.changed_row_count} 行。"
        )
        self._update_controls()

    @Slot(str)
    def _show_batch_error(self, detail: str) -> None:
        self._prepared_batch = None
        self._batch_completion_message = "批量预览未生成，工作副本保持不变。"
        if not self._close_requested:
            self.batch_panel.show_notice(detail, "warning")
            self._show_error_dialog("无法生成批量预览", self._batch_completion_message, detail)

    @Slot()
    def _show_batch_cancelled(self) -> None:
        self._prepared_batch = None
        self._batch_completion_message = "批量预览已取消，工作副本保持不变。"
        if not self._close_requested:
            self.batch_panel.clear_preview()
            self.batch_panel.show_notice(self._batch_completion_message)

    @Slot()
    def _finish_batch_preview(self) -> None:
        self._batch_worker = None
        self._batch_snapshot = None
        self._update_controls()
        if self._batch_completion_message:
            self.statusBar().showMessage(self._batch_completion_message, 10000)
        if self._close_requested:
            self.close()

    @Slot()
    def _cancel_batch_task(self) -> None:
        if self._batch_worker is not None:
            self._batch_worker.requestInterruption()
        elif self._scan_for_batch:
            self._cancel_sheet_scan()

    @Slot()
    def _apply_batch_cleaning(self) -> None:
        if self._is_busy() or not self._can_apply_batch():
            return
        preview = self._prepared_batch
        if not confirm_batch_apply(self, preview.result):
            return
        if self._is_busy() or self._prepared_batch is not preview or not self._can_apply_batch():
            self.statusBar().showMessage("批量预览已变化，请重新生成后再应用。", 6000)
            return
        try:
            action = self.data_manager.apply_batch_preview(preview)
        except Exception as error:
            self._show_error_dialog("无法应用批量清洗", "应用失败，工作副本未被修改。", str(error))
            return
        if action is not None:
            self._refresh_after_change(
                None, f"批量清洗完成，共修改 {preview.result.changed_cell_count} 个单元格，"
                f"涉及 {preview.result.changed_row_count} 行。",
            )

    @Slot(int)
    def _request_batch_duplicate(self, column: int) -> None:
        if self._is_busy():
            return
        self._select_column(column)
        self._start_annotation_preview(use_sheet_scan=True)

    def _refresh_sheet_overview(self) -> None:
        state = self.data_manager.current_data
        cache = state.scan_cache
        stale = cache is not None and cache.revision != state.revision
        self.sheet_overview_panel.set_sheet(
            self.data_manager.current_sheet, len(state.working_df),
            cache.result if cache is not None else None, stale=stale,
        )
        self.results_tabs.setTabText(
            self.results_tabs.indexOf(self.sheet_overview_panel),
            "数据质量总览 · 待更新" if stale else "数据质量总览",
        )

    @Slot()
    def _start_sheet_scan(self) -> None:
        if self._is_busy() or self._preview is None:
            return
        try:
            snapshot = self.data_manager.prepare_sheet_scan()
        except Exception as error:
            self._show_error_dialog("无法扫描工作表", "无法准备当前工作副本，数据保持不变。", str(error))
            return
        worker = SheetAnalysisWorker(snapshot, self)
        self._scan_worker = worker
        self._scan_snapshot = snapshot
        self._scan_completion_message = None
        self._scan_cancel_requested = False
        self.sheet_overview_panel.begin_scan(snapshot.dataframe.shape[1])
        self.results_tabs.setCurrentWidget(self.batch_panel if self._scan_for_batch else self.sheet_overview_panel)
        worker.progress.connect(self._show_scan_progress)
        worker.succeeded.connect(self._show_sheet_scan)
        worker.failed.connect(self._show_scan_error)
        worker.cancelled.connect(self._show_scan_cancelled)
        worker.finished.connect(self._finish_sheet_scan)
        worker.finished.connect(worker.deleteLater)
        self._update_controls()
        self.statusBar().showMessage("正在扫描当前工作表……")
        worker.start()

    @Slot(int, int, str)
    def _show_scan_progress(self, completed: int, total: int, column_name: str) -> None:
        if not self._close_requested:
            self.sheet_overview_panel.show_progress(completed, total, column_name)
            if self._scan_for_batch:
                self.batch_panel.show_progress(completed, total, column_name)

    @Slot(object)
    def _show_sheet_scan(self, result: SheetAnalysis) -> None:
        if self._close_requested:
            return
        if self._scan_cancel_requested:
            self._show_scan_cancelled()
            return
        try:
            self.data_manager.save_sheet_scan(self._scan_snapshot, result)
        except DataStateError as error:
            self._show_scan_error(str(error))
            return
        self._refresh_sheet_overview()
        self._refresh_batch_workspace()
        self._scan_completion_message = (
            f"扫描完成 · 字段数：{result.column_count} · 发现问题字段：{result.problem_column_count}"
            f" · 正常字段：{result.normal_column_count}"
        )
        self._update_controls()

    @Slot()
    def _cancel_sheet_scan(self) -> None:
        if self._scan_worker is not None:
            self._scan_cancel_requested = True
            self._scan_worker.requestInterruption()
            self.sheet_overview_panel.show_cancel_pending()

    @Slot()
    def _show_scan_cancelled(self) -> None:
        if self._close_requested:
            return
        self._refresh_sheet_overview()
        cached = self.data_manager.current_data.scan_cache
        stale = cached is not None and cached.revision != self.data_manager.current_data.revision
        message = "扫描已取消，未发布不完整结果。"
        if cached is not None:
            message += "当前保留上次完整扫描；数据修改后仍需重新扫描。" if stale else "当前显示上次完整扫描结果。"
        self.sheet_overview_panel.show_notice(message, "warning" if stale else "neutral")
        self._scan_completion_message = "扫描已取消"
        if self._scan_for_batch:
            self.batch_panel.show_notice(message)

    @Slot(str)
    def _show_scan_error(self, detail: str) -> None:
        self._scan_completion_message = "工作表扫描失败，数据保持不变；请重新扫描。"
        if not self._close_requested:
            self._refresh_sheet_overview()
            self.sheet_overview_panel.show_notice(self._scan_completion_message, "warning")
            if self._scan_for_batch:
                self.batch_panel.show_notice(self._scan_completion_message, "warning")
            self._show_error_dialog("数据质量扫描失败", self._scan_completion_message, detail)

    @Slot()
    def _finish_sheet_scan(self) -> None:
        self._scan_worker = None
        self._scan_snapshot = None
        self._scan_for_batch = False
        self._update_controls()
        if self._scan_completion_message:
            self.statusBar().showMessage(self._scan_completion_message, 10000)
        self._scan_completion_message = None
        if self._close_requested:
            self.close()

    @Slot(int)
    def _select_overview_column(self, column: int) -> None:
        if self._is_busy() or not 0 <= column < self._model.columnCount():
            return
        self._select_column(column)
        if self._model.rowCount():
            self.table.scrollTo(self._model.index(0, column))
        # 复用现有后台单列入口，避免全表缓存数百万条问题明细。
        self._start_analysis()

    @Slot()
    def _start_analysis(self, completion_message: str | None = None) -> None:
        if self._preview is None or self.selected_column is None:
            return
        if self._is_busy():
            return
        self._analysis_completion_message = completion_message
        self.diagnosis_panel.set_column(self._model.column_name(self.selected_column), self._model.rowCount())
        self.cleaning_panel.set_column(
            self._model.column_name(self.selected_column), f"{get_column_letter(self.selected_column + 1)}2",
            is_generated=self.selected_column >= self.data_manager.original_df.shape[1],
        )
        # 只复制当前列作为后台任务快照，不改动源 DataFrame。
        column = self.data_manager.working_df.iloc[:, self.selected_column].copy(deep=True)
        worker = ColumnAnalysisWorker(column, self)
        self._analysis_worker = worker
        worker.succeeded.connect(self._show_analysis_result)
        worker.failed.connect(self._show_analysis_error)
        worker.finished.connect(self._finish_analysis)
        worker.finished.connect(worker.deleteLater)
        self._update_controls()
        self.statusBar().showMessage("正在检测当前列，请稍候……")
        worker.start()

    @Slot(object)
    def _show_analysis_result(self, result: ColumnAnalysis) -> None:
        self.diagnosis_panel.show_result(result)
        self.cleaning_panel.set_analysis(result)
        self.results_tabs.setCurrentWidget(self.diagnosis_panel)

    @Slot(str)
    def _show_analysis_error(self, detail: str) -> None:
        if self._analysis_completion_message is not None:
            self._analysis_completion_message = "工作副本已更新，重新检测失败；请点击“开始检测”重试。"
        if not self._close_requested:
            self._show_error_dialog("数据检测失败", "无法完成当前列检测，请查看错误详情。", detail)

    @Slot()
    def _finish_analysis(self) -> None:
        self._analysis_worker = None
        self._update_controls()
        if self._analysis_completion_message:
            self.statusBar().showMessage(self._analysis_completion_message, 10000)
        else:
            self.statusBar().clearMessage()
        self._analysis_completion_message = None
        if self._close_requested:
            self.close()

    @Slot()
    def _fit_columns(self) -> None:
        # 一次调整列宽，不切换为自动锁定模式，调整后仍可手动拖动。
        self._current_table().resizeColumnsToContents()

    @Slot(QModelIndex)
    def _show_cell_preview(self, index: QModelIndex) -> None:
        if not index.isValid():
            return
        dialog = CellPreviewDialog(
            self._model.column_name(index.column()),
            index.row() + 1,
            self._model.data(index),
            self,
        )
        dialog.exec()
        dialog.deleteLater()

    @Slot()
    def _open_excel(self) -> None:
        file_path, _ = QFileDialog.getOpenFileName(
            self, "打开 Excel", "", "Excel 工作簿 (*.xlsx)"
        )
        if file_path:
            self._start_load(file_path)

    @Slot(str)
    def _switch_sheet(self, sheet_name: str) -> None:
        if self._is_busy() or self._preview is None or sheet_name == self._preview.sheet_name:
            return
        if sheet_name in self.data_manager.sheet_data:
            self._display_sheet(self.data_manager.select_sheet(sheet_name))
        else:
            self._start_load(str(self._preview.file_path), sheet_name)

    def _start_load(self, file_path: str, sheet_name: str | None = None) -> None:
        if self._is_busy():
            return
        self._loading_new_workbook = sheet_name is None
        self.statusBar().showMessage("正在读取 Excel，请稍候……")
        worker = ExcelLoadWorker(file_path, sheet_name, self)
        self._worker = worker
        self._update_controls()
        worker.succeeded.connect(self._show_preview)
        worker.failed.connect(self._show_load_error)
        worker.finished.connect(self._finish_load)
        worker.finished.connect(worker.deleteLater)
        worker.start()

    @Slot(object)
    def _show_preview(self, preview: ExcelPreview) -> None:
        try:
            if self._loading_new_workbook or self.data_manager.file_path != preview.file_path.resolve():
                self.data_manager.open_workbook(preview)
            else:
                self.data_manager.add_sheet(preview)
        except Exception as error:
            self._show_load_error("无法建立数据副本，原有数据保持不变。", str(error))
            return
        self._display_sheet(preview)

    def _display_sheet(self, preview: ExcelPreview) -> None:
        self._preview = preview
        # 更新下拉框时阻止信号，避免触发重复读取。
        with QSignalBlocker(self.sheet_combo):
            self.sheet_combo.clear()
            self.sheet_combo.addItems(preview.sheet_names)
            self.sheet_combo.setCurrentText(preview.sheet_name)

        self._model.set_dataframe(self.data_manager.working_df)
        self.preview_stack.setCurrentWidget(self.table)
        self.table.scrollToTop()
        self.table.horizontalScrollBar().setValue(0)
        self.selected_column = None
        self.selected_column_label.setText("当前选择列：未选择")
        self.diagnosis_panel.set_column(None, self._model.rowCount())
        self.cleaning_panel.set_column(None, None)
        self._refresh_sheet_overview()
        self._refresh_batch_workspace()
        self.results_tabs.setCurrentWidget(self.raw_data_page)
        self.file_label.setText(f"当前文件名：{preview.file_path.name}")
        self.file_label.setToolTip(str(preview.file_path))
        self.sheet_label.setText(f"当前工作表：{preview.sheet_name}")
        self.row_count_label.setText(f"总行数：{self._model.rowCount()}")
        self.column_count_label.setText(f"总列数：{self._model.columnCount()}")
        self._update_modification_state()

    @Slot(str, str)
    def _show_load_error(self, message: str, detail: str) -> None:
        # 失败时保留上次成功读取的数据，并恢复下拉框的工作表名称。
        if self._preview is not None:
            with QSignalBlocker(self.sheet_combo):
                self.sheet_combo.setCurrentText(self._preview.sheet_name)
        if self._close_requested:
            return
        self._show_error_dialog("无法读取 Excel", message, detail)

    def _show_error_dialog(self, title: str, message: str, detail: str) -> None:
        dialog = QMessageBox(self)
        dialog.setWindowTitle(title)
        dialog.setIcon(QMessageBox.Icon.Warning)
        dialog.setText(message)
        dialog.setStandardButtons(QMessageBox.StandardButton.Ok)
        dialog.button(QMessageBox.StandardButton.Ok).setText("确定")
        if detail:
            dialog.setDetailedText(detail)
        dialog.exec()

    @Slot()
    def _finish_load(self) -> None:
        self._worker = None
        self._loading_new_workbook = False
        self._update_controls()
        self.statusBar().clearMessage()
        if self._close_requested:
            self.close()

    @Slot(int)
    def _select_column(self, column: int) -> None:
        if not 0 <= column < self._model.columnCount():
            return
        if self._is_busy():
            return
        if column != self.selected_column:
            self.diagnosis_panel.set_column(self._model.column_name(column), self._model.rowCount())
            self.cleaning_panel.set_column(
                self._model.column_name(column), f"{get_column_letter(column + 1)}2",
                is_generated=column >= self.data_manager.original_df.shape[1],
            )
        self.selected_column = column
        self.selected_column_label.setText(f"当前选择列：{self._model.column_name(column)}")
        self._update_controls()
        # 空工作表也能记录字段选择；有数据时同步高亮整列。
        if self._model.rowCount():
            self.table.selectionModel().select(
                self._model.index(0, column),
                QItemSelectionModel.SelectionFlag.ClearAndSelect
                | QItemSelectionModel.SelectionFlag.Columns,
            )

    def closeEvent(self, event: QCloseEvent) -> None:
        if self._is_busy():
            if self._batch_worker is not None:
                self._batch_worker.requestInterruption()
            if self._scan_worker is not None:
                self._scan_worker.requestInterruption()
            # 等待后台任务结束后关闭，避免仍在运行的 QThread 被销毁。
            self._close_requested = True
            self.setEnabled(False)
            self.statusBar().showMessage("正在完成当前任务，完成后自动关闭……")
            event.ignore()
            return
        event.accept()
