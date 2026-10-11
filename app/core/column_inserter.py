"""插入对照列时更新普通 A1 引用及常见工作表结构；复杂依赖明确拒绝。"""

from bisect import bisect_left
from collections import OrderedDict
from copy import copy, deepcopy
from dataclasses import dataclass
import re

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.print_settings import PrintArea


class ColumnInsertionError(ValueError):
    pass


@dataclass(frozen=True)
class ColumnMap:
    after_columns: tuple[int, ...]

    def column(self, old_column: int) -> int:
        return old_column + bisect_left(self.after_columns, old_column)

    def coordinate(self, reference: str) -> str:
        match = re.fullmatch(r"(\$?)([A-Za-z]{1,3})(\$?)([1-9][0-9]*)", reference)
        if match is None:
            raise ColumnInsertionError("存在无法识别的单元格引用，请选择“原列更新”。")
        column = self.column(column_index_from_string(match[2]))
        if column > 16384:
            raise ColumnInsertionError("新增列会超出 Excel 最大列数，请选择“原列更新”。")
        return f"{match[1]}{get_column_letter(column)}{match[3]}{match[4]}"

    def crosses(self, start: int, end: int) -> bool:
        return any(start <= column < end for column in self.after_columns)

    def parts(self, area: CellRange) -> list[str]:
        """拆分适用范围，保留对原始列的验证/条件格式，避免施加到新增列。"""
        cuts = [area.min_col] + [
            column + 1 for column in self.after_columns if area.min_col <= column < area.max_col
        ] + [area.max_col + 1]
        return [
            CellRange(min_col=self.column(start), max_col=self.column(end - 1),
                      min_row=area.min_row, max_row=area.max_row).coord
            for start, end in zip(cuts, cuts[1:])
        ]

    def area(self, reference: str, include_added: bool = False) -> str:
        area = CellRange(reference)
        maximum = self.column(area.max_col)
        if include_added and area.max_col in self.after_columns:
            maximum += 1
        return CellRange(min_col=self.column(area.min_col), max_col=maximum,
                         min_row=area.min_row, max_row=area.max_row).coord


def _sheet_name(prefix: str) -> str:
    if prefix.startswith("'") and prefix.endswith("'"):
        return prefix[1:-1].replace("''", "'")
    return prefix


def translate_formula(formula: str, context: str, maps: dict[str, ColumnMap]) -> str:
    """按插入位置更新引用，绝对引用同样跟随原始单元格，文本常量保持不变。"""
    try:
        tokens = Tokenizer(formula).items
    except Exception as error:
        raise ColumnInsertionError("原工作簿存在无法解析的公式，请选择“原列更新”。") from error
    lookup = {name.casefold(): mapping for name, mapping in maps.items()}
    modified = False
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN":
            function = token.value[:-1].upper().removeprefix("_XLFN.")
            if function in {"INDIRECT", "OFFSET", "ADDRESS"}:
                raise ColumnInsertionError("原工作簿含动态地址公式，无法安全插入对照列，请选择“原列更新”。")
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        value = token.value
        prefix, address = value.rsplit("!", 1) if "!" in value else ("", value)
        if "[" in value or "]" in value or ":" in _sheet_name(prefix):
            raise ColumnInsertionError("存在外部、结构化或跨工作表区间引用，请选择“原列更新”。")
        mapping = lookup.get((_sheet_name(prefix) if prefix else context).casefold())
        if mapping is None:
            continue
        endpoints = address.split(":")
        cell = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*"
        column = r"\$?[A-Za-z]{1,3}"
        row = r"\$?[1-9][0-9]*"
        if all(re.fullmatch(cell, item) for item in endpoints) and len(endpoints) <= 2:
            bounds = [column_index_from_string(re.search(r"[A-Za-z]+", item)[0]) for item in endpoints]
            row_bounds = [int(re.search(r"[0-9]+$", item)[0]) for item in endpoints]
            if any(number > 16384 for number in bounds) or any(number > 1048576 for number in row_bounds):
                continue  # 合法的自定义名称可能看起来像超出边界的 A1 地址。
            if mapping.crosses(min(bounds), max(bounds)):
                raise ColumnInsertionError(
                    "原公式引用范围跨过了待新增列，插入会改变计算范围，请选择“原列更新”。"
                )
            shifted = ":".join(mapping.coordinate(item) for item in endpoints)
        elif len(endpoints) == 2 and all(re.fullmatch(column, item) for item in endpoints):
            bounds = [column_index_from_string(item.lstrip("$")) for item in endpoints]
            if mapping.crosses(min(bounds), max(bounds)):
                raise ColumnInsertionError("原公式的整列范围跨过待新增列，请选择“原列更新”。")
            if any(mapping.column(bound) > 16384 for bound in bounds):
                raise ColumnInsertionError("原公式引用会超出 Excel 最大列数，请选择“原列更新”。")
            shifted = ":".join(
                ("$" if item.startswith("$") else "") + get_column_letter(mapping.column(bound))
                for item, bound in zip(endpoints, bounds)
            )
        elif len(endpoints) == 2 and all(re.fullmatch(row, item) for item in endpoints):
            raise ColumnInsertionError("原公式引用整行，插入对照列可能改变结果，请选择“原列更新”。")
        else:
            if any(character in address for character in ("@", "#", ":", "!")):
                raise ColumnInsertionError("原公式含未支持的区域引用，请选择“原列更新”。")
            continue  # 定义名称的实际地址会单独更新。
        updated = (prefix + "!" if prefix else "") + shifted
        if updated != value:
            token.value = updated
            modified = True
    return "=" + "".join(token.value for token in tokens) if modified else formula


def _expression(value: str | None, sheet: str, maps: dict[str, ColumnMap]) -> str | None:
    if not value:
        return value
    formula = value if value.startswith("=") else "=" + value
    translated = translate_formula(formula, sheet, maps)
    return translated if value.startswith("=") else translated[1:]


def _shift_title_columns(reference: str | None, mapping: ColumnMap) -> str | None:
    if not reference:
        return reference
    first, last = reference.replace("$", "").split(":")
    start, end = column_index_from_string(first), column_index_from_string(last)
    maximum = mapping.column(end) + (end in mapping.after_columns)
    return f"${get_column_letter(mapping.column(start))}:${get_column_letter(maximum)}"


def insert_comparison_columns(workbook, maps: dict[str, ColumnMap]) -> None:
    """先计算所有引用和结构更新，再插列；调用方只在临时副本上执行。"""
    formula_updates, name_updates, metadata = [], [], {}
    for sheet in workbook:
        # 暂不尝试猜测图形锚点、表格结构化引用或透视表依赖。
        if sheet._charts or sheet._images or sheet._pivots:
            raise ColumnInsertionError("工作簿含图形或透视表，无法安全插入对照列，请选择“原列更新”。")
        if sheet.title in maps and sheet.tables:
            raise ColumnInsertionError("待处理工作表含 Excel 表格对象，请选择“原列更新”。")
        for cell in tuple(sheet._cells.values()):
            if cell.data_type == "f":
                if not isinstance(cell.value, str):
                    raise ColumnInsertionError("工作簿含数组或数据表公式，请选择“原列更新”。")
                translated = translate_formula(cell.value, sheet.title, maps)
                if translated != cell.value:
                    formula_updates.append((sheet.title, cell.row, cell.column, translated))
        for name in sheet.defined_names.values():
            name_updates.append((name, _expression(name.attr_text, sheet.title, maps)))
        mapping = maps.get(sheet.title, ColumnMap(()))
        if sheet.max_column + len(mapping.after_columns) > 16384:
            raise ColumnInsertionError("新增列会超出 Excel 最大列数，请选择“原列更新”。")
        merges = []
        for area in sheet.merged_cells.ranges:
            if mapping.crosses(area.min_col, area.max_col):
                raise ColumnInsertionError("新增列会拆开合并区域，请选择“原列更新”。")
            merges.append(mapping.area(area.coord))
        dimensions = {}
        originals = [(key, copy(value)) for key, value in sheet.column_dimensions.items()]
        for key, value in originals:
            first = value.min or column_index_from_string(key)
            last = value.max or first
            if mapping.column(last) > 16384:
                raise ColumnInsertionError("原列格式范围会超出 Excel 最大列数，请选择“原列更新”。")
            for part in mapping.parts(CellRange(min_col=first, max_col=last, min_row=1, max_row=1)):
                area = CellRange(part)
                shifted = copy(value)
                shifted.index = get_column_letter(area.min_col)
                shifted.min, shifted.max = area.min_col, area.max_col
                dimensions[shifted.index] = shifted
        for column in mapping.after_columns:
            source = next((value for key, value in originals
                           if (value.min or column_index_from_string(key)) <= column
                           <= (value.max or value.min or column_index_from_string(key))), None)
            if source is not None:
                shifted = copy(source)
                shifted.index = get_column_letter(mapping.column(column) + 1)
                shifted.min = shifted.max = mapping.column(column) + 1
                shifted.hidden = False
                dimensions[shifted.index] = shifted
        auto_filter = deepcopy(sheet.auto_filter)
        if auto_filter.ref:
            original_area = CellRange(auto_filter.ref)
            auto_filter.ref = mapping.area(auto_filter.ref, include_added=True)
            for item in auto_filter.filterColumn:
                old = original_area.min_col + item.colId
                item.colId = mapping.column(old) - mapping.column(original_area.min_col)
            if auto_filter.sortState:
                auto_filter.sortState.ref = mapping.area(auto_filter.sortState.ref, include_added=True)
                for condition in auto_filter.sortState.sortCondition:
                    condition.ref = mapping.area(condition.ref)
        validations = deepcopy(sheet.data_validations)
        for validation in validations.dataValidation:
            validation.sqref = " ".join(
                part for area in validation.sqref.ranges for part in mapping.parts(area)
            )
            validation.formula1 = _expression(validation.formula1, sheet.title, maps)
            validation.formula2 = _expression(validation.formula2, sheet.title, maps)
        conditional = OrderedDict()
        for key, rules in sheet.conditional_formatting._cf_rules.items():
            updated_key = deepcopy(key)
            updated_key.sqref = " ".join(
                part for area in key.sqref.ranges for part in mapping.parts(area)
            )
            updated_rules = deepcopy(rules)
            for rule in updated_rules:
                rule.formula = [_expression(value, sheet.title, maps) for value in rule.formula]
            conditional[updated_key] = updated_rules
        metadata[sheet.title] = (
            merges, dimensions, auto_filter, validations, conditional,
            mapping.coordinate(sheet.freeze_panes) if sheet.freeze_panes else None,
            [mapping.area(area.coord, include_added=True) for area in PrintArea.from_string(sheet.print_area).ranges],
            _shift_title_columns(sheet.print_title_cols, mapping),
        )
    for name in workbook.defined_names.values():
        if name.attr_text:
            expression = name.attr_text if name.attr_text.startswith("=") else "=" + name.attr_text
            for token in Tokenizer(expression).items:
                if (token.type == "OPERAND" and token.subtype == "RANGE" and "!" not in token.value
                        and re.fullmatch(r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]*(?::\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)?", token.value)):
                    raise ColumnInsertionError("工作簿含未限定工作表的相对定义名称，请选择“原列更新”。")
        name_updates.append((name, _expression(name.attr_text, workbook.active.title, maps)))
    for sheet in workbook:
        sheet_name = sheet.title
        mapping = maps.get(sheet_name, ColumnMap(()))
        for column in reversed(mapping.after_columns):
            sheet.insert_cols(column + 1)
        merges, dimensions, auto_filter, validations, conditional, freeze, print_area, title_cols = metadata[sheet_name]
        sheet.merged_cells.ranges.clear()
        for area in merges:
            sheet.merged_cells.add(area)
        sheet.column_dimensions.clear()
        sheet.column_dimensions.update(dimensions)
        sheet.auto_filter = auto_filter
        sheet.data_validations = validations
        sheet.conditional_formatting._cf_rules = conditional
        sheet.freeze_panes = freeze
        sheet.print_area = print_area
        sheet.print_title_cols = title_cols
    for sheet_name, row, column, formula in formula_updates:
        mapping = maps.get(sheet_name)
        workbook[sheet_name].cell(row, mapping.column(column) if mapping else column).value = formula
    for name, expression in name_updates:
        name.attr_text = expression
