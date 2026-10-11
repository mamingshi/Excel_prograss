"""复制源工作簿，仅将当前差异写入临时副本，成功后发布为新文件。"""

from copy import copy
import os
import re
from pathlib import Path
import shutil
import tempfile
import warnings
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.xml.functions import Element, xmlfile
import pandas as pd

from app.core.file_loader import ExcelFileLoader, file_sha256
from app.models.export_data import CellChange, ExportPlan, ExportResult, ExportSnapshot, ExportMode
from app.core.cleaner import apply_cleaning_rules
from app.core.type_utils import values_equal
from app.core.annotation_builder import build_duplicate_formula, build_annotation_values
from app.core.column_inserter import ColumnMap, ColumnInsertionError, insert_comparison_columns
from app.core.export_formula_builder import build_export_formula, build_passthrough_formula


NO_CHANGES_MESSAGE = "当前没有任何已应用的清洗修改或辅助列，无需导出。"


class ExcelExportError(Exception):
    def __init__(self, message: str, detail: str = "") -> None:
        super().__init__(message)
        self.detail = detail


def dataframe_index_to_excel_row(position: int) -> int:
    """传入数据行的位置，而非 DataFrame index 标签；Excel 首行为表头。"""
    if position < 0:
        raise ExcelExportError("数据行位置不能为负数。")
    return position + 2


def dataframe_column_to_excel_column(position: int) -> int:
    """读取器保留 Excel 全部列的原始顺序，因此按位置映射，支持重名字段。"""
    if position < 0:
        raise ExcelExportError("数据列位置不能为负数。")
    return position + 1


def _values_equal(before: object, after: object) -> bool:
    return values_equal(before, after)


def _validate_sheet_structure(sheet) -> None:
    original, working = sheet.original_df, sheet.working_df
    width = original.shape[1]
    if (len(original) != len(working) or not original.index.equals(working.index)
            or working.shape[1] != width + len(sheet.annotations)
            or not original.columns.equals(working.columns[:width])):
        raise ExcelExportError("工作副本的行列结构已变化，当前版本无法安全导出。", sheet.sheet_name)
    used_names = {str(name).casefold() for name in original.columns}
    for position, annotation in enumerate(sheet.annotations):
        if (not 0 <= annotation.source_column_index < width
                or annotation.column_name != working.columns[width + position]
                or annotation.column_name.casefold() in used_names):
            raise ExcelExportError("辅助列结构与已应用记录不一致，请重新生成。", sheet.sheet_name)
        used_names.add(annotation.column_name.casefold())
        expected = build_annotation_values(working.iloc[:, annotation.source_column_index], annotation.kind,
                                           date_confirmed=annotation.date_confirmed)
        expected.name = annotation.column_name
        if not expected.equals(working.iloc[:, width + position]):
            raise ExcelExportError("辅助列标签与当前源字段不一致，请重新生成。", sheet.sheet_name)


def get_changes(snapshot: ExportSnapshot) -> tuple[CellChange, ...]:
    """遍历已加载 Sheet 的当前差异，不读取撤销日志，也不修改快照。"""
    changes = []
    names = set()
    for sheet in snapshot.sheets:
        if sheet.sheet_name in names:
            raise ExcelExportError("导出快照中存在重复工作表。")
        names.add(sheet.sheet_name)
        original, working = sheet.original_df, sheet.working_df
        _validate_sheet_structure(sheet)
        for column in range(original.shape[1]):
            before, after = original.iloc[:, column], working.iloc[:, column]
            if before.equals(after):
                continue
            for row, (old_value, new_value) in enumerate(zip(before, after)):
                if not _values_equal(old_value, new_value):
                    changes.append(CellChange(
                        sheet.sheet_name, dataframe_index_to_excel_row(row),
                        dataframe_column_to_excel_column(column), str(original.columns[column]),
                        old_value, new_value,
                    ))
    return tuple(changes)


def prepare_export(snapshot: ExportSnapshot, mode: ExportMode = ExportMode.UPDATE_COLUMNS) -> ExportPlan:
    changes = get_changes(snapshot)
    if not changes and not any(sheet.annotations for sheet in snapshot.sheets):
        raise ExcelExportError(NO_CHANGES_MESSAGE)
    return ExportPlan(snapshot, changes, ExportMode(mode))


def suggest_export_path(source_path: str | Path) -> Path:
    source = Path(source_path)
    number = 0
    while True:
        suffix = f"_{number}" if number else ""
        target = source.with_name(f"{source.stem}_已清洗{suffix}.xlsx")
        if not os.path.lexists(target):
            return target
        number += 1


def validate_export_paths(source_path: str | Path, target_path: str | Path) -> tuple[Path, Path]:
    source, target = Path(source_path).resolve(), Path(target_path).resolve()
    if source.suffix.lower() != ".xlsx" or target.suffix.lower() != ".xlsx":
        raise ExcelExportError("当前版本仅支持导出新的 .xlsx 文件。")
    if not source.is_file():
        raise ExcelExportError("源文件不存在或已被移动，请重新打开文件。")
    if (os.path.normcase(source) == os.path.normcase(target)
            or (target.exists() and source.samefile(target))):
        raise ExcelExportError("不能覆盖原 Excel 文件，请选择一个新的文件名。")
    if os.path.lexists(target_path) or os.path.lexists(target):
        raise ExcelExportError("目标文件已经存在，请选择一个新的文件名；已有文件不会被覆盖。")
    if not target.parent.is_dir():
        raise ExcelExportError("保存目录不存在，请选择一个已有的文件夹。")
    return source, target


def _validate_source_data(path: Path, snapshot: ExportSnapshot) -> None:
    """沿用导入约定核对表头、列顺序和原始数据，拒绝坐标错位。"""
    with pd.ExcelFile(path, engine="openpyxl") as workbook:
        for sheet in snapshot.sheets:
            if sheet.sheet_name not in workbook.sheet_names:
                raise ExcelExportError("源工作簿中找不到工作表，已取消导出。", sheet.sheet_name)
            loaded = ExcelFileLoader.read_sheet(workbook, sheet.sheet_name)
            original = sheet.original_df.reset_index(drop=True)
            if original.shape != loaded.shape or not original.columns.equals(loaded.columns):
                raise ExcelExportError("DataFrame 与 Excel 的行列结构不一致，已取消导出。", sheet.sheet_name)
            # 导入把真正空单元格表示为 ""；允许测试/业务快照使用 None 或 NaN 表示同一空格位。
            if not original.equals(loaded) and not original.fillna("").equals(loaded.fillna("")):
                raise ExcelExportError("原始快照与源 Excel 数据不一致，请重新打开文件。", sheet.sheet_name)


def _reject_loss_warnings(captured: list) -> None:
    messages = [str(item.message) for item in captured]
    if any(any(token in message.lower() for token in (
            "not supported", "unsupported", "will be removed", "will be lost"
    )) for message in messages):
        raise ExcelExportError("文件含有当前导出器无法保留的 Excel 扩展，已取消导出。", "\n".join(messages))


def _publish_new_file(temporary: Path, target: Path) -> None:
    """发布时仍不能覆盖：Windows rename 拒绝已存在文件，其他平台使用独占硬链接。"""
    if os.name == "nt":
        os.rename(temporary, target)
    else:
        os.link(temporary, target)
        temporary.unlink()


_WHITESPACE_TEXT = re.compile(rb'<(?P<tag>(?:[\w.-]+:)?t)(?P<attributes>\s[^>]*|)>(?P<value>[ \t\r\n]+)</(?P=tag)>')


def _preserve_xml_text(path: Path) -> None:
    """保留回车和纯空白字符串，修正 XML 写入后 Excel 可能丢失的文本语义。"""
    with ZipFile(path) as archive:
        affected = {
            name for name in archive.namelist()
            if name.endswith(".xml") and (b"\r" in archive.read(name) or _WHITESPACE_TEXT.search(archive.read(name)))
        }
        if not affected:
            return
        descriptor, name = tempfile.mkstemp(prefix=".excel-cleaner-", suffix=".xml.tmp.xlsx", dir=path.parent)
        os.close(descriptor)
        normalized = Path(name)
        try:
            # 不同 XML 后端/版本的换行行为不同；用同一后端探测，避免把正常 LF 误加为 CRLF。
            probe = Element("text")
            probe.text = "\n"
            with xmlfile(str(normalized)) as writer:
                writer.write(probe)
            expands_newlines = b"\r\n" in normalized.read_bytes()
            with ZipFile(normalized, "w") as destination:
                for info in archive.infolist():
                    content = archive.read(info.filename)
                    if info.filename in affected:
                        if expands_newlines and info.filename.startswith("xl/worksheets/"):
                            content = content.replace(b"\r\n", b"\n")
                        def mark_space(match):
                            attributes = match["attributes"]
                            if b"xml:space=" not in attributes:
                                attributes += b' xml:space="preserve"'
                            return b"<" + match["tag"] + attributes + b">" + match["value"] + b"</" + match["tag"] + b">"
                        content = _WHITESPACE_TEXT.sub(mark_space, content)
                        content = content.replace(b"\r", b"&#13;")
                    destination.writestr(info, content)
            # 替换的只是本任务自己创建的临时文件，不是目标文件或源文件。
        except BaseException:
            normalized.unlink(missing_ok=True)
            raise
    try:
        os.replace(normalized, path)
    finally:
        normalized.unlink(missing_ok=True)


def _write_updated_columns(workbook, plan: ExportPlan) -> None:
    for change in plan.changes:
        cell = workbook[change.sheet_name].cell(change.excel_row, change.excel_column)
        value = change.new_value
        if bool(pd.isna(value)):
            value = None
        elif hasattr(value, "item"):
            value = value.item()
        cell.value = value
        if isinstance(value, str):
            # 清洗后以 '=' 开头的原文本仍是文本，不能意外转换为 Excel 公式。
            cell.data_type = "s"


def _write_comparison_columns(workbook, plan: ExportPlan) -> dict[str, ColumnMap]:
    if not plan.changes:
        return {}
    columns = {}
    for change in plan.changes:
        columns.setdefault(change.sheet_name, set()).add(change.excel_column)
    maps = {name: ColumnMap(tuple(sorted(indices))) for name, indices in columns.items()}
    snapshots = {sheet.sheet_name: sheet for sheet in plan.snapshot.sheets}
    formulas = {}
    literal_outputs = set()
    for name, mapping in maps.items():
        snapshot = snapshots[name]
        sheet = workbook[name]
        for column in mapping.after_columns:
            rules = snapshot.column_rules.get(column - 1)
            if not rules:
                raise ExcelExportError("缺少已应用的清洗规则，无法生成公式，请选择“原列更新”。", name)
            # 每列只解析一次公式模板，较多数据行时避免重复构造和解析长公式。
            template_reference = f"{get_column_letter(mapping.column(column))}{dataframe_index_to_excel_row(len(snapshot.original_df) - 1)}"
            try:
                use_values = any(step.has_type_rules for step in rules)
                template = None if use_values else build_export_formula(rules, template_reference)
            except ValueError as error:
                raise ExcelExportError(str(error), name) from error
            for position in range(len(snapshot.original_df)):
                row = dataframe_index_to_excel_row(position)
                original = snapshot.original_df.iat[position, column - 1]
                working = snapshot.working_df.iat[position, column - 1]
                reference = f"{get_column_letter(mapping.column(column))}{row}"
                if sheet.cell(row, column).data_type == "f":
                    formulas[name, row, column] = build_passthrough_formula(reference)
                    continue
                value = original
                if not use_values and isinstance(value, str) and "\x01" in value:
                    raise ExcelExportError("原文本含有当前公式不支持的控制字符，请选择“原列更新”。", name)
                for step in rules:
                    value = apply_cleaning_rules(value, step)
                if not _values_equal(value, working):
                    raise ExcelExportError("已应用规则与工作数据不一致，无法生成公式，请选择“原列更新”。", name)
                if use_values:
                    formulas[name, row, column] = working
                    literal_outputs.add((name, row, column))
                else:
                    formulas[name, row, column] = template.replace(template_reference, reference)
    try:
        insert_comparison_columns(workbook, maps)
    except ColumnInsertionError as error:
        raise ExcelExportError(str(error)) from error
    for name, mapping in maps.items():
        sheet = workbook[name]
        snapshot = snapshots[name]
        used_names = {str(cell.value).casefold() for cell in sheet[1] if cell.value is not None}
        for column in mapping.after_columns:
            source_column = mapping.column(column)
            target_column = source_column + 1
            base = str(snapshot.original_df.columns[column - 1]) + "_清洗后"
            title, number = base, 0
            while title.casefold() in used_names:
                number += 1
                title = f"{base}_{number}"
            used_names.add(title.casefold())
            header = sheet.cell(1, target_column)
            header._style = copy(sheet.cell(1, source_column)._style)
            header.value = title
            header.data_type = "s"
            for position in range(len(snapshot.original_df)):
                row = dataframe_index_to_excel_row(position)
                cell = sheet.cell(row, target_column)
                cell._style = copy(sheet.cell(row, source_column)._style)
                cell.value = formulas[name, row, column]
                if (name, row, column) in literal_outputs:
                    # 类型转换由 Python 保证口径，输出已确认的值；不能将 "=..." 文本当公式。
                    cell.number_format = "General"
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
    if workbook.calculation is None:
        workbook.calculation = CalcProperties()
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return maps


def _write_annotation_columns(workbook, plan: ExportPlan, maps: dict[str, ColumnMap]) -> None:
    """在实际数据区域右侧追加，单独导出辅助列也不触碰任何原始单元格。"""
    for snapshot in plan.snapshot.sheets:
        if not snapshot.annotations:
            continue
        sheet = workbook[snapshot.sheet_name]
        mapping = maps.get(snapshot.sheet_name, ColumnMap(()))
        first_column = snapshot.original_df.shape[1] + len(mapping.after_columns) + 1
        last_column = first_column + len(snapshot.annotations) - 1
        last_row = len(snapshot.working_df) + 1
        if last_column > 16384:
            raise ExcelExportError("导出后列数超过 Excel 最大列数，已取消导出。", snapshot.sheet_name)
        for area in sheet.merged_cells.ranges:
            if area.min_col <= last_column and area.max_col >= first_column and area.min_row <= last_row:
                raise ExcelExportError("新增辅助列与原有合并区域重叠，已取消导出。", snapshot.sheet_name)
        # 不把格式化的尾部空列视为业务数据，也不覆盖右侧已有文字或公式。
        for row in range(1, last_row + 1):
            for column in range(first_column, last_column + 1):
                if sheet.cell(row, column).value is not None:
                    raise ExcelExportError("辅助列目标位置存在原有数据，已取消导出。", snapshot.sheet_name)
        for position, annotation in enumerate(snapshot.annotations):
            target_column = first_column + position
            source_original = annotation.source_column_index + 1
            source_column = mapping.column(source_original)
            if source_original in mapping.after_columns:
                # 对照模式下原始脏数据保留；辅助列应引用代表 working_df 的清洗结果列。
                source_column += 1
            header = sheet.cell(1, target_column)
            header._style = copy(sheet.cell(1, source_column)._style)
            header.value = annotation.column_name
            header.data_type = "s"
            sheet.column_dimensions[get_column_letter(target_column)].width = max(
                18, min(48, len(annotation.column_name) * 1.5 + 2),
            )
            for row in range(2, last_row + 1):
                cell = sheet.cell(row, target_column)
                cell._style = copy(sheet.cell(row, source_column)._style)
                cell.number_format = "General"
                if annotation.kind == "duplicate":
                    cell.value = build_duplicate_formula(source_column - 1, row, last_row)
                else:
                    cell.value = snapshot.working_df.iat[row - 2, snapshot.original_df.shape[1] + position]
                    cell.data_type = "s"
    if plan.annotation_column_count:
        if workbook.calculation is None:
            workbook.calculation = CalcProperties()
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True


def export_workbook(plan: ExportPlan, target_path: str | Path) -> ExportResult:
    """所有保存操作只针对随机临时文件；源文件只读，失败不发布目标文件。"""
    if not plan.changes and not plan.annotation_column_count:
        raise ExcelExportError(NO_CHANGES_MESSAGE)
    temporary = None
    workbook = None
    try:
        # 重新计算以防调用方误改计划；界面确认后的版本检查由数据管理器提供。
        if get_changes(plan.snapshot) != plan.changes:
            raise ExcelExportError("导出快照已变化，请重新确认导出。")
        source, target = validate_export_paths(plan.snapshot.source_path, target_path)
        digest = file_sha256(source)
        if plan.snapshot.source_digest is not None and digest != plan.snapshot.source_digest:
            raise ExcelExportError("源文件在打开后发生变化，请重新打开文件后再导出。")
        descriptor, name = tempfile.mkstemp(prefix=".excel-cleaner-", suffix=".tmp.xlsx", dir=target.parent)
        os.close(descriptor)
        temporary = Path(name)
        # 不继承源文件的只读权限，且绝不在源文件上调用 save。
        shutil.copyfile(source, temporary)
        if file_sha256(temporary) != digest:
            raise ExcelExportError("源文件在复制期间发生变化，已取消导出，请重新打开文件。")
        with warnings.catch_warnings(record=True) as captured:
            warnings.simplefilter("always")
            _validate_source_data(temporary, plan.snapshot)
            # 必须保留公式表达式，不以公式的缓存结果覆盖原公式。
            workbook = load_workbook(temporary, data_only=False, keep_links=True, rich_text=True)
            _reject_loss_warnings(captured)
            for change in plan.changes:
                sheet = workbook[change.sheet_name]
                if (change.excel_row < 2 or change.excel_row > sheet.max_row
                        or change.excel_column < 1 or change.excel_column > sheet.max_column):
                    raise ExcelExportError("单元格坐标超出源工作表范围，已取消导出。", change.sheet_name)
                cell = sheet.cell(change.excel_row, change.excel_column)
                if cell.data_type == "f":
                    raise ExcelExportError(
                        "该单元格包含 Excel 公式，当前版本不会自动清洗公式单元格。",
                        f"工作表：{change.sheet_name}，单元格：{cell.coordinate}",
                    )
                if isinstance(cell, MergedCell):
                    raise ExcelExportError("不能修改合并区域中的非首个单元格，已取消导出。", change.sheet_name)
            maps = {}
            if plan.mode == ExportMode.UPDATE_COLUMNS:
                _write_updated_columns(workbook, plan)
            elif plan.mode == ExportMode.ADD_FORMULA_COLUMNS:
                maps = _write_comparison_columns(workbook, plan)
            else:
                raise ExcelExportError("请选择有效的导出模式。")
            _write_annotation_columns(workbook, plan, maps)
            workbook.save(temporary)
            _reject_loss_warnings(captured)
        workbook.close()
        workbook = None
        _preserve_xml_text(temporary)
        if file_sha256(source) != digest:
            raise ExcelExportError("源文件在导出期间发生变化，已取消导出，请重新打开文件。")
        _publish_new_file(temporary, target)
        temporary = None
        return ExportResult(target, plan.sheet_count, plan.cell_count, plan.mode,
                            plan.column_count if plan.mode == ExportMode.ADD_FORMULA_COLUMNS else 0,
                            plan.annotation_column_count)
    except ExcelExportError:
        raise
    except FileExistsError as error:
        raise ExcelExportError("目标文件已经存在，导出已取消；已有文件不会被覆盖。", str(error)) from error
    except PermissionError as error:
        raise ExcelExportError("无法保存文件，请检查目录权限，或关闭占用文件的 Excel 后重试。", str(error)) from error
    except FileNotFoundError as error:
        raise ExcelExportError("源文件或保存目录不存在，请重新选择。", str(error)) from error
    except Exception as error:
        raise ExcelExportError("导出失败，源文件保持不变；请检查文件有效性及保存位置后重试。", str(error)) from error
    finally:
        if workbook is not None:
            workbook.close()
        if temporary is not None:
            try:
                temporary.unlink(missing_ok=True)
            except OSError as error:
                raise ExcelExportError(
                    "导出未完成，源文件保持不变，但临时文件无法删除，请关闭占用程序后删除。",
                    str(temporary),
                ) from error
